"""
_docx_base.py — Shared utilities for template generators.

Provides common functions for building DOCX papers from JSON configs.
Extracted from IEEEgen.py patterns to be reused by other template generators.
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

XSL_CANDIDATES = [
    Path("/usr/share/docbook-xsl/docbook-xsl-1.79.2/fo/docbook.xsl"),
    Path("/usr/share/xml/docbook/stylesheet/docbook-xsl/fo/docbook.xsl"),
]

NS_MAP = {
    b"http://purl.oclc.org/ooxml/wordprocessingml/main": b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing": b"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    b"http://purl.oclc.org/ooxml/officeDocument/math": b"http://schemas.openxmlformats.org/officeDocument/2006/math",
}

_XSLT = None


def _strict_to_trans(data: bytes) -> bytes:
    for old, new in NS_MAP.items():
        data = data.replace(old, new)
    return data


def _pt2tw(pt: float) -> int:
    return int(round(pt * 20))


def open_template(template_path: Path) -> Document:
    doc = Document()
    with zipfile.ZipFile(template_path) as archive:
        raw = archive.read("word/styles.xml")
    tmpl_styles = etree.fromstring(_strict_to_trans(raw))
    cur = doc.part.styles._element
    for child in list(cur):
        cur.remove(child)
    for child in list(tmpl_styles):
        cur.append(child)
    try:
        if hasattr(doc.part.styles, "_styles"):
            doc.part.styles._styles = None
    except Exception:
        pass
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc


def finalize_doc(doc: Document) -> None:
    pass


def set_para_style(paragraph, style_name: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style_name)


def para(
    doc: Document,
    style_id: str | None = None,
    align=None,
    sb: float | None = None,
    sa: float | None = None,
    fi: float | None = None,
    li: float | None = None,
):
    paragraph = doc.add_paragraph()
    if style_id:
        set_para_style(paragraph, style_id)
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    if sb is not None:
        pf.space_before = Pt(sb)
    if sa is not None:
        pf.space_after = Pt(sa)
    if fi is not None:
        pf.first_line_indent = Pt(fi)
    if li is not None:
        pf.left_indent = Pt(li)
    return paragraph


def build_sectpr(
    num_cols: int,
    col_space_pt: float,
    top_pt: float,
    bottom_pt: float,
    left_pt: float,
    right_pt: float,
    section_type: str | None = "continuous",
    w_pt: float = 595.3,
    h_pt: float = 841.9,
    header_pt: float = 36.0,
    footer_pt: float = 36.0,
    title_pg: bool = False,
):
    ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    def wq(tag):
        return f"{{{ns_w}}}{tag}"
    sectpr = etree.Element(wq("sectPr"))
    if section_type:
        sec_type = etree.SubElement(sectpr, wq("type"))
        sec_type.set(wq("val"), section_type)
    pg_sz = etree.SubElement(sectpr, wq("pgSz"))
    pg_sz.set(wq("w"), str(_pt2tw(w_pt)))
    pg_sz.set(wq("h"), str(_pt2tw(h_pt)))
    pg_mar = etree.SubElement(sectpr, wq("pgMar"))
    pg_mar.set(wq("top"), str(_pt2tw(top_pt)))
    pg_mar.set(wq("right"), str(_pt2tw(right_pt)))
    pg_mar.set(wq("bottom"), str(_pt2tw(bottom_pt)))
    pg_mar.set(wq("left"), str(_pt2tw(left_pt)))
    pg_mar.set(wq("header"), str(_pt2tw(header_pt)))
    pg_mar.set(wq("footer"), str(_pt2tw(footer_pt)))
    pg_mar.set(wq("gutter"), "0")
    cols = etree.SubElement(sectpr, wq("cols"))
    if num_cols > 1:
        cols.set(wq("num"), str(num_cols))
    cols.set(wq("space"), str(_pt2tw(col_space_pt)))
    if title_pg:
        etree.SubElement(sectpr, wq("titlePg"))
    return sectpr


def embed_sectpr(doc: Document, sectpr_el, style_id: str | None = None):
    paragraph = para(doc, style_id=style_id)
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(sectpr_el)
    return paragraph


def setup_main_sectpr(
    doc: Document,
    w_pt: float = 595.3,
    h_pt: float = 841.9,
    top_pt: float = 54.0,
    bottom_pt: float = 72.0,
    left_pt: float = 44.65,
    right_pt: float = 44.65,
    header_pt: float = 36.0,
    footer_pt: float = 36.0,
    col_space_pt: float = 36.0,
    num_cols: int = 1,
    section_type: str = "continuous",
) -> None:
    body = doc.element.body
    sectpr = body.find(qn("w:sectPr"))
    if sectpr is None:
        sectpr = OxmlElement("w:sectPr")
        body.append(sectpr)
    for tag in ("w:cols", "w:pgSz", "w:pgMar", "w:type", "w:titlePg"):
        for old in sectpr.findall(qn(tag)):
            sectpr.remove(old)
    pg_sz = OxmlElement("w:pgSz")
    pg_sz.set(qn("w:w"), str(_pt2tw(w_pt)))
    pg_sz.set(qn("w:h"), str(_pt2tw(h_pt)))
    sectpr.append(pg_sz)
    pg_mar = OxmlElement("w:pgMar")
    pg_mar.set(qn("w:top"), str(_pt2tw(top_pt)))
    pg_mar.set(qn("w:right"), str(_pt2tw(right_pt)))
    pg_mar.set(qn("w:bottom"), str(_pt2tw(bottom_pt)))
    pg_mar.set(qn("w:left"), str(_pt2tw(left_pt)))
    pg_mar.set(qn("w:header"), str(_pt2tw(header_pt)))
    pg_mar.set(qn("w:footer"), str(_pt2tw(footer_pt)))
    pg_mar.set(qn("w:gutter"), "0")
    sectpr.append(pg_mar)
    cols = OxmlElement("w:cols")
    if num_cols > 1:
        cols.set(qn("w:num"), str(num_cols))
    cols.set(qn("w:space"), str(_pt2tw(col_space_pt)))
    sectpr.append(cols)
    sec_type = OxmlElement("w:type")
    sec_type.set(qn("w:val"), section_type)
    sectpr.append(sec_type)


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag == "oMath":
        paragraph._p.append(omml)
    elif tag == "oMathPara":
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def append_text_run(
    paragraph, text: str, bold: bool = False, italic: bool = False, underline: bool = False
):
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline


def append_line_break(paragraph):
    paragraph.add_run().add_break()


def _normalize_text_commands(text: str) -> str:
    text = text.replace("\\n", "\n")
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def append_rich_text(paragraph, text: str):
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            append_line_break(paragraph)
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = paragraph.add_run(token["value"])
                run.italic = True
            continue
        append_text_run(
            paragraph,
            token["value"],
            bold=token["bold"],
            italic=token["italic"],
            underline=token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [""]


def body_paragraphs(doc: Document, text: str, style_id: str = "BodyText"):
    paragraphs = []
    for block in _split_body_blocks(text):
        paragraph = para(doc, style_id=style_id)
        append_rich_text(paragraph, block)
        paragraphs.append(paragraph)
    return paragraphs


def roman(number) -> str:
    try:
        value = int(str(number).strip())
    except Exception:
        return str(number).strip()
    numerals = [
        (1000, "M"),
        (900, "CM"),
        (500, "D"),
        (400, "CD"),
        (100, "C"),
        (90, "XC"),
        (50, "L"),
        (40, "XL"),
        (10, "X"),
        (9, "IX"),
        (5, "V"),
        (4, "IV"),
        (1, "I"),
    ]
    result = ""
    for val, numeral in numerals:
        while value >= val:
            result += numeral
            value -= val
    return result


def _resolve_path(path_text: str, json_path: Path) -> Path:
    p = Path(path_text)
    if p.is_absolute():
        return p
    return json_path.parent / p


def _set_full_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        el = tc_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def _set_table_borders(table, full: bool = False):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    if full:
        visible_sides = {"top", "left", "bottom", "right", "insideH", "insideV"}
    else:
        visible_sides = {"top", "bottom", "insideH", "insideV"}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        if edge in visible_sides:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
            el.set(qn("w:sz"), "0")


def _style_cell_paragraph(paragraph, style_id: str, align=WD_ALIGN_PARAGRAPH.CENTER):
    set_para_style(paragraph, style_id)
    paragraph.alignment = align


def _add_equation_line(doc: Document, formula: str, number: str | None = None, style_id: str = "Equation0"):
    p = para(doc, style_id=style_id)
    omml = _latex_to_omml(formula)
    if omml is not None:
        tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
        if tag == "oMath":
            p._p.append(omml)
        elif tag == "oMathPara":
            p._p.append(omml)
        else:
            wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
            wrapper.append(omml)
            p._p.append(wrapper)
    else:
        p.add_run(formula).italic = True
    if number:
        p.add_run(f"   ({number})")


def _add_prompt_box_with_text(doc: Document, text: str, full_borders: bool = True):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Normal Table"
    _set_table_borders(table, full=full_borders)
    cell = table.cell(0, 0)
    _set_full_cell_borders(cell)
    paragraph = cell.paragraphs[0]
    set_para_style(paragraph, "BodyText")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    append_rich_text(paragraph, text)
    return table


def _render_content_item(doc: Document, item: dict, json_path: Path, cfg: dict | None = None):
    if cfg is None:
        cfg = {}
    item_id = str(item.get("id", "")).lower()
    body_style = cfg.get("body", "BodyText")
    fig_caption_style = cfg.get("figure_caption", "figurecaption")
    eq_style = cfg.get("equation", "Equation0")
    table_head_style = cfg.get("table_head", "tablehead")
    table_col_head_style = cfg.get("table_col_head", "tablecolhead")
    table_copy_style = cfg.get("table_copy", "tablecopy")
    max_fig_width = cfg.get("max_fig_width_cm", 8.4)
    full_borders = cfg.get("full_borders", False)
    table_auto_label = cfg.get("table_auto_label", False)
    figure_auto_label = cfg.get("figure_auto_label", False)

    if item_id == "text":
        text = str(item.get("text", ""))
        if text:
            body_paragraphs(doc, text, style_id=body_style)

    elif item_id in ("gambar", "image"):
        image_number = str(item.get("ImageNumber", "")).strip()
        path_text = str(item.get("Path", "")).strip()
        prompt = str(item.get("Prompt", "")).strip()
        title = str(item.get("Title", "")).strip()
        caption_text = title

        image_path = _resolve_path(path_text, json_path) if path_text else None
        try:
            width_cm = float(item.get("WidthCm", max_fig_width))
        except Exception:
            width_cm = max_fig_width
        width_cm = max(1.0, min(width_cm, max_fig_width))

        if image_path is not None and image_path.is_file():
            paragraph = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, sb=6, sa=2)
            paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
        else:
            prompt_body = prompt if prompt else f"Figure {image_number} not found"
            if title:
                fallback_text = f"[PROMPT UNTUK AI GAMBAR: {title}. {prompt_body}]"
            else:
                fallback_text = f"[PROMPT UNTUK AI GAMBAR: {prompt_body}]"
            _add_prompt_box_with_text(doc, fallback_text, full_borders=full_borders)

        if image_number and caption_text:
            if figure_auto_label:
                cap = para(doc, style_id=fig_caption_style, align=WD_ALIGN_PARAGRAPH.CENTER)
                append_rich_text(cap, caption_text)
            else:
                cap = para(doc, style_id=fig_caption_style, align=WD_ALIGN_PARAGRAPH.CENTER)
                append_rich_text(cap, f"Fig. {image_number}. {caption_text}".strip())

    elif item_id in ("rumus", "formula"):
        formula_number = str(item.get("FormulaNumber", "")).strip()
        formula_text = str(item.get("text", "") or item.get("latex", "")).strip()
        if formula_text:
            _add_equation_line(doc, formula_text, formula_number if formula_number else None, style_id=eq_style)

    elif item_id in ("tabel", "table"):
        table_number = str(item.get("TableNumber", "")).strip()
        title = str(item.get("Title", "")).strip()
        headers = list(item.get("Headers", []))
        rows = list(item.get("Rows", []))

        if not headers:
            return

        if table_auto_label:
            caption = para(doc, style_id=table_head_style)
            append_rich_text(caption, title if title else "")
        else:
            caption = para(doc, style_id=table_head_style)
            label = f"TABLE {roman(table_number)}" if table_number else "TABLE"
            text = f"{label}. {title}" if title else label
            append_rich_text(caption, text)

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.style = "Normal Table"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_borders(table, full=full_borders)

        for column_index, value in enumerate(headers):
            cell = table.rows[0].cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _style_cell_paragraph(paragraph, table_col_head_style)
            append_rich_text(paragraph, str(value))
            for run in paragraph.runs:
                run.bold = True

        for row_index, row_data in enumerate(rows, start=1):
            for column_index, value in enumerate(row_data):
                if column_index >= len(headers):
                    break
                cell = table.rows[row_index].cells[column_index]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                _style_cell_paragraph(paragraph, table_copy_style)
                append_rich_text(paragraph, str(value))

        para(doc, sa=4)


def render_sections(doc: Document, config: dict, json_path: Path, base_dir: Path, cfg: dict | None = None):
    if cfg is None:
        cfg = {}
    heading1_style = cfg.get("heading1", "Heading1")
    heading2_style = cfg.get("heading2", "Heading2")
    body_style = cfg.get("body", "BodyText")
    section_heading_format = cfg.get("section_heading_format", "roman_dot")
    subsection_no_prefix = cfg.get("subsection_no_prefix", False)

    section_keys = []
    for key in config.keys():
        if key.startswith("section") and key.replace("section", "").isdigit():
            section_keys.append(key)
    section_keys.sort(key=lambda x: int(x.replace("section", "")))

    for section_key in section_keys:
        section = config[section_key]
        section_number = str(section.get("number", "")).strip()
        if not section_number:
            num_part = section_key.replace("section", "")
            if num_part.isdigit():
                section_number = num_part
        section_title = str(section.get("title", "")).strip().upper()

        if section_heading_format == "plain_upper":
            heading_text = section_title
        elif section_number:
            heading_text = f"{roman(section_number)}. {section_title}"
        else:
            heading_text = section_title

        h = para(doc, style_id=heading1_style)
        h.add_run(heading_text)

        content = section.get("content", "")
        if isinstance(content, str) and content.strip():
            body_paragraphs(doc, content.strip(), style_id=body_style)
        elif isinstance(content, list):
            for item in content:
                if isinstance(item, str):
                    body_paragraphs(doc, item.strip(), style_id=body_style)
                elif isinstance(item, dict):
                    _render_content_item(doc, item, json_path, cfg)

        subsection_keys = [
            key
            for key in section.keys()
            if (key.startswith("sub") and len(key) > 3 and key[3:].isalnum())
            or re.match(r"^section\d+[a-z]+$", key)
        ]
        for key in subsection_keys:
            subsection = section[key]
            subsection_letter = str(subsection.get("letter", "")).strip()
            if not subsection_letter and key:
                m = re.match(r"^(?:sub|section)\d+([a-z]+)$", key)
                if m:
                    subsection_letter = m.group(1).upper()
            subsection_title = str(subsection.get("title", "")).strip()

            if subsection_no_prefix:
                heading_text = subsection_title
            elif subsection_letter and subsection_title:
                heading_text = f"{subsection_letter}. {subsection_title}"
            elif subsection_title:
                heading_text = subsection_title
            else:
                heading_text = ""

            if heading_text:
                h2 = para(doc, style_id=heading2_style)
                h2.add_run(heading_text)

            content_items = subsection.get("content", [])
            if isinstance(content_items, list):
                for item in content_items:
                    if isinstance(item, str):
                        body_paragraphs(doc, item.strip(), style_id=body_style)
                    elif isinstance(item, dict):
                        _render_content_item(doc, item, json_path, cfg)
            elif isinstance(content_items, str):
                if content_items.strip():
                    body_paragraphs(doc, content_items.strip(), style_id=body_style)


def run_generator(base_dir: Path, template_path: Path, build_fn):
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        if not json_arg.exists():
            print(f"[ERR] File not found: {json_arg}")
            return
        result = build_fn(json_arg, output_arg)
        print(f"[OK] Generated: {result.name}")
    else:
        json_files = list(base_dir.glob("*.json"))
        if not json_files:
            print("[ERR] No JSON files found")
            return
        for json_path in sorted(json_files):
            if json_path.name.lower() in ("package.json", "tsconfig.json", "settings.json"):
                continue
            try:
                data = json.loads(json_path.read_text(encoding="utf-8"))
                if not isinstance(data, dict):
                    continue
                result = build_fn(json_path)
                print(f"[OK] Generated: {result.name}")
            except Exception as e:
                print(f"[ERR] {json_path.name}: {e}")
