"""
Data extraction worker — async processing of uploaded files.

Flow:
  1. Extract text from each file (PDF/Excel/CSV/DOCX)
  2. Send combined text + user prompt to AI for formatting
  3. AI returns: tables + chart_recommendations + analysis per table + analysis per chart
  4. Auto-generate chart PNGs from recommendations
  5. Save everything to DB + filesystem
  6. Publish SSE progress events throughout

Uses AiJob(kind="data_extract") for persistence.
"""

from __future__ import annotations

import json
import logging
import os
import shutil
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

import threading
import redis
import traceback

log = logging.getLogger(__name__)

_REDIS = None
_REDIS_LOCK = threading.Lock()

def get_redis(_retry=0, _max_retries=3):
    global _REDIS
    if _REDIS is None:
        with _REDIS_LOCK:
            if _REDIS is None:
                try:
                    _REDIS = redis.Redis.from_url(os.getenv("REDIS_URL", "redis://localhost:6379/0"), decode_responses=True)
                except Exception:
                    return None
    # Health check: if connection went stale (Redis restart etc.), reconnect
    try:
        _REDIS.ping()
    except Exception:
        with _REDIS_LOCK:
            try:
                _REDIS.close()
            except Exception:
                pass
            _REDIS = None
        if _retry >= _max_retries:
            log.warning("get_redis: ping failed after %d retries, giving up", _max_retries)
            return None
        return get_redis(_retry + 1, _max_retries)
    return _REDIS


def _publish(job_id, payload):
    _r = get_redis()
    if not _r:
        return
    try:
        _r.publish(f"datajob:{job_id}:progress", json.dumps(payload, default=str))
    except Exception:
        pass
    try:
        _r.setex(f"datajob:{job_id}:last", 86400, json.dumps(payload, default=str))
    except Exception:
        pass


def _is_cancelled(job_id):
    _r = get_redis()
    if not _r:
        return False
    try:
        return bool(_r.get(f"datajob:{job_id}:cancel"))
    except Exception:
        return False


def _clean_redis_job_keys(job_id: str) -> int:
    """Delete all Redis keys for a data job. Called on cleanup/error/cancel."""
    count = 0
    try:
        r = get_redis()
        if r:
            for suffix in ("last", "cancel"):
                key = f"datajob:{job_id}:{suffix}"
                if r.delete(key):
                    count += 1
    except Exception:
        pass
    return count


def _cleanup_stale_worker_state(job_id: str):
    """Check if job is still in running/queued state in DB.
    If so, mark as error and clean Redis.
    If job is already in terminal state (done/error/cancelled), skip.
    """
    try:
        from database.models import AiJob, db, safe_commit
        j = AiJob.query.get(job_id)
    except Exception:
        return
    if not j:
        _clean_redis_job_keys(job_id)
        return
    if j.status in ("done", "error", "cancelled"):
        # Already terminal — just clean Redis to be safe
        _clean_redis_job_keys(job_id)
        return
    # Still running/queued → mark as error
    j.status = "error"
    j.stage = "error"
    j.error = "Worker thread terminated unexpectedly."
    j.finished_at = datetime.now(timezone.utc)
    try:
        safe_commit()
    except Exception:
        try:
            db.session.rollback()
        except Exception:
            pass
    _publish(job_id, {
        "status": "error",
        "stage": "error",
        "progress": 0,
        "error": "Worker thread terminated unexpectedly.",
        "message": "Worker thread terminated.",
    })
    _clean_redis_job_keys(job_id)
    log.warning("DataJob %s: stale worker state cleaned (was %s)", job_id, j.status)


def run_data_job(app, job_id, file_paths, file_names, paper_id, user_id, user_prompt):
    """Main worker entry point. Runs in a daemon thread."""
    log.debug("DATA_WORKER_DEBUG: run_data_job ENTER — job_id=%s file_count=%d text_only=%s",
              job_id, len(file_paths), len(file_paths) == 0)
    # Publish immediately so SSE has something to show right away
    # (overwrites the initial 'queued' state from create_data_job)
    _publish(job_id, {
        "status": "running",
        "stage": "starting",
        "progress": 1,
        "message": "Memulai pemrosesan data...",
        "file_count": len(file_paths),
        "text_only": len(file_paths) == 0,
    })
    log.debug("DATA_WORKER_DEBUG: After initial _publish — job_id=%s", job_id)
    with app.app_context():
        try:
            from database.models import AiJob, Paper, PaperImage, db, safe_commit

            job = AiJob.query.get(job_id)
            if not job:
                log.error("DataJob %s not found", job_id)
                _clean_redis_job_keys(job_id)
                return

            # Mark running
            job.status = "running"
            # Skip extraction stage entirely when no files (text-only mode)
            if file_paths:
                job.stage = "extracting"
                job.progress = 5
                try:
                    safe_commit()
                except Exception:
                    db.session.rollback()
                    raise
                _publish(job_id, {"status": "running", "stage": "extracting", "progress": 5, "message": "Memulai ekstraksi file..."})

                if _is_cancelled(job_id):
                    _finish_cancelled(job)
                    return

                # ── Step 1: Extract text from files ──────────────────────────
                extracted_texts = []
                for i, (fpath, fname) in enumerate(zip(file_paths, file_names)):
                    if _is_cancelled(job_id):
                        _finish_cancelled(job)
                        return

                    prog = 5 + int(15 * (i + 1) / len(file_paths))
                    _publish(job_id, {
                        "status": "running",
                        "stage": "extracting",
                        "progress": prog,
                        "message": f"Mengekstrak {fname} ({i+1}/{len(file_paths)})...",
                        "current_file": fname,
                        "file_index": i,
                        "file_total": len(file_paths),
                    })

                    try:
                        text = _extract_file_text(fpath, fname)
                        if text:
                            extracted_texts.append(f"## File: {fname}\n{text}")
                    except Exception as e:
                        log.warning("Failed to extract %s: %s", fname, e)
                        extracted_texts.append(f"## File: {fname}\n[Gagal mengekstrak: {e}]")

                # Cleanup temp files
                for fpath in file_paths:
                    try:
                        Path(fpath).unlink(missing_ok=True)
                    except Exception:
                        pass
                try:
                    parent = Path(file_paths[0]).parent
                    if parent.name.startswith("datajob_"):
                        shutil.rmtree(parent, ignore_errors=True)
                except Exception:
                    pass

                if not extracted_texts:
                    _finish_error(job, "Tidak ada teks yang berhasil diekstrak dari file.")
                    return

                combined_text = "\n\n".join(extracted_texts)
                if combined_text and len(combined_text) > 15000:
                    combined_text = combined_text[:15000] + "\n\n... [data terpotong]"

                _publish(job_id, {"status": "running", "stage": "extracting", "progress": 20,
                                 "message": f"Berhasil mengekstrak {len(file_paths)} file."})
            else:
                # Text-only mode: skip extraction, go straight to AI formatting
                combined_text = ""
                _publish(job_id, {"status": "running", "stage": "ai_formatting", "progress": 20,
                                 "message": "Mode teks — langsung ke AI formatting..."})

            if _is_cancelled(job_id):
                _finish_cancelled(job)
                return

            # ── Step 2: AI formatting ────────────────────────────────────────
            # FIX: Do NOT close session - just commit and let SQLAlchemy manage connections
            try:
                safe_commit()
            except Exception:
                db.session.rollback()
                raise
            _publish(job_id, {"status": "running", "stage": "ai_formatting", "progress": 25, "message": "Menganalisis data dengan AI..."})

            from tools.data.dataFormating import format_data_with_ai

            # Heartbeat thread: publish progress every 10s during AI call
            # so SSE stream doesn't appear stuck and frontend knows job is alive
            _ai_stop = threading.Event()
            def _heartbeat():
                prog = 25
                while not _ai_stop.wait(10):
                    prog = min(prog + 2, 48)  # slowly advance to 48%
                    _publish(job_id, {"status": "running", "stage": "ai_formatting", "progress": prog, "message": "AI sedang menganalisis data..."})
            _hb_thread = threading.Thread(target=_heartbeat, daemon=True)
            _hb_thread.start()

            # Overall deadline timer: if AI call hangs > 5 minutes, forcibly
            # publish error and stop heartbeat. Prevents permanent stuck state
            # when all endpoints are slow/unreachable.
            _AI_DEADLINE = 300  # 5 minutes max for AI formatting
            _ai_timeout_flag = [False]

            def _ai_deadline_handler():
                _ai_timeout_flag[0] = True
                _ai_stop.set()
                _publish(job_id, {
                    "status": "error",
                    "stage": "ai_formatting",
                    "progress": 49,
                    "error": "AI formatting timeout (>5 menit) — semua endpoint tidak responsif.",
                    "message": "AI formatting timeout.",
                })

            _deadline_timer = threading.Timer(_AI_DEADLINE, _ai_deadline_handler)
            _deadline_timer.daemon = True
            _deadline_timer.start()

            try:
                if file_paths:
                    # File mode: combine user prompt with extracted file text
                    ai_prompt = combined_text
                    if user_prompt:
                        ai_prompt = f"### Instruksi User:\n{user_prompt}\n\n### Data dari file:\n{combined_text}"
                    result = format_data_with_ai(ai_prompt, filename=", ".join(file_names))
                else:
                    # Text-only mode: send user prompt directly
                    if not user_prompt:
                        _finish_error(job, "Tidak ada instruksi atau file yang diberikan.")
                        return
                    result = format_data_with_ai(user_prompt, filename="Text Input")
            finally:
                _deadline_timer.cancel()
                _ai_stop.set()

            # If AI deadline timer fired before we got a result
            if _ai_timeout_flag[0]:
                _finish_error(job, "AI formatting timeout (>5 menit) — semua endpoint tidak responsif.")
                return

            tables = result.get("tables", [])
            recs = result.get("chart_recommendations", [])

            # Text-only mode: allow chart creation even without tables.
            # Users can request charts directly via "Generate dari Teks".
            is_text_only = len(file_paths) == 0

            if not tables and is_text_only and recs:
                # Build synthetic table from chart recommendations so
                # _generate_chart_from_rec has data to work with.
                tables = _build_synthetic_tables_from_charts(recs, user_prompt)
                if not tables:
                    _finish_error(job, "AI tidak bisa memformat data dari file ini.", result=result)
                    return
                log.info("Text-only mode: synthesised %d table(s) from chart recs", len(tables))
            elif not tables:
                _finish_error(job, "AI tidak bisa memformat data dari file ini.", result=result)
                return
            summary = result.get("summary", {})
            model_used = result.get("_model_used", "")

            _publish(job_id, {
                "status": "running",
                "stage": "ai_formatting",
                "progress": 50,
                "message": f"AI menghasilkan {len(tables)} tabel, {len(recs)} rekomendasi grafik.",
                "tables_count": len(tables),
                "charts_count": len(recs),
            })

            if _is_cancelled(job_id):
                _finish_cancelled(job)
                return

            # ── Step 3: Generate charts ───────────────────────────────────────
            # FIX: Do NOT close session - just commit and let SQLAlchemy manage connections
            try:
                safe_commit()
            except Exception:
                db.session.rollback()
                raise
            job.stage = "generating_charts"
            job.progress = 55
            try:
                safe_commit()
            except Exception:
                db.session.rollback()
                raise

            paper = Paper.query.filter_by(id=paper_id, user_id=user_id).first()
            from tools.editor.utils import safe_paper_image_dir
            paper_dir = safe_paper_image_dir(paper_id)
            if paper_dir:
                paper_dir.mkdir(parents=True, exist_ok=True)

            created_charts = []
            total_recs = len(recs) if recs else 0

            for i, rec in enumerate(recs):
                if _is_cancelled(job_id):
                    _finish_cancelled(job)
                    return

                progress = 55 + int(35 * (i + 1) / max(total_recs, 1))
                _publish(job_id, {
                    "status": "running",
                    "stage": "generating_charts",
                    "progress": progress,
                    "message": f"Membuat grafik {i+1}/{total_recs}: {rec.get('title', 'Grafik')}...",
                    "chart_index": i,
                    "chart_total": total_recs,
                })

                try:
                    chart_info = _generate_chart_from_rec(rec, tables, paper_id, user_id, paper, paper_dir)
                    if chart_info:
                        created_charts.append(chart_info)
                except Exception as e:
                    log.warning("Chart generation failed for '%s': %s", rec.get("title", "?"), e)

            # ── Step 4: Finalize ──────────────────────────────────────────────
            _publish(job_id, {"status": "running", "stage": "finalizing", "progress": 95, "message": "Menyimpan hasil..."})

            # Re-attach job to session (may have been detached by chart gen session ops)
            job = AiJob.query.get(job_id)
            if not job:
                log.error("DataJob %s not found at finalize", job_id)
                return

            table_results = []
            for t in tables:
                table_results.append({
                    "name": t.get("name", "Tabel"),
                    "columns": t.get("columns", []),
                    "rows": t.get("rows", []),
                    "column_types": t.get("column_types", []),
                    "analysis": t.get("analysis", ""),
                    "description": t.get("description", ""),
                })

            job_result = {
                "file_names": file_names,
                "file_count": len(file_paths),
                "tables": table_results,
                "charts": created_charts,
                "summary": summary,
                "model_used": model_used,
                "user_prompt": user_prompt,
            }

            job.status = "done"
            job.stage = "complete"
            job.progress = 100
            job.finished_at = datetime.now(timezone.utc)
            job.result = job_result
            try:
                safe_commit()
            except Exception:
                db.session.rollback()
                raise

            _publish(job_id, {
                "status": "done",
                "stage": "complete",
                "progress": 100,
                "message": f"Selesai! {len(table_results)} tabel, {len(created_charts)} grafik berhasil dibuat.",
                "result": job_result,
            })

            log.info("DataJob %s done: %d tables, %d charts", job_id, len(table_results), len(created_charts))

        except Exception as e:
            log.exception("DataJob %s failed: %s", job_id, e)
            try:
                from database.models import AiJob, db, safe_commit
                j = AiJob.query.get(job_id)
                if j:
                    _finish_error(j, "Internal processing error. Please try again.")
            except Exception:
                pass
        finally:
            # Guard: if the thread was killed (gunicorn restart) or an
            # unexpected exit path left the job in running/queued state,
            # publish error + clean Redis so the job doesn't appear stuck.
            try:
                _cleanup_stale_worker_state(job_id)
            except Exception:
                pass


def _finish_cancelled(job):
    from database.models import db, safe_commit
    job.status = "cancelled"
    job.stage = "cancelled"
    job.finished_at = datetime.now(timezone.utc)
    try:
        safe_commit()
    except Exception:
        db.session.rollback()
        raise
    _publish(job.id, {"status": "cancelled", "progress": job.progress, "stage": "cancelled"})


def _finish_error(job, error_msg, result=None):
    from database.models import db, safe_commit
    job.status = "error"
    job.stage = "error"
    job.error = error_msg
    job.finished_at = datetime.now(timezone.utc)
    if result:
        job.result = result
    try:
        safe_commit()
    except Exception:
        db.session.rollback()
        raise
    _publish(job.id, {"status": "error", "progress": job.progress, "stage": "error", "error": error_msg})


def _extract_file_text(filepath, filename):
    """Extract text from a file (PDF, Excel, CSV, DOCX)."""
    ext = Path(filepath).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(filepath)
    elif ext in (".xlsx", ".xls"):
        return _extract_excel(filepath)
    elif ext in (".csv", ".tsv"):
        return _extract_csv(filepath, ext)
    elif ext in (".docx", ".doc"):
        return _extract_docx(filepath)
    elif ext in (".pptx", ".ppt"):
        return _extract_pptx(filepath)
    else:
        return f"[Format tidak didukung: {ext}]"


def _extract_pdf(filepath):
    try:
        import fitz
        doc = fitz.open(filepath)
        try:
            texts = []
            for page in doc:
                texts.append(page.get_text())
            return "\n".join(texts).strip()
        finally:
            doc.close()
    except ImportError:
        try:
            from pdfminer.high_level import extract_text
            return extract_text(filepath) or ""
        except ImportError:
            return "[PyMuPDF/pdfminer tidak tersedia]"
    except Exception as e:
        return "[Error extracting PDF: file tidak bisa dibaca]" 


def _extract_excel(filepath):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True)
        try:
            parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
                if not rows:
                    continue
                header = rows[0]
                parts.append(f"### Sheet: {sheet_name}\n")
                parts.append("| " + " | ".join(header) + " |")
                parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                for row in rows[1:]:
                    # Copy row before padding to avoid mutating shared list entries
                    row = list(row)
                    row += [""] * (len(header) - len(row))
                    parts.append("| " + " | ".join(row[:len(header)]) + " |")
                parts.append("")
            return "\n".join(parts).strip()
        finally:
            wb.close()
    except ImportError:
        return "[openpyxl tidak tersedia]"
    except Exception as e:
        return "[Error extracting Excel: file tidak bisa dibaca]" 


def _extract_csv(filepath, ext):
    import csv
    delimiter = "\t" if ext == ".tsv" else ","
    try:
        with open(filepath, "r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f, delimiter=delimiter)
            rows = list(reader)
        if not rows:
            return "[File kosong]"
        header = rows[0]
        parts = []
        parts.append("| " + " | ".join(header) + " |")
        parts.append("| " + " | ".join(["---"] * len(header)) + " |")
        for row in rows[1:100]:
            while len(row) < len(header):
                row.append("")
            parts.append("| " + " | ".join(row[:len(header)]) + " |")
        return "\n".join(parts).strip()
    except Exception as e:
        return f"[Error extracting CSV: {e}]"


def _extract_docx(filepath):
    try:
        from docx import Document
        doc = Document(filepath)
        texts = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                texts.append(" | ".join(cells))
        return "\n".join(texts).strip()
    except ImportError:
        return "[python-docx tidak tersedia]"
    except Exception as e:
        return f"[Error extracting DOCX: {e}]"


def _extract_pptx(filepath):
    try:
        from pptx import Presentation
        prs = Presentation(filepath)
        texts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            texts.append(f"--- Slide {slide_num} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        texts.append(" | ".join(cells))
        return "\n".join(texts).strip()
    except ImportError:
        return "[python-pptx tidak tersedia]"
    except Exception as e:
        return f"[Error extracting PPTX: {e}]"


def _generate_chart_from_rec(rec, tables, paper_id, user_id, paper, paper_dir):
    """Generate a chart from an AI recommendation and save to DB."""
    from tools.data.chart_generator import ChartSpec, generate_chart
    from database.models import PaperImage, db, safe_commit

    table_idx = rec.get("table_index", 0)
    if not isinstance(table_idx, int) or table_idx < 0 or table_idx >= len(tables):
        table_idx = 0
    tbl = tables[table_idx]

    x_col_name = rec.get("x_column", "")
    y_col_names = rec.get("y_columns", [])
    columns = tbl.get("columns", [])
    rows = tbl.get("rows", [])

    if not columns or not rows:
        return None

    x_col = columns.index(x_col_name) if x_col_name in columns else -1
    y_cols = [columns.index(n) for n in y_col_names if n in columns]

    if not y_cols:
        return None

    x_data = [rows[r][x_col] if x_col >= 0 and x_col < len(rows[r]) else str(r + 1) for r in range(len(rows))]
    data = []
    for yc in y_cols:
        series = []
        for row in rows:
            val = row[yc] if yc < len(row) else 0
            try:
                series.append(float(str(val).replace(",", "")))
            except (ValueError, TypeError):
                series.append(0)
        data.append(series)

    kind = rec.get("kind", "bar")
    title = rec.get("title", "Grafik")
    settings = rec.get("suggested_settings", {})

    spec = ChartSpec(
        kind=kind,
        title=title,
        xlabel=x_col_name,
        ylabel=", ".join(y_col_names),
        data=data,
        series_labels=y_col_names,
        x_data=x_data,
        color_palette=settings.get("color_palette", "academic"),
        theme=settings.get("theme", "clean"),
        show_data_labels=settings.get("show_data_labels", False),
    )

    out_path = Path(generate_chart(paper_id, spec, user_id=user_id, judul_paper=paper.title if paper else None))

    if paper_dir:
        filename = out_path.name
        dest = paper_dir / filename
        shutil.move(str(out_path), str(dest))
    else:
        filename = out_path.name

    # Cleanup: remove temp chart dir if empty
    try:
        temp_chart_dir = out_path.parent
        if temp_chart_dir.exists() and not any(temp_chart_dir.iterdir()):
            temp_chart_dir.rmdir()
    except Exception:
        pass

    img = PaperImage(
        paper_id=paper_id,
        user_id=user_id,
        filename=filename,
        original_name=f"chart-{kind}-{filename}",
        file_path=f"{paper_id}/{filename}",
    )
    db.session.add(img)
    db.session.flush()

    return {
        "image_id": img.id,
        "filename": filename,
        "url": f"/api/images/{paper_id}/{filename}",
        "kind": kind,
        "title": title,
        "analysis": rec.get("analysis", ""),
        "table_index": table_idx,
        "x_column": x_col_name,
        "y_columns": y_col_names,
    }


# ── Synthetic table builder (text-only fallback) ─────────────────────────────

def _build_synthetic_tables_from_charts(recs: list, user_prompt: str) -> list:
    """Build synthetic tables from chart recommendations + user prompt data.

    When AI returns chart recs without tables (text-only mode), try to
    extract structured data from the user prompt so chart generation works.
    """
    import re

    # Strategy 1: Parse "Name (details): Key1=Val1Unit, Key2=Val2Unit" patterns
    rows = []
    columns = None
    kv_pattern = re.compile(r'(\w[\w\s]*?)\s*=\s*([\d.]+)\s*\w*')
    all_metrics = set()

    # Split prompt at each "Config\d+" to avoid text between configs
    segments = re.split(r'(Config\d+|Konfigurasi\s*\d+)', user_prompt, flags=re.IGNORECASE)
    for i in range(1, len(segments), 2):
        config_name = segments[i].strip()
        rest = segments[i + 1].strip() if i + 1 < len(segments) else ""
        # Extract details from parenthesized block
        det_m = re.match(r'\s*\(([^)]+)\)\s*:\s*(.*)', rest)
        if not det_m:
            continue
        details = det_m.group(1).strip()
        metrics_str = det_m.group(2).strip()
        # Clean trailing punctuation only, don't split on decimal points
        metrics_str = re.sub(r'[,.;]+$', '', metrics_str)

        metrics = {}
        for km in kv_pattern.finditer(metrics_str):
            key = km.group(1).strip()
            val = km.group(2)
            metrics[key] = val
            all_metrics.add(key)

        row = {"Konfigurasi": config_name, "Detail": details, **metrics}
        rows.append(row)

    if rows and all_metrics:
        columns = ["Konfigurasi", "Detail"] + sorted(all_metrics)

    # Strategy 2: If prompt parsing failed, try chart recs analysis field
    if not rows:
        for rec in recs:
            analysis = rec.get("analysis", "")
            x = rec.get("x_column", "")
            ys = rec.get("y_columns", [])
            if analysis and x and ys:
                # Try extracting key=value from analysis text
                for km in kv_pattern.finditer(analysis):
                    key = km.group(1).strip()
                    val = km.group(2)
                    rows.append({"label": key, x: val} if x else {"label": key, **{y: val for y in ys}})
                if rows:
                    columns = list(rows[0].keys()) if rows else None

    if not rows or not columns:
        return []

    # Normalize rows to match columns
    normalized = []
    for row in rows:
        normalized.append([str(row.get(c, "")) for c in columns])

    table = {
        "name": "Data (diextrak dari instruksi)",
        "description": "Data yang diekstrak otomatis dari instruksi teks.",
        "columns": columns,
        "column_types": _detect_column_types(columns, normalized),
        "rows": normalized,
        "analysis": "Data diekstrak dari instruksi user.",
    }
    return [table]


def _detect_column_types(columns: list, rows: list) -> list:
    """Auto-detect column types: numerik vs kategori."""
    types = []
    for ci in range(len(columns)):
        sample = [r[ci] for r in rows[:20] if ci < len(r) and r[ci] not in ("", None)]
        if not sample:
            types.append("kategori")
            continue
        try:
            [float(v.replace(",", "")) for v in sample]
            types.append("numerik")
        except (ValueError, TypeError):
            types.append("kategori")
    return types
