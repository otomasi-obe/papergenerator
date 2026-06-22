"""
ROTASIgen.py - Generate ROTASI DOCX from JSON using the original ROTASI.docx.

The output starts from the original document package so styles, numbering,
header, footer, settings, theme, fonts, and section properties stay intact.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "ROTASI.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

PAGE_W_TWIPS = 11906
MARGIN_LEFT = 1418
MARGIN_RIGHT = 851
BODY_WIDTH_PT = (PAGE_W_TWIPS - MARGIN_LEFT - MARGIN_RIGHT) / 20
MAX_FIGURE_WIDTH_CM = 14.0
ABSTRACT_TABLE_WIDTH_TWIPS = 8504

IND_HEAD_LEFT = 284
IND_HEAD_HANG = 284
IND_BODY_FIRST = 426

ROTASI_STYLE_ID: dict[str, str] = {
    "icsm_title": "icsmtitle",
    "icsm_authors": "icsmauthors",
    "icsm_addresses": "icsmaddresses",
    "icsm_bodytext": "icsmbodytext",
    "icsm_heading1": "icsmheading1",
}

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None

_DYNAMIC_PARTS = {
    "[Content_Types].xml",
    "docProps/core.xml",
    "word/document.xml",
    "word/_rels/document.xml.rels",
}
_DYNAMIC_PREFIXES = ("word/media/",)


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_para_style(paragraph, style_name: str) -> None:
    xml_id = ROTASI_STYLE_ID.get(style_name, style_name)
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), xml_id)


def _ensure_para_rpr(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    rpr = ppr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    return rpr


def _set_para_default_rpr(
    paragraph,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    lang: str | None = None,
    size_pt: float | None = None,
    font_name: str | None = None,
) -> None:
    rpr = _ensure_para_rpr(paragraph)

    def set_on_off(tag: str, value: bool | None) -> None:
        if value is None:
            return
        el = rpr.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            rpr.append(el)
        el.set(qn("w:val"), "1" if value else "0")

    set_on_off("b", bold)
    set_on_off("i", italic)

    if size_pt is not None:
        value = str(int(round(size_pt * 2)))
        for tag in ("sz", "szCs"):
            sz_el = rpr.find(qn(f"w:{tag}"))
            if sz_el is None:
                sz_el = OxmlElement(f"w:{tag}")
                rpr.append(sz_el)
            sz_el.set(qn("w:val"), value)

    if font_name is not None:
        fonts_el = rpr.find(qn("w:rFonts"))
        if fonts_el is None:
            fonts_el = OxmlElement("w:rFonts")
            rpr.append(fonts_el)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts_el.set(qn(f"w:{attr}"), font_name)

    if lang is not None:
        lang_el = rpr.find(qn("w:lang"))
        if lang_el is None:
            lang_el = OxmlElement("w:lang")
            rpr.append(lang_el)
        lang_el.set(qn("w:val"), lang)


def _set_spacing_single(para, after: int = 0) -> None:
    ppr = para._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    sp.set(qn("w:after"), str(after))
    sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto")


def _set_run_lang(run, lang: str) -> None:
    rpr = run._r.get_or_add_rPr()
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def _set_run_format(
    run,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    size_pt: float | None = None,
    font_name: str | None = None,
    lang: str | None = None,
    superscript: bool | None = None,
) -> None:
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size_pt is not None:
        _set_run_size(run, size_pt)
    if font_name is not None:
        run.font.name = font_name
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), font_name)
    if lang is not None:
        _set_run_lang(run, lang)
    if superscript is not None:
        run.font.superscript = superscript


def _sp0(para) -> None:
    ppr = para._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")


def _set_spacing_before_after(para, before_pt: float = 0, after_pt: float = 0) -> None:
    ppr = para._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    sp.set(qn("w:before"), str(int(round(before_pt * 20))))
    sp.set(qn("w:after"), str(int(round(after_pt * 20))))


def _sp_after0(para) -> None:
    ppr = para._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    sp.set(qn("w:after"), "0")


def _jc(para, val: str = "both") -> None:
    ppr = para._p.get_or_add_pPr()
    jc_el = ppr.find(qn("w:jc"))
    if jc_el is None:
        jc_el = OxmlElement("w:jc")
        ppr.append(jc_el)
    jc_el.set(qn("w:val"), val)


def _ind(para, left=None, hanging=None, first_line=None) -> None:
    ppr = para._p.get_or_add_pPr()
    ind_el = ppr.find(qn("w:ind"))
    if ind_el is None:
        ind_el = OxmlElement("w:ind")
        ppr.append(ind_el)
    if left is not None:
        ind_el.set(qn("w:left"), str(left))
    if hanging is not None:
        ind_el.set(qn("w:hanging"), str(hanging))
    if first_line is not None:
        ind_el.set(qn("w:firstLine"), str(first_line))


def _para(doc: Document, style_id: str | None = None):
    paragraph = doc.add_paragraph()
    if style_id:
        _set_para_style(paragraph, style_id)
    return paragraph


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _normalize_text(text: str) -> str:
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    text = _normalize_text(text)
    buf: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush():
        nonlocal buf
        value = "".join(buf)
        buf = []
        if value:
            yield {
                "kind": "text",
                "value": value,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(text):
        char = text[index]
        if char == "\n":
            yield from flush()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(text):
            cmd = text[index + 1]
            if cmd == "\\":
                buf.append("\\")
                index += 2
                continue
            if cmd == "b":
                next_char = text[index + 2] if index + 2 < len(text) else ""
                if next_char.islower():
                    buf.append("\\b")
                    index += 2
                    continue
                yield from flush()
                bold = not bold
                index += 2
                continue
            if cmd == "i":
                next_char = text[index + 2] if index + 2 < len(text) else ""
                if next_char.islower():
                    buf.append("\\i")
                    index += 2
                    continue
                yield from flush()
                italic = not italic
                index += 2
                continue
            if cmd == "u":
                next_char = text[index + 2] if index + 2 < len(text) else ""
                if next_char.islower():
                    buf.append("\\u")
                    index += 2
                    continue
                yield from flush()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = text.find("$", index + 1)
            if closing != -1:
                yield from flush()
                formula = text[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buf.append(char)
        index += 1

    yield from flush()


def _append_rich_text(
    paragraph,
    text: str,
    *,
    default_bold: bool | None = None,
    default_italic: bool | None = None,
    default_size_pt: float | None = None,
    default_font_name: str | None = None,
    default_lang: str | None = None,
) -> None:
    for tok in _iter_rich_tokens(text):
        if tok["kind"] == "linebreak":
            paragraph.add_run().add_break()
        elif tok["kind"] == "math":
            if not _append_inline_math(paragraph, tok["value"]):
                run = paragraph.add_run(tok["value"])
                _set_run_format(
                    run,
                    bold=default_bold,
                    italic=True,
                    size_pt=default_size_pt,
                    font_name=default_font_name,
                    lang=default_lang,
                )
        else:
            run = paragraph.add_run(tok["value"])
            _set_run_format(
                run,
                bold=tok["bold"] if tok.get("bold") else default_bold,
                italic=tok["italic"] if tok.get("italic") else default_italic,
                size_pt=default_size_pt,
                font_name=default_font_name,
                lang=default_lang,
            )
            if tok.get("underline"):
                run.underline = True


def _split_blocks(text: str) -> list[str]:
    normalized = _normalize_text(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [""]


def _normalize_keywords(raw_keywords) -> list[str]:
    if isinstance(raw_keywords, list):
        return [str(k).strip() for k in raw_keywords if str(k).strip()]
    if isinstance(raw_keywords, str):
        return [k.strip() for k in raw_keywords.split(",") if k.strip()]
    return []


def _should_preserve_from_template(name: str) -> bool:
    if name in _DYNAMIC_PARTS:
        return False
    return not any(name.startswith(prefix) for prefix in _DYNAMIC_PREFIXES)


def _restore_template_package_parts(output_path: Path, template_path: Path) -> None:
    temp_path = output_path.with_suffix(output_path.suffix + ".tmp")
    with zipfile.ZipFile(template_path) as template_zip, zipfile.ZipFile(output_path) as output_zip:
        template_names = set(template_zip.namelist())
        output_names = set(output_zip.namelist())
        ordered_names = []
        seen = set()
        for name in list(output_zip.namelist()) + list(template_zip.namelist()):
            if name not in seen:
                seen.add(name)
                ordered_names.append(name)

        with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as rebuilt_zip:
            for name in ordered_names:
                if name in template_names and _should_preserve_from_template(name):
                    rebuilt_zip.writestr(name, template_zip.read(name))
                elif name in output_names:
                    rebuilt_zip.writestr(name, output_zip.read(name))
                elif name in template_names:
                    rebuilt_zip.writestr(name, template_zip.read(name))

    temp_path.replace(output_path)


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    candidate = json_path.parent / path
    if candidate.exists():
        return candidate
    return BASE_DIR / path


def _set_cell_shading(cell, fill: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_single_cell_table_width(table, width_twips: int) -> None:
    tbl = table._tbl
    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    grid_col = OxmlElement("w:gridCol")
    grid_col.set(qn("w:w"), str(width_twips))
    tbl_grid.append(grid_col)

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_twips))
            tc_w.set(qn("w:type"), "dxa")


def _set_horizontal_cell_borders(cell, top: bool = False, bottom: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    spec = {
        "top": (
            {"val": "single", "sz": "4", "space": "0", "color": "auto"} if top else {"val": "none"}
        ),
        "bottom": (
            {"val": "single", "sz": "4", "space": "0", "color": "auto"}
            if bottom
            else {"val": "none"}
        ),
        "left": {"val": "none"},
        "right": {"val": "none"},
    }
    for edge, attrs in spec.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        for key, value in attrs.items():
            el.set(qn(f"w:{key}"), value)


def _set_run_size(run, size_pt: float) -> None:
    run.font.size = Pt(size_pt)


def _add_title(doc: Document, config: dict) -> None:
    title = str(config.get("title", "Untitled")).strip() or "Untitled"

    p = _para(doc, "icsm_title")
    _sp0(p)
    _set_para_default_rpr(p, bold=True, lang="en-US", size_pt=12)
    run = p.add_run(title)
    _set_run_format(run, bold=True, size_pt=12, font_name="Times New Roman", lang="en-US")

    spacer = _para(doc, "icsm_title")
    _sp0(spacer)
    _set_para_default_rpr(spacer, bold=True, lang="en-US", size_pt=12)


def _add_authors(doc: Document, config: dict) -> None:
    authors = config.get("authors", [])
    if not authors:
        return

    affiliation_map: dict[str, tuple[str, str]] = {}
    author_entries = []

    for index, author in enumerate(authors):
        name = str(author.get("name", "")).strip()
        affiliation = str(author.get("affiliation", "")).strip()
        location = str(author.get("location", "")).strip()
        email = str(author.get("email", "")).strip()
        aff_text = ", ".join(part for part in (affiliation, location) if part)
        aff_key = aff_text or affiliation or location
        letter = ""
        if aff_key:
            if aff_key not in affiliation_map:
                letter = chr(ord("a") + len(affiliation_map))
                affiliation_map[aff_key] = (letter, aff_text or aff_key)
            else:
                letter = affiliation_map[aff_key][0]
        author_entries.append(
            {
                "name": name,
                "letter": letter,
                "email": email,
                "corresponding": index == 0,
            }
        )

    names_para = _para(doc, "icsm_authors")
    _sp_after0(names_para)
    _set_para_default_rpr(names_para, bold=True, lang="en-US", size_pt=10)
    for index, author in enumerate(author_entries):
        if index:
            sep = names_para.add_run(", ")
            _set_run_format(sep, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
        name_run = names_para.add_run(author["name"])
        _set_run_format(name_run, bold=True, size_pt=10, font_name="Times New Roman", lang="id-ID")
        if author["letter"]:
            sup = names_para.add_run(author["letter"])
            _set_run_format(
                sup,
                bold=True,
                size_pt=10,
                font_name="Times New Roman",
                lang="en-US",
                superscript=True,
            )
        if author["corresponding"]:
            star = names_para.add_run("*")
            _set_run_format(star, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")

    for letter, text in affiliation_map.values():
        p = _para(doc, "icsm_addresses")
        _sp_after0(p)
        _set_para_default_rpr(p, italic=False, lang="id-ID", size_pt=10)
        letter_run = p.add_run(letter)
        _set_run_format(
            letter_run,
            bold=False,
            italic=False,
            size_pt=10,
            font_name="Times New Roman",
            lang="en-US",
            superscript=True,
        )
        body_run = p.add_run(text)
        _set_run_format(
            body_run,
            bold=False,
            italic=False,
            size_pt=10,
            font_name="Times New Roman",
            lang="id-ID",
        )

    emails = list(dict.fromkeys(a["email"] for a in author_entries if a["email"]))
    if emails:
        p = _para(doc, "icsm_addresses")
        _sp_after0(p)
        _set_para_default_rpr(p, italic=False, lang="en-US", size_pt=10)
        run = p.add_run("*E-mail: " + "; ".join(emails))
        _set_run_format(
            run, bold=False, italic=False, size_pt=10, font_name="Times New Roman", lang="en-US"
        )

    spacer = _para(doc, "icsm_title")
    _sp0(spacer)
    _jc(spacer, "left")
    _set_para_default_rpr(spacer, lang="en-US", size_pt=10)


def _add_abstract_table(
    doc: Document, header_label: str, abstract_text: str, keywords_label: str, keywords: list[str]
) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders_match_template(table)
    _set_single_cell_table_width(table, ABSTRACT_TABLE_WIDTH_TWIPS)

    tbl = table._tbl
    tbl_pr = tbl.find(qn("w:tblPr"))
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "0")
    tbl_w.set(qn("w:type"), "auto")

    shd = tbl_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tbl_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "D9D9D9")

    cell = table.cell(0, 0)
    _set_cell_shading(cell, "D9D9D9")

    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    jc_el = tr_pr.find(qn("w:jc"))
    if jc_el is None:
        jc_el = OxmlElement("w:jc")
        tr_pr.append(jc_el)
    jc_el.set(qn("w:val"), "center")

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing_single(p1)
    _set_para_default_rpr(p1, bold=True, lang="en-US", size_pt=10, font_name="Times New Roman")
    r1 = p1.add_run(header_label)
    _set_run_format(r1, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing_single(p2)
    _set_para_default_rpr(p2, bold=True, lang="en-US", size_pt=10, font_name="Times New Roman")

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_spacing_single(p3)
    _set_para_default_rpr(p3, lang="en-US", size_pt=10, font_name="Times New Roman")
    if abstract_text:
        _append_rich_text(
            p3,
            abstract_text,
            default_bold=False,
            default_italic=False,
            default_size_pt=10,
            default_font_name="Times New Roman",
            default_lang="en-US",
        )

    p4 = cell.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_spacing_single(p4)
    _set_para_default_rpr(p4, bold=True, lang="en-US", size_pt=10, font_name="Times New Roman")
    r4 = p4.add_run(" ")
    _set_run_format(r4, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")

    p5 = cell.add_paragraph()
    _set_spacing_single(p5)
    _set_para_default_rpr(p5, lang="en-US", size_pt=10, font_name="Times New Roman")
    r5 = p5.add_run(f"{keywords_label}: ")
    _set_run_format(r5, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
    if keywords:
        r6 = p5.add_run(", ".join(keywords))
        _set_run_format(r6, bold=False, size_pt=10, font_name="Times New Roman", lang="en-US")


def _add_abstracts(doc: Document, config: dict) -> None:
    abstract_generic = str(config.get("abstract", "")).strip()
    abstract_en = str(config.get("abstractEN", "")).strip() or abstract_generic
    abstract_id = str(config.get("abstractID", config.get("abstrak", ""))).strip()
    keywords = _normalize_keywords(config.get("keywords", []))

    if abstract_en:
        _add_abstract_table(doc, "Abstract", abstract_en, "Keywords", keywords)

        sep = _para(doc, "icsm_heading1")
        _sp0(sep)
        _jc(sep)
        _set_para_default_rpr(sep, bold=False, lang="id-ID")

    if abstract_id:
        _add_abstract_table(doc, "Abstrak", abstract_id, "Kata kunci", keywords)

        sep2 = _para(doc, "icsm_heading1")
        _sp0(sep2)
        _jc(sep2)
        _set_para_default_rpr(sep2, bold=False, lang="id-ID")


def _add_prompt_box(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Normal Table"
    _set_table_borders_match_template(table)
    cell = table.cell(0, 0)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "bottom", "left", "right"):
        el = tc_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
    p = cell.paragraphs[0]
    _set_para_style(p, "icsm_heading1")
    _sp0(p)
    _jc(p)
    _set_para_default_rpr(p, bold=False, lang="en-US")
    _append_rich_text(
        p,
        text,
        default_bold=False,
        default_italic=False,
        default_size_pt=10,
        default_font_name="Times New Roman",
        default_lang="en-US",
    )


def _add_figure(doc: Document, item: dict, json_path: Path) -> None:

    image_number = str(item.get("ImageNumber", "")).strip()
    path_text = str(item.get("Path", "")).strip()
    prompt = str(item.get("Prompt", "")).strip()
    title = str(item.get("Title", "")).strip()

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except Exception:
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    image_path = _resolve_path(path_text, json_path) if path_text else None

    if image_path and image_path.is_file():
        p_img = _para(doc, "icsm_heading1")
        _sp0(p_img)
        _jc(p_img, "center")
        p_img.add_run().add_picture(str(image_path), width=Cm(width_cm))
    else:
        fallback = prompt if prompt else f"[Gambar {image_number}: {title}]"
        _add_prompt_box(doc, fallback)

    if image_number and title:
        cap = _para(doc, "icsm_heading1")
        _set_spacing_before_after(cap, before_pt=3, after_pt=6)
        _jc(cap, "center")
        _set_para_default_rpr(cap, lang="id-ID")
        label_run = cap.add_run("Gambar")
        _set_run_format(label_run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
        num_run = cap.add_run(f" {image_number}. ")
        _set_run_format(num_run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
        title_run = cap.add_run(title)
        _set_run_format(
            title_run, bold=False, size_pt=10, font_name="Times New Roman", lang="id-ID"
        )


def _add_equation_line(doc: Document, formula: str, number: str | None = None) -> None:
    p = _para(doc, "icsm_heading1")
    _set_spacing_before_after(p, before_pt=6, after_pt=6)
    _jc(p, "both")
    _set_para_default_rpr(p, bold=False, lang="en-US")
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Pt(BODY_WIDTH_PT / 2), WD_TAB_ALIGNMENT.CENTER)
    tab_stops.add_tab_stop(Pt(BODY_WIDTH_PT), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("\t")
    if not _append_inline_math(p, formula):
        run = p.add_run(formula)
        _set_run_format(
            run, bold=False, italic=True, size_pt=10, font_name="Times New Roman", lang="en-US"
        )
    if number is not None:
        run = p.add_run(f"\t({number})")
        _set_run_format(
            run, bold=False, italic=False, size_pt=10, font_name="Times New Roman", lang="en-US"
        )


def _add_equation(doc: Document, item: dict) -> None:
    number = str(item.get("FormulaNumber", "")).strip() or None
    formula = str(item.get("text", "") or item.get("latex", "")).strip()
    if formula:
        _add_equation_line(doc, formula, number)


def _add_table(doc: Document, item: dict) -> None:
    table_number = str(item.get("TableNumber", "")).strip()
    title = str(item.get("Title", "")).strip()
    headers = list(item.get("Headers", []))
    rows = list(item.get("Rows", []))
    if not headers:
        return

    cap = _para(doc, "icsm_heading1")
    _set_spacing_before_after(cap, before_pt=6, after_pt=3)
    _jc(cap)
    _set_para_default_rpr(cap, bold=False, lang="id-ID")
    label_run = cap.add_run("Tabel")
    _set_run_format(label_run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
    if table_number:
        num_run = cap.add_run(f" {table_number}.")
        _set_run_format(num_run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")
    if title:
        title_run = cap.add_run(f" {title}")
        _set_run_format(
            title_run, bold=False, size_pt=10, font_name="Times New Roman", lang="id-ID"
        )

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders_match_template(table)

    for ci, val in enumerate(headers):
        cell = table.rows[0].cells[ci]
        _set_horizontal_cell_borders(cell, top=True, bottom=True)
        cell.text = ""
        p = cell.paragraphs[0]
        _set_para_style(p, "icsm_heading1")
        _sp0(p)
        _jc(p, "center")
        _set_para_default_rpr(p, lang="en-US")
        run = p.add_run(str(val))
        _set_run_format(run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")

    for ri, row_data in enumerate(rows, start=1):
        is_last = ri == len(rows)
        for ci, val in enumerate(row_data):
            if ci >= len(headers):
                break
            cell = table.rows[ri].cells[ci]
            _set_horizontal_cell_borders(cell, top=False, bottom=is_last)
            cell.text = ""
            p = cell.paragraphs[0]
            _set_para_style(p, "icsm_heading1")
            _sp0(p)
            _jc(p, "center")
            _set_para_default_rpr(p, bold=False, lang="en-US")
            _append_rich_text(
                p,
                str(val),
                default_bold=False,
                default_italic=False,
                default_size_pt=10,
                default_font_name="Times New Roman",
                default_lang="en-US",
            )

    spacer = _para(doc, "icsm_heading1")
    _sp0(spacer)
    _jc(spacer)
    _set_para_default_rpr(spacer, bold=False, lang="en-US")


def _body_paragraphs(doc: Document, text: str) -> None:
    for block in _split_blocks(text):
        p = _para(doc, "icsm_heading1")
        _sp0(p)
        _jc(p)
        _ind(p, first_line=IND_BODY_FIRST)
        _set_para_default_rpr(p, bold=False, lang="en-US")
        _append_rich_text(
            p,
            block,
            default_bold=False,
            default_italic=False,
            default_size_pt=10,
            default_font_name="Times New Roman",
            default_lang="en-US",
        )


def _add_section_heading(doc: Document, section_number: int, title: str) -> None:
    p = _para(doc, "icsm_heading1")
    _sp0(p)
    _jc(p)
    _ind(p, left=IND_HEAD_LEFT, hanging=IND_HEAD_HANG)
    _set_para_default_rpr(p, lang="en-US")
    _append_rich_text(
        p,
        f"{section_number}. {title}",
        default_bold=True,
        default_italic=False,
        default_size_pt=10,
        default_font_name="Times New Roman",
        default_lang="en-US",
    )


def _add_subsection_heading(doc: Document, prefix: str, title: str) -> None:
    p = _para(doc, "icsm_heading1")
    _sp0(p)
    _jc(p)
    _ind(p, left=IND_HEAD_LEFT, hanging=IND_HEAD_HANG)
    _set_para_default_rpr(p, bold=False, lang="en-US")
    _append_rich_text(
        p,
        f"{prefix} {title}",
        default_bold=False,
        default_italic=False,
        default_size_pt=10,
        default_font_name="Times New Roman",
        default_lang="en-US",
    )


def _render_content_item(doc: Document, item: dict, json_path: Path) -> None:
    item_id = str(item.get("id", "")).lower()
    if item_id == "text":
        text = str(item.get("text", ""))
        if text:
            _body_paragraphs(doc, text)
    elif item_id in ("gambar", "image"):
        _add_figure(doc, item, json_path)
    elif item_id in ("rumus", "formula"):
        _add_equation(doc, item)
    elif item_id in ("tabel", "table"):
        _add_table(doc, item)


def _iter_subsection_keys(section: dict, section_num: int) -> list[str]:
    keys: list[tuple[str, str]] = []
    for key, value in section.items():
        if not isinstance(value, dict):
            continue
        match = re.match(rf"^section{section_num}([a-z]+)$", key)
        if match:
            keys.append((match.group(1), key))
        elif key.startswith("sub") and len(key) > 3 and key[3:].isalnum():
            keys.append((key[3:], key))
    return [key for _, key in sorted(keys)]


def _render_subsection(doc: Document, subsection: dict, json_path: Path, prefix: str) -> None:
    title = str(subsection.get("title", "")).strip()
    if title:
        _add_subsection_heading(doc, prefix, title)

    for item in subsection.get("content", []):
        if isinstance(item, dict):
            _render_content_item(doc, item, json_path)
        elif isinstance(item, str) and item.strip():
            _body_paragraphs(doc, item.strip())


def _render_sections(doc: Document, config: dict, json_path: Path) -> None:
    section_keys = sorted(
        [k for k in config.keys() if re.match(r"^section\d+$", k)],
        key=lambda key: int(key.replace("section", "")),
    )
    for section_index, section_key in enumerate(section_keys, start=1):
        section = config[section_key]
        section_title = str(section.get("title", "")).strip()
        if section_title:
            _add_section_heading(doc, section_index, section_title)

        content = section.get("content", [])
        if isinstance(content, str) and content.strip():
            _body_paragraphs(doc, content.strip())
        elif isinstance(content, list):
            for item in content:
                if isinstance(item, str) and item.strip():
                    _body_paragraphs(doc, item.strip())
                elif isinstance(item, dict):
                    _render_content_item(doc, item, json_path)

        subsection_keys = _iter_subsection_keys(section, section_index)
        for subsection_index, subsection_key in enumerate(subsection_keys, start=1):
            subsection = section[subsection_key]
            if isinstance(subsection, dict):
                _render_subsection(
                    doc, subsection, json_path, prefix=f"{section_index}.{subsection_index}"
                )


def _add_references(doc: Document, config: dict) -> None:
    ref_block = config.get("references") or config.get("section_references")
    references: list = []
    if ref_block:
        references = ref_block.get("content", [])
    if not references:
        return

    heading = _para(doc, "icsm_heading1")
    _sp0(heading)
    _jc(heading)
    _set_para_default_rpr(heading, bold=False, lang="en-US")
    run = heading.add_run("Daftar Pustaka")
    _set_run_format(run, bold=True, size_pt=10, font_name="Times New Roman", lang="en-US")

    for index, ref in enumerate(references, start=1):
        if isinstance(ref, dict):
            ref_id = str(ref.get("id", index)).strip() or str(index)
            ref_text = str(ref.get("text", "")).strip()
        else:
            ref_text = str(ref).strip()
            match = re.match(r"^\s*\[(.+?)\]\s*(.*)", ref_text, re.DOTALL)
            if match:
                ref_id = match.group(1).strip()
                ref_text = match.group(2).strip()
            else:
                ref_id = str(index)

        p = _para(doc, "icsm_heading1")
        _sp0(p)
        _jc(p)
        _ind(p, left=IND_HEAD_LEFT, hanging=IND_HEAD_HANG)
        _set_para_default_rpr(p, bold=False, lang="id-ID")
        _append_rich_text(
            p,
            f"[{ref_id}] {ref_text}",
            default_bold=False,
            default_italic=False,
            default_size_pt=10,
            default_font_name="Times New Roman",
            default_lang="id-ID",
        )


def _set_core_properties(doc: Document, config: dict) -> None:
    try:
        props = doc.core_properties
        props.title = str(config.get("title", "")).strip()
        props.author = ", ".join(
            str(author.get("name", "")).strip()
            for author in config.get("authors", [])
            if str(author.get("name", "")).strip()
        )
        props.keywords = ", ".join(_normalize_keywords(config.get("keywords", [])))
    except Exception:
        pass


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _clear_document_body(doc)

    _set_core_properties(doc, config)
    _add_title(doc, config)
    _add_authors(doc, config)
    _add_abstracts(doc, config)
    _render_sections(doc, config, Path(json_path))
    _add_references(doc, config)

    doc.save(str(final_output))
    _restore_template_package_parts(final_output, Path(template_path))
    return final_output


def main() -> None:
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        if not json_arg.exists():
            print(f"ERROR: File not found: {json_arg}")
            return
        result = build_document(json_arg, output_arg, template_arg)
        print(f"Generated: {result.name}")
        return

    skip_names = {"package.json", "tsconfig.json", "settings.json"}
    json_files = [f for f in sorted(BASE_DIR.glob("*.json")) if f.name.lower() not in skip_names]
    if not json_files:
        print("No JSON files found")
        return
    for json_file in json_files:
        try:
            data = json.loads(json_file.read_text(encoding="utf-8"))
            if not isinstance(data, dict):
                print(f"Skip {json_file.name}: not a JSON object")
                continue
            result = build_document(json_file)
            print(f"Generated: {result.name}")
        except json.JSONDecodeError as exc:
            print(f"Skip {json_file.name}: invalid JSON - {exc}")
        except Exception as exc:
            print(f"Error {json_file.name}: {exc}")
    print("\nDone!")


def _set_table_borders_match_template(table) -> None:
    """Set border tabel sesuai pattern template original: HORIZONTAL_ONLY (academic).

    Pattern booktabs / 3-line table: top + bottom visible, left/right/insideV nil,
    insideH visible (untuk garis di bawah header).
    Auto-injected oleh _fix_table_borders_v2.py.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    visible_sides = {"top", "bottom", "insideH"}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        if edge in visible_sides:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
            el.set(qn("w:sz"), "0")


if __name__ == "__main__":
    main()
