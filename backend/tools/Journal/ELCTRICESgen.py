"""
ELCTRICESgen.py — Generator dokumen ELCTRICES dari JSON.

Menggunakan ELCTRICES.docx sebagai template asli agar header, footer,
styles, numbering, dan section properties tetap mengikuti dokumen sumber.
Semua konten diambil dari _PLC-MediapipeID.json.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "ELCTRICES.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

MAX_FIGURE_WIDTH_CM = 13.5

NS_MAP_STRICT = {
    b"http://purl.oclc.org/ooxml/wordprocessingml/main":
        b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships":
        b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main":
        b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing":
        b"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    b"http://purl.oclc.org/ooxml/officeDocument/math":
        b"http://schemas.openxmlformats.org/officeDocument/2006/math",
}

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None


@dataclass
class RenderState:
    figure_count: int = 0
    table_count: int = 0


def _set_ai_prompt_color_red(doc):
    """Scan output DOCX, set warna text MERAH untuk semua run di paragraf
    yang berisi pola '[PROMPT UNTUK AI GAMBAR'.
    Idempotent dan aman dipanggil setelah doc.save() / sebelum save."""
    RED = RGBColor(0xFF, 0x00, 0x00)
    for p in doc.paragraphs:
        text = p.text or ""
        if "[PROMPT UNTUK AI GAMBAR" in text or "[PROMPT AI GAMBAR" in text:
            for r in p.runs:
                try:
                    r.font.color.rgb = RED
                except Exception:
                    pass
    # Juga scan paragraf di dalam tabel (kalau ada)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text or ""
                    if "[PROMPT UNTUK AI GAMBAR" in text or "[PROMPT AI GAMBAR" in text:
                        for r in p.runs:
                            try:
                                r.font.color.rgb = RED
                            except Exception:
                                pass


def _wq(tag: str) -> str:
    return f"{{{NS_W}}}{tag}"


def _strict_to_trans(data: bytes) -> bytes:
    for old, new in NS_MAP_STRICT.items():
        data = data.replace(old, new)
    return data


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_document_final_sectpr(doc: Document, sectpr: etree._Element) -> None:
    body = doc._element.body
    current = body.find(qn("w:sectPr"))
    if current is not None:
        body.remove(current)
    body.append(deepcopy(sectpr))


def _apply_sample_ppr(paragraph, sample_ppr: etree._Element | None) -> None:
    current = paragraph._p.find(qn("w:pPr"))
    if current is not None:
        paragraph._p.remove(current)
    if sample_ppr is not None:
        paragraph._p.insert(0, deepcopy(sample_ppr))


def _apply_sample_rpr(run, sample_rpr: etree._Element | None) -> None:
    current = run._r.find(qn("w:rPr"))
    if current is not None:
        run._r.remove(current)
    if sample_rpr is not None:
        run._r.insert(0, deepcopy(sample_rpr))


def _new_paragraph(doc: Document, sample_ppr: etree._Element | None = None):
    paragraph = doc.add_paragraph()
    if sample_ppr is not None:
        _apply_sample_ppr(paragraph, sample_ppr)
    return paragraph


def _add_sample_run(
    paragraph,
    text: str,
    sample_rpr: etree._Element | None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
):
    run = paragraph.add_run(text)
    _apply_sample_rpr(run, sample_rpr)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    return run


def _normalize_omml_math(omml: etree._Element, half_points: int = 20) -> etree._Element:
    for math_run in omml.findall(f".//{{{MATH_NS}}}r"):
        rpr = math_run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            math_run.insert(0, rpr)
        _set_math_run_defaults(rpr, half_points=half_points)

    for ctrl_pr in omml.findall(f".//{{{MATH_NS}}}ctrlPr"):
        rpr = ctrl_pr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ctrl_pr.insert(0, rpr)
        _set_math_run_defaults(rpr, half_points=half_points)

    return omml


def _set_math_run_defaults(rpr: etree._Element, half_points: int = 20) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Cambria Math")

    for tag in ("sz", "szCs"):
        size = rpr.find(qn(f"w:{tag}"))
        if size is None:
            size = OxmlElement(f"w:{tag}")
            rpr.append(size)
        size.set(qn("w:val"), str(half_points))

    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    if lang.get(qn("w:val")) is None:
        lang.set(qn("w:val"), "en-US")


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    omml = _normalize_omml_math(omml, half_points=20)
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _normalize_text_commands(text: str) -> str:
    # Replace literal \n → newline, but ONLY when not followed by a-z
    # (LaTeX commands like \nu, \nabla, \neg, \notin must be preserved).
    text = re.sub(r'\\n(?![a-z])', '\n', text)
    # Same for \t — preserve \tau, \theta, \times, \tan, \text etc.
    text = re.sub(r'\\t(?![a-z])', '\t', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r'\*\*(.+?)\*\*', r'\\b\1\\b', text, flags=re.DOTALL)
    text = re.sub(r'\*([^*\n]+?)\*', r'\\i\1\\i', text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1:closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1

    yield from flush_buffer()


def _append_rich_text(paragraph, text: str, sample_rpr: etree._Element | None) -> None:
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            paragraph.add_run().add_break()
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = _add_sample_run(paragraph, token["value"], sample_rpr)
                run.italic = True
            continue
        _add_sample_run(
            paragraph,
            token["value"],
            sample_rpr,
            bold=token["bold"],
            italic=token["italic"],
            underline=token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    blocks = [part for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return blocks or [normalized.strip()]


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _text_of_run(run_el: etree._Element) -> str:
    return "".join(t.text or "" for t in run_el.findall(_wq("t")))


def _load_template_samples(template_path: Path) -> dict[str, etree._Element | None]:
    with zipfile.ZipFile(template_path) as archive:
        raw = _strict_to_trans(archive.read("word/document.xml"))
    doc_root = etree.fromstring(raw)
    body = doc_root.find(_wq("body"))
    if body is None:
        raise RuntimeError("Template ELCTRICES tidak memiliki body XML yang valid.")

    children = list(body)

    def body_paragraph(index: int) -> etree._Element:
        para = children[index]
        if para.tag != _wq("p"):
            raise RuntimeError(f"Elemen body[{index}] bukan paragraf.")
        return para

    def clone_ppr(index: int):
        ppr = body_paragraph(index).find(_wq("pPr"))
        return deepcopy(ppr) if ppr is not None else None

    def clone_rpr(index: int, *, contains: str | None = None, fallback: int = 0):
        runs = body_paragraph(index).findall(_wq("r"))
        if contains is not None:
            for run in runs:
                if contains in _text_of_run(run):
                    rpr = run.find(_wq("rPr"))
                    return deepcopy(rpr) if rpr is not None else None
        if not runs:
            return None
        fallback = max(0, min(fallback, len(runs) - 1))
        rpr = runs[fallback].find(_wq("rPr"))
        return deepcopy(rpr) if rpr is not None else None

    def clone_sectpr(index: int):
        ppr = body_paragraph(index).find(_wq("pPr"))
        if ppr is None:
            return None
        sectpr = ppr.find(_wq("sectPr"))
        return deepcopy(sectpr) if sectpr is not None else None

    title_break_ppr = clone_ppr(21)
    if title_break_ppr is None:
        raise RuntimeError("Template ELCTRICES tidak memiliki section break paragraf judul.")
    sectpr = title_break_ppr.find(_wq("sectPr"))
    if sectpr is not None:
        pgnum = sectpr.find(_wq("pgNumType"))
        if pgnum is not None:
            pgnum.set(_wq("start"), "1")

    body_sectpr = body.find(_wq("sectPr"))
    body_sectpr = deepcopy(body_sectpr) if body_sectpr is not None else clone_sectpr(103)
    if body_sectpr is None:
        raise RuntimeError("Template ELCTRICES tidak memiliki section break body penutup.")
    # cols=1 continuous break that closes the body section (matches original idx 103)
    body_close_ppr = clone_ppr(103)

    return {
        "title_id_ppr": clone_ppr(1),
        "title_id_rpr": clone_rpr(1, contains="Sistem"),
        "blank_center_ppr": clone_ppr(4),
        "title_en_ppr": clone_ppr(3),
        "title_en_rpr": clone_rpr(3, contains="Baby"),
        "author_ppr": clone_ppr(5),
        "author_name_rpr": clone_rpr(5, contains="Penulis A"),
        "author_sup_rpr": clone_rpr(5, contains="1"),
        "affiliation_primary_ppr": clone_ppr(6),
        "affiliation_primary_sup_rpr": clone_rpr(6, contains="1,2"),
        "affiliation_primary_text_rpr": clone_rpr(6, contains="Program Studi"),
        "affiliation_secondary_ppr": clone_ppr(7),
        "affiliation_secondary_sup_rpr": clone_rpr(7, contains="3"),
        "affiliation_secondary_text_rpr": clone_rpr(7, contains="Program Studi"),
        "email_ppr": clone_ppr(8),
        "email_rpr": clone_rpr(6, contains="Program Studi"),
        "abstrak_heading_ppr": clone_ppr(11),
        "abstrak_heading_rpr": clone_rpr(11, contains="ABSTRAK"),
        "abstract_id_ppr": clone_ppr(13),
        "abstract_id_open_rpr": clone_rpr(13, contains="("),
        "abstract_id_word_rpr": clone_rpr(13, contains="Abstrak"),
        "abstract_id_close_rpr": clone_rpr(13, contains=")."),
        "abstract_id_body_rpr": clone_rpr(13, contains=" Tinggi"),
        "keywords_id_ppr": clone_ppr(14),
        "keywords_id_label_rpr": clone_rpr(14, contains="Kata kunci"),
        "keywords_id_body_rpr": clone_rpr(14, contains=":"),
        "abstract_en_heading_ppr": clone_ppr(16),
        "abstract_en_heading_rpr": clone_rpr(16, contains="ABSTRACT"),
        "abstract_en_ppr": clone_ppr(18),
        "abstract_en_open_rpr": clone_rpr(18, contains="("),
        "abstract_en_word_rpr": clone_rpr(18, contains="Abstract"),
        "abstract_en_close_rpr": clone_rpr(18, contains=")."),
        "abstract_en_body_rpr": clone_rpr(18, contains=" Baby"),
        "keywords_en_ppr": clone_ppr(19),
        "keywords_en_label_rpr": clone_rpr(19, contains="Keywords"),
        "keywords_en_body_rpr": clone_rpr(19, contains=":"),
        "section_break_ppr": title_break_ppr,
        "section_heading_ppr": clone_ppr(22),
        "section_heading_rpr": clone_rpr(22, contains="PENDAHULUAN"),
        "body_first_ppr": clone_ppr(23),
        "body_first_rpr": clone_rpr(23, contains="Pendahuluan"),
        "body_ppr": clone_ppr(24),
        "body_rpr": clone_rpr(24, contains="Untuk menunjukkan"),
        "subsection_heading_ppr": clone_ppr(51),
        "subsection_label_rpr": clone_rpr(51, contains="A."),
        "subsection_title_rpr": clone_rpr(51, contains="Sub Bagian 1"),
        "subsection_body_ppr": clone_ppr(52),
        "subsection_body_rpr": clone_rpr(52, contains="Xxxx"),
        "equation_ppr": clone_ppr(34),
        "equation_spacer_rpr": clone_rpr(34, fallback=1),
        "equation_number_rpr": clone_rpr(34, contains="(1)"),
        "figure_caption_ppr": clone_ppr(41),
        "figure_caption_label_rpr": clone_rpr(41, contains="Gambar 1."),
        "figure_caption_text_rpr": clone_rpr(41, contains="Judul"),
        "table_caption_ppr": clone_ppr(45),
        "table_caption_rpr": clone_rpr(45, contains="Tabel 1."),
        "reference_heading_ppr": clone_ppr(59),
        "reference_heading_rpr": clone_rpr(59, contains="DAFTAR PUSTAKA"),
        "reference_item_ppr": clone_ppr(63),
        "reference_item_rpr": clone_rpr(63, contains="Penulis1 A"),
        "body_sectpr": body_sectpr,
        "body_close_ppr": body_close_ppr,
    }


def _pick_first(config: dict, keys: tuple[str, ...], default: str = "") -> str:
    for key in keys:
        value = config.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if text:
            return text
    return default


def _title_texts(config: dict) -> tuple[str, str]:
    base = _pick_first(config, ("title",), "Untitled Paper")
    title_id = _pick_first(config, ("title_id", "title_indonesian", "title_ina", "judul"), base)
    title_en = _pick_first(config, ("title_en", "title_english"), base)
    return title_id, title_en


def _abstract_texts(config: dict) -> tuple[str, str]:
    block = config.get("Abstract", {}) if isinstance(config.get("Abstract"), dict) else {}
    base = _pick_first(config, ("abstract",), str(block.get("Indonesian") or block.get("English") or "").strip())
    abstract_id = _pick_first(
        config,
        ("abstract_id", "abstract_indonesian", "abstrak"),
        str(block.get("Indonesian") or base).strip(),
    )
    abstract_en = _pick_first(
        config,
        ("abstract_en", "abstract_english"),
        str(block.get("English") or base).strip(),
    )
    return abstract_id, abstract_en


def _keyword_lists(config: dict) -> tuple[list[str], list[str]]:
    block = config.get("Abstract", {}) if isinstance(config.get("Abstract"), dict) else {}

    def normalize(value) -> list[str]:
        if isinstance(value, list):
            return [str(item).strip() for item in value if str(item).strip()]
        text = str(value or "").strip()
        if not text:
            return []
        return [item.strip() for item in text.split(",") if item.strip()]

    base = normalize(config.get("keywords") or block.get("KeywordsIndonesian") or block.get("KeywordsEnglish"))
    keywords_id = normalize(config.get("keywords_id") or config.get("keywords_indonesian") or block.get("KeywordsIndonesian")) or base
    keywords_en = normalize(config.get("keywords_en") or config.get("keywords_english") or block.get("KeywordsEnglish")) or base
    return keywords_id, keywords_en


def _author_entries(config: dict) -> list[dict[str, str]]:
    authors = config.get("authors", [])
    if not isinstance(authors, list):
        return []
    entries = []
    for author in authors:
        if not isinstance(author, dict):
            continue
        name = str(author.get("name", "")).strip()
        if not name:
            continue
        entries.append(
            {
                "name": name,
                "affiliation": str(author.get("affiliation", "")).strip(),
                "location": str(author.get("location", "")).strip(),
                "email": str(author.get("email", "")).strip(),
            }
        )
    return entries


def _shorten_text(text: str, max_len: int) -> str:
    compact = " ".join(text.split())
    if len(compact) <= max_len:
        return compact
    return compact[: max_len - 3].rstrip() + "..."


def _running_title(config: dict) -> str:
    _, title_en = _title_texts(config)
    return _shorten_text(title_en, 60)


def _running_authors(config: dict) -> str:
    authors = _author_entries(config)
    if not authors:
        return ""
    names = ", ".join(entry["name"] for entry in authors)
    return _shorten_text(names, 80)


def _render_title_block(doc: Document, config: dict, samples: dict[str, etree._Element | None]) -> None:
    title_id, title_en = _title_texts(config)
    abstract_id, abstract_en = _abstract_texts(config)
    keywords_id, keywords_en = _keyword_lists(config)
    authors = _author_entries(config)

    _add_sample_run(_new_paragraph(doc, samples["title_id_ppr"]), title_id, samples["title_id_rpr"])
    _new_paragraph(doc, samples["blank_center_ppr"])
    _add_sample_run(_new_paragraph(doc, samples["title_en_ppr"]), title_en, samples["title_en_rpr"])
    _new_paragraph(doc, samples["blank_center_ppr"])

    if authors:
        affiliation_map: dict[tuple[str, str], list[int]] = {}
        for index, author in enumerate(authors, start=1):
            key = (author["affiliation"], author["location"])
            affiliation_map.setdefault(key, []).append(index)

        author_para = _new_paragraph(doc, samples["author_ppr"])
        for index, author in enumerate(authors, start=1):
            if index > 1:
                _add_sample_run(author_para, ", ", samples["author_name_rpr"])
            _add_sample_run(author_para, author["name"], samples["author_name_rpr"])
            _add_sample_run(author_para, str(index), samples["author_sup_rpr"])

        aff_items = list(affiliation_map.items())
        for aff_index, ((affiliation, location), numbers) in enumerate(aff_items):
            sample_ppr = samples["affiliation_primary_ppr"] if aff_index == 0 else samples["affiliation_secondary_ppr"]
            sample_sup = samples["affiliation_primary_sup_rpr"] if aff_index == 0 else samples["affiliation_secondary_sup_rpr"]
            sample_txt = samples["affiliation_primary_text_rpr"] if aff_index == 0 else samples["affiliation_secondary_text_rpr"]
            paragraph = _new_paragraph(doc, sample_ppr)
            _add_sample_run(paragraph, ",".join(str(number) for number in numbers), sample_sup)
            _add_sample_run(paragraph, " ", sample_sup)
            parts = [part for part in (affiliation, location) if part]
            _add_sample_run(paragraph, ", ".join(parts), sample_txt)

        emails = [author["email"] for author in authors if author["email"]]
        if emails:
            email_para = _new_paragraph(doc, samples["email_ppr"])
            _add_sample_run(email_para, "; ".join(emails), samples["email_rpr"])
    else:
        _new_paragraph(doc, samples["blank_center_ppr"])

    _new_paragraph(doc, samples["blank_center_ppr"])

    abstrak_heading = _new_paragraph(doc, samples["abstrak_heading_ppr"])
    _add_sample_run(abstrak_heading, "ABSTRAK", samples["abstrak_heading_rpr"], bold=True)
    _new_paragraph(doc, samples["blank_center_ppr"])

    abstract_id_para = _new_paragraph(doc, samples["abstract_id_ppr"])
    _add_sample_run(abstract_id_para, "(", samples["abstract_id_open_rpr"])
    _add_sample_run(abstract_id_para, "Abstrak", samples["abstract_id_word_rpr"])
    _add_sample_run(abstract_id_para, "). ", samples["abstract_id_close_rpr"])
    _append_rich_text(abstract_id_para, abstract_id, samples["abstract_id_body_rpr"])

    if keywords_id:
        kw_para = _new_paragraph(doc, samples["keywords_id_ppr"])
        _add_sample_run(kw_para, "Kata kunci", samples["keywords_id_label_rpr"])
        _add_sample_run(kw_para, ": ", samples["keywords_id_body_rpr"])
        _append_rich_text(kw_para, ", ".join(keywords_id), samples["keywords_id_body_rpr"])

    _new_paragraph(doc, samples["blank_center_ppr"])

    abstract_en_heading = _new_paragraph(doc, samples["abstract_en_heading_ppr"])
    _add_sample_run(abstract_en_heading, "ABSTRACT", samples["abstract_en_heading_rpr"], bold=True, italic=True)
    _new_paragraph(doc, samples["blank_center_ppr"])

    abstract_en_para = _new_paragraph(doc, samples["abstract_en_ppr"])
    _add_sample_run(abstract_en_para, "(", samples["abstract_en_open_rpr"])
    _add_sample_run(abstract_en_para, "Abstract", samples["abstract_en_word_rpr"])
    _add_sample_run(abstract_en_para, "). ", samples["abstract_en_close_rpr"])
    _append_rich_text(abstract_en_para, abstract_en, samples["abstract_en_body_rpr"])

    if keywords_en:
        kw_en_para = _new_paragraph(doc, samples["keywords_en_ppr"])
        _add_sample_run(kw_en_para, "Keywords", samples["keywords_en_label_rpr"], italic=True)
        _add_sample_run(kw_en_para, ": ", samples["keywords_en_body_rpr"], italic=True)
        _append_rich_text(kw_en_para, ", ".join(keywords_en), samples["keywords_en_body_rpr"])

    _new_paragraph(doc, samples["section_break_ppr"])


def _add_body_text(
    doc: Document,
    text: str,
    samples: dict[str, etree._Element | None],
    *,
    first: bool,
    subsection: bool = False,
) -> bool:
    blocks = [block for block in _split_body_blocks(text) if block.strip()]
    if not blocks:
        return False

    for index, block in enumerate(blocks):
        if subsection:
            paragraph = _new_paragraph(doc, samples["subsection_body_ppr"])
            _append_rich_text(paragraph, block, samples["subsection_body_rpr"])
        else:
            paragraph = _new_paragraph(doc, samples["body_first_ppr"] if first and index == 0 else samples["body_ppr"])
            _append_rich_text(
                paragraph,
                block,
                samples["body_first_rpr"] if first and index == 0 else samples["body_rpr"],
            )
    return True


def _subsection_label(index: int) -> str:
    if 1 <= index <= 26:
        return chr(ord("A") + index - 1)
    return str(index)


def _add_section_heading(doc: Document, title: str, samples: dict[str, etree._Element | None]) -> None:
    paragraph = _new_paragraph(doc, samples["section_heading_ppr"])
    _add_sample_run(paragraph, title.upper(), samples["section_heading_rpr"], bold=True)


def _add_subsection_heading(
    doc: Document,
    title: str,
    samples: dict[str, etree._Element | None],
    label: str,
) -> None:
    paragraph = _new_paragraph(doc, samples["subsection_heading_ppr"])
    _add_sample_run(paragraph, f"{label}. ", samples["subsection_label_rpr"])
    _add_sample_run(paragraph, title, samples["subsection_title_rpr"], italic=True)


def _set_table_width(table, total_width: float, column_count: int) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    tbl_style = tbl_pr.find(qn("w:tblStyle"))
    if tbl_style is None:
        tbl_style = OxmlElement("w:tblStyle")
        tbl_pr.insert(0, tbl_style)
    tbl_style.set(qn("w:val"), "Table1")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), f"{total_width:.1f}")
    tbl_w.set(qn("w:type"), "dxa")

    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = tbl_borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            tbl_borders.append(border)
        border.set(qn("w:color"), "000000")
        border.set(qn("w:space"), "0")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:val"), "single")

    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    column_width = total_width / max(column_count, 1)
    width_text = f"{column_width:.3f}".rstrip("0").rstrip(".")
    for _ in range(column_count):
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), width_text)
        tbl_grid.append(grid_col)


def _set_cell_format(cell, *, header: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        margin = tc_mar.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), "100.0")
        margin.set(qn("w:type"), "dxa")

    v_align = tc_pr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tc_pr.append(v_align)
    v_align.set(qn("w:val"), "top")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header else WD_ALIGN_PARAGRAPH.CENTER
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")

    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLine"), "0")

    ppr_rpr = ppr.find(qn("w:rPr"))
    if ppr_rpr is None:
        ppr_rpr = OxmlElement("w:rPr")
        ppr.append(ppr_rpr)
    for tag, value in (("sz", "20"), ("szCs", "20")):
        el = ppr_rpr.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            ppr_rpr.append(el)
        el.set(qn("w:val"), value)
    if header:
        for tag in ("b", "bCs"):
            el = ppr_rpr.find(qn(f"w:{tag}"))
            if el is None:
                el = OxmlElement(f"w:{tag}")
                ppr_rpr.append(el)
            el.set(qn("w:val"), "1")


def _add_prompt_box(doc: Document, text: str, samples: dict[str, etree._Element | None]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_width(table, 8503.5, 1)
    cell = table.cell(0, 0)
    _set_cell_format(cell, header=False)
    paragraph = cell.paragraphs[0]
    _apply_sample_ppr(paragraph, samples["body_ppr"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _append_rich_text(paragraph, text, samples["body_rpr"])
    # Set RED color for AI prompt
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)


def _next_figure_number(item: dict, state: RenderState) -> str:
    raw = str(item.get("ImageNumber") or item.get("number") or "").strip()
    if raw:
        return raw
    state.figure_count += 1
    return str(state.figure_count)


def _next_table_number(item: dict, state: RenderState) -> str:
    raw = str(item.get("TableNumber") or item.get("NumberiOrLetter") or item.get("number") or "").strip()
    if raw:
        return raw
    state.table_count += 1
    return str(state.table_count)


def _add_figure(doc: Document, item: dict, json_path: Path, samples: dict[str, etree._Element | None], state: RenderState) -> None:
    title = str(item.get("Title") or item.get("title") or "").strip()
    path_text = str(item.get("Path") or item.get("path") or "").strip()
    prompt = str(item.get("Prompt") or "").strip()
    number = _next_figure_number(item, state)
    state.figure_count = max(state.figure_count, int(number)) if number.isdigit() else state.figure_count

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except Exception:
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    image_path = _resolve_path(path_text, json_path) if path_text else None
    if image_path is not None and image_path.is_file():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    else:
        # Combine title and prompt like AIMS (with period separator)
        if title and prompt:
            prompt_text = f"{title}. {prompt}"
        else:
            prompt_text = prompt or title or f"Gambar {number}"
        _add_prompt_box(doc, f"[PROMPT UNTUK AI GAMBAR: {prompt_text}]", samples)

    if title:
        caption = _new_paragraph(doc, samples["figure_caption_ppr"])
        _add_sample_run(caption, f"Gambar {number}. ", samples["figure_caption_label_rpr"])
        _add_sample_run(caption, title, samples["figure_caption_text_rpr"])


def _add_table(doc: Document, item: dict, samples: dict[str, etree._Element | None], state: RenderState) -> None:
    headers = list(item.get("Headers") or item.get("headers") or [])
    rows = list(item.get("Rows") or item.get("rows") or [])
    if not headers:
        return

    number = _next_table_number(item, state)
    try:
        state.table_count = max(state.table_count, int(number))
    except Exception:
        pass

    title = str(item.get("Title") or item.get("title") or "").strip()
    if title:
        caption = _new_paragraph(doc, samples["table_caption_ppr"])
        _add_sample_run(caption, f"Tabel {number}. {title}", samples["table_caption_rpr"])

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_width(table, 8503.5, len(headers))

    for column_index, value in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        cell.text = ""
        _set_cell_format(cell, header=True)
        paragraph = cell.paragraphs[0]
        _append_rich_text(paragraph, str(value), samples["body_rpr"])
        for run in paragraph.runs:
            run.bold = True

    for row_index, row_data in enumerate(rows, start=1):
        for column_index in range(len(headers)):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            _set_cell_format(cell, header=False)
            if column_index < len(row_data):
                paragraph = cell.paragraphs[0]
                _append_rich_text(paragraph, str(row_data[column_index]), samples["body_rpr"])

    doc.add_paragraph()


def _add_equation_group(doc: Document, item: dict, samples: dict[str, etree._Element | None]) -> None:
    formulas = [str(value).strip() for value in item.get("Lines", []) if str(value).strip()]
    single = str(item.get("latex") or item.get("text") or "").strip()
    if not formulas and single:
        formulas = [single]
    if not formulas:
        return

    number = str(item.get("FormulaNumber") or item.get("NumberiOrLetter") or "").strip() or None
    for index, formula in enumerate(formulas):
        paragraph = _new_paragraph(doc, samples["equation_ppr"])
        if not _append_inline_math(paragraph, formula):
            fallback = _add_sample_run(paragraph, formula, samples["equation_spacer_rpr"])
            fallback.italic = True
        if number is not None and index == len(formulas) - 1:
            spacer = _add_sample_run(paragraph, " ", samples["equation_spacer_rpr"])
            spacer.add_tab()
            spacer.add_tab()
            spacer.add_text = None
            spacer = _add_sample_run(paragraph, " ", samples["equation_spacer_rpr"])
            spacer.add_tab()
            _add_sample_run(paragraph, f"({number})", samples["equation_number_rpr"])


def _iter_point_entries(item: dict):
    items = item.get("Items")
    if isinstance(items, list) and items:
        for entry in items:
            if isinstance(entry, dict):
                yield str(entry.get("Text") or "").strip(), str(entry.get("Label") or "").strip()
            else:
                yield str(entry).strip(), ""
        return
    text = str(item.get("Text") or item.get("text") or "").strip()
    if text:
        yield text, str(item.get("Label") or "").strip()


def _add_point_list(doc: Document, item: dict, samples: dict[str, etree._Element | None], *, subsection: bool) -> None:
    list_type = str(item.get("ListType") or ("number" if item.get("Numbered") else "bullet")).lower()
    for index, (text, label) in enumerate(_iter_point_entries(item), start=1):
        if not text:
            continue
        paragraph = _new_paragraph(doc, samples["subsection_body_ppr"] if subsection else samples["body_ppr"])
        prefix = label or (f"{index}. " if list_type in {"number", "numbering", "ordered"} else "- ")
        _add_sample_run(paragraph, prefix, samples["subsection_body_rpr"] if subsection else samples["body_rpr"])
        _append_rich_text(paragraph, text, samples["subsection_body_rpr"] if subsection else samples["body_rpr"])


def _render_content_item(
    doc: Document,
    item: dict,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
    *,
    first_text: bool,
    subsection: bool,
) -> bool:
    item_id = str(item.get("id") or "").lower().strip()
    if item_id == "text":
        text = str(item.get("text") or "").strip()
        if text:
            return _add_body_text(doc, text, samples, first=first_text, subsection=subsection)
        return False
    if item_id in {"gambar", "image", "figure"}:
        _add_figure(doc, item, json_path, samples, state)
        return False
    if item_id in {"tabel", "table"}:
        _add_table(doc, item, samples, state)
        return False
    if item_id in {"rumus", "formula", "equation"}:
        _add_equation_group(doc, item, samples)
        return False
    if item_id in {"poin", "list", "bullet"}:
        _add_point_list(doc, item, samples, subsection=subsection)
        return False
    return False


def _render_content_sequence(
    doc: Document,
    content,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
    *,
    subsection: bool,
) -> None:
    first_text = True
    if isinstance(content, str):
        _add_body_text(doc, content, samples, first=True, subsection=subsection)
        return
    if not isinstance(content, list):
        return

    for item in content:
        if isinstance(item, str):
            if _add_body_text(doc, item, samples, first=first_text, subsection=subsection):
                first_text = False
            continue
        if isinstance(item, dict):
            if _render_content_item(
                doc,
                item,
                json_path,
                samples,
                state,
                first_text=first_text,
                subsection=subsection,
            ):
                first_text = False


def _render_sections(doc: Document, config: dict, json_path: Path, samples: dict[str, etree._Element | None]) -> None:
    state = RenderState()
    section_keys = sorted(
        [key for key in config.keys() if re.fullmatch(r"section\d+", key)],
        key=lambda key: int(key.replace("section", "")),
    )

    for section_key in section_keys:
        section = config.get(section_key)
        if not isinstance(section, dict):
            continue

        title = str(section.get("title") or "").strip()
        if title:
            _add_section_heading(doc, title, samples)

        _render_content_sequence(doc, section.get("content", []), json_path, samples, state, subsection=False)

        subsection_keys = sorted(
            [
                key for key, value in section.items()
                if isinstance(value, dict) and re.fullmatch(rf"{section_key}[a-z]+", key)
            ],
            key=lambda key: key[len(section_key):],
        )
        for subsection_index, subsection_key in enumerate(subsection_keys, start=1):
            subsection_value = section[subsection_key]
            subtitle = str(subsection_value.get("title") or "").strip()
            if subtitle:
                _add_subsection_heading(doc, subtitle, samples, _subsection_label(subsection_index))
            _render_content_sequence(
                doc,
                subsection_value.get("content", []),
                json_path,
                samples,
                state,
                subsection=True,
            )


def _reference_texts(config: dict) -> tuple[str, list[str]]:
    references = config.get("references")
    if isinstance(references, dict):
        title = str(references.get("title") or "DAFTAR PUSTAKA").strip() or "DAFTAR PUSTAKA"
        content = references.get("content", [])
        items = []
        if isinstance(content, list):
            for entry in content:
                if isinstance(entry, dict):
                    text = str(entry.get("text") or entry.get("Text") or "").strip()
                else:
                    text = str(entry).strip()
                if text:
                    items.append(text)
        return title, items
    return "DAFTAR PUSTAKA", []


def _strip_reference_label(text: str) -> str:
    return re.sub(r"^\s*\[\d+\]\s*", "", text).strip()


def _add_references(doc: Document, config: dict, samples: dict[str, etree._Element | None]) -> None:
    title, items = _reference_texts(config)
    if not items:
        return

    heading = _new_paragraph(doc, samples["reference_heading_ppr"])
    _add_sample_run(heading, title.upper(), samples["reference_heading_rpr"], bold=True)

    for item in items:
        paragraph = _new_paragraph(doc, samples["reference_item_ppr"])
        _append_rich_text(paragraph, _strip_reference_label(item), samples["reference_item_rpr"])


def _replace_paragraph_text(paragraph, text: str) -> None:
    sample_rpr = None
    for run in paragraph.runs:
        if run._r.find(qn("w:rPr")) is not None:
            sample_rpr = deepcopy(run._r.find(qn("w:rPr")))
            break
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    _apply_sample_rpr(run, sample_rpr)


def _update_running_headers(doc: Document, config: dict) -> None:
    short_title = _running_title(config)
    short_authors = _running_authors(config)
    if not short_title and not short_authors:
        return

    for section in doc.sections:
        if short_title and section.header.paragraphs:
            _replace_paragraph_text(section.header.paragraphs[0], short_title)
        if short_authors and section.even_page_header.paragraphs:
            _replace_paragraph_text(section.even_page_header.paragraphs[0], short_authors)


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path)
        if output_path is not None
        else Path(json_path).parent / f"{JOURNAL_NAME}_{Path(json_path).stem}.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    samples = _load_template_samples(template_path)

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _clear_document_body(doc)
    _set_document_final_sectpr(doc, samples["body_sectpr"])

    _render_title_block(doc, config, samples)
    _render_sections(doc, config, Path(json_path), samples)
    _new_paragraph(doc, samples["body_close_ppr"])  # Second section break (cols=1 continuous, closes body)
    _add_references(doc, config, samples)
    _update_running_headers(doc, config)

    _set_ai_prompt_color_red(doc)
    doc.save(str(final_output))
    print(f"Generated: {final_output}")
    return final_output


def main() -> None:
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        result = build_document(json_arg, output_arg, template_arg)
        print(f"Selesai: {result}")
        return

    json_files = sorted(
        path for path in BASE_DIR.glob("*.json")
        if path.name.lower() not in {"package.json", "tsconfig.json", "settings.json"}
    )
    for json_file in json_files:
        build_document(json_file)


if __name__ == "__main__":
    main()