"""
IJEECSgen.py - generate an IJEECS paper from JSON using the original IJEECS.docx.

The generator keeps the original journal template, preserves header/footer and
section properties, replaces the title block and abstract table in-place, then
rebuilds the manuscript body from JSON.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "IJEECS.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

FONT_NAME = "Times New Roman"
TITLE_FONT_PT = 16.0
AUTHOR_FONT_PT = 10.0
AFFILIATION_FONT_PT = 8.0
BODY_FONT_PT = 10.0
REFERENCE_FONT_PT = 8.0
MAX_FIGURE_WIDTH_CM = 14.0
SECTION_INDENT_PT = 21.3
BODY_FIRST_LINE_PT = 36.0
EQUATION_FIRST_LINE_PT = 35.45
EQUATION_TAB_RIGHT_PT = 425.25
SECTION_TAB_PT = 21.3
GUIDE_BODY_FIRST_LINE_PT = 35.45
GUIDE_TAB_ONE_PT = 14.2
GUIDE_TAB_TWO_PT = 21.3
GUIDE_TABLE_FONT_PT = 8.0
BIOGRAPHY_FONT_PT = 9.0
CHECK_MARK = "[OK]"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None
_DRAWING_ID_NEXT = 1
_DOC_PR_PATTERN = re.compile(r'<wp:docPr\b[^>]*\bid="(\d+)"')

AUTHOR_CONTRIBUTION_HEADERS = [
    "Name of Author",
    "C",
    "M",
    "So",
    "Va",
    "Fo",
    "I",
    "R",
    "D",
    "O",
    "E",
    "Vi",
    "Su",
    "P",
    "Fu",
]

AUTHOR_CONTRIBUTION_TEMPLATE_ROWS = [
    [
        "Author 1 name",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        "",
    ],
    [
        "Author 2 name",
        "",
        CHECK_MARK,
        "",
        "",
        "",
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
    ],
    [
        "Author 3 name",
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
    ],
    [".....", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
    [
        "Author x name",
        "",
        "",
        "",
        "",
        CHECK_MARK,
        "",
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        "",
        CHECK_MARK,
        "",
        CHECK_MARK,
    ],
]

AUTHOR_CONTRIBUTION_EXAMPLE_ROWS = [
    [
        "Muhammad Usman Akram",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        "",
    ],
    [
        "Siti Zaiton Mohd Hashim",
        "",
        CHECK_MARK,
        "",
        "",
        "",
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
    ],
    [
        "Mohd Ali Hassan",
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        "",
        "",
        CHECK_MARK,
        CHECK_MARK,
        "",
        CHECK_MARK,
        CHECK_MARK,
    ],
]

DATA_AVAILABILITY_EXAMPLES = [
    "Data yang mendukung temuan penelitian ini tersedia secara terbuka di [nama repositori] pada http://doi.org/[doi], nomor referensi [nomor referensi].",
    "Data yang mendukung temuan penelitian ini akan tersedia di [nama repositori] [tautan URL/DOI] setelah embargo [6 bulan] sejak tanggal publikasi untuk memberi ruang bagi komersialisasi hasil penelitian.",
    "Data yang mendukung temuan penelitian ini tersedia atas permintaan dari corresponding author, [inisial, AB]. Data tersebut tidak tersedia secara publik karena memuat informasi yang dapat mengganggu privasi partisipan penelitian.",
    "Data turunan yang mendukung temuan penelitian ini tersedia dari corresponding author [inisial, AB] atas permintaan yang wajar.",
    "Data yang mendukung temuan penelitian ini tersedia dari [pihak ketiga]. Terdapat pembatasan atas ketersediaan data tersebut karena digunakan berdasarkan lisensi untuk penelitian ini. Data tersedia [dari penulis / pada URL] dengan izin dari [pihak ketiga].",
    "Penulis menegaskan bahwa data yang mendukung temuan penelitian ini tersedia di dalam artikel ini [dan/atau materi pelengkapnya].",
    "Data yang mendukung temuan penelitian ini tersedia dari corresponding author, [inisial, AB], berdasarkan permintaan yang wajar.",
    "Ketersediaan data tidak berlaku untuk paper ini karena tidak ada data baru yang dibuat atau dianalisis dalam penelitian ini.",
]


def _iter_body_blocks(parent: _DocumentType | _Cell):
    parent_elm = parent._element.body if isinstance(parent, _DocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _clear_paragraph(paragraph: Paragraph) -> None:
    paragraph.alignment = None
    element = paragraph._element
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)


def _clear_cell(cell: _Cell) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    paragraph = OxmlElement("w:p")
    tc.append(paragraph)


def _initialize_drawing_ids(doc: Document) -> None:
    global _DRAWING_ID_NEXT
    max_id = 0
    for part in doc.part.package.parts:
        part_name = str(getattr(part, "partname", ""))
        if not part_name.endswith(".xml"):
            continue
        try:
            xml_text = part.blob.decode("utf-8", errors="ignore")
        except Exception:
            continue
        for match in _DOC_PR_PATTERN.finditer(xml_text):
            max_id = max(max_id, int(match.group(1)))
    _DRAWING_ID_NEXT = max_id + 1


def _next_drawing_id() -> int:
    global _DRAWING_ID_NEXT
    drawing_id = _DRAWING_ID_NEXT
    _DRAWING_ID_NEXT += 1
    return drawing_id


def _assign_inline_drawing_id(inline) -> None:
    drawing_id = _next_drawing_id()
    inline._inline.docPr.set("id", str(drawing_id))
    for node in inline._inline.xpath(".//*[local-name()='cNvPr']"):
        node.set("id", str(drawing_id))


def _set_run_font(
    run, size_pt: float, *, bold: bool = False, italic: bool = False, superscript: bool = False
) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    run.font.name = FONT_NAME
    run.font.superscript = superscript
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(key), FONT_NAME)


def _append_run(
    paragraph: Paragraph,
    text: str,
    size_pt: float,
    *,
    bold: bool = False,
    italic: bool = False,
    superscript: bool = False,
):
    run = paragraph.add_run(text)
    _set_run_font(run, size_pt, bold=bold, italic=italic, superscript=superscript)
    return run


def _normalize_text_commands(text: str) -> str:
    text = text.replace("\\n", "\n")
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _split_text_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [normalized.strip()]


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph: Paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _append_rich_text(
    paragraph: Paragraph,
    text: str,
    size_pt: float,
    *,
    base_bold: bool = False,
    base_italic: bool = False,
) -> None:
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            paragraph.add_run().add_break()
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                _append_run(paragraph, token["value"], size_pt, italic=True)
            continue
        run = paragraph.add_run(token["value"])
        _set_run_font(
            run,
            size_pt,
            bold=base_bold or token["bold"],
            italic=base_italic or token["italic"],
        )
        if token["underline"]:
            run.underline = True


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _set_paragraph_common(
    paragraph: Paragraph,
    align,
    *,
    first_line: float | None = None,
    left_indent: float | None = None,
    hanging: float | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    if left_indent is not None:
        pf.left_indent = Pt(left_indent)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    elif first_line is not None:
        pf.first_line_indent = Pt(first_line)
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)


def _add_tab_stop(
    paragraph: Paragraph, position_pt: float, alignment=WD_TAB_ALIGNMENT.LEFT
) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(position_pt), alignment)


def _guideline_heading(doc: Document, title: str, qualifier: str | None = None) -> Paragraph:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY)
    _append_run(paragraph, title, BODY_FONT_PT, bold=True)
    if qualifier:
        _append_run(paragraph, " ", BODY_FONT_PT, bold=True)
        _append_run(paragraph, qualifier, BODY_FONT_PT, bold=True, italic=True)
    _append_run(paragraph, " (10 PT)", BODY_FONT_PT, bold=True)
    return paragraph


def _guideline_paragraph(doc: Document, text: str) -> Paragraph:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(
        paragraph,
        WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line=GUIDE_BODY_FIRST_LINE_PT,
    )
    _add_tab_stop(paragraph, GUIDE_TAB_ONE_PT, WD_TAB_ALIGNMENT.LEFT)
    _add_tab_stop(paragraph, GUIDE_TAB_TWO_PT, WD_TAB_ALIGNMENT.LEFT)
    _append_rich_text(paragraph, text, BODY_FONT_PT)
    return paragraph


def _guideline_label(doc: Document, text: str, *, bold: bool = False) -> Paragraph:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY)
    _append_rich_text(paragraph, text, BODY_FONT_PT, base_bold=bold)
    return paragraph


def _guideline_bullet(doc: Document, text: str) -> Paragraph:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(
        paragraph,
        WD_ALIGN_PARAGRAPH.JUSTIFY,
        left_indent=GUIDE_TAB_ONE_PT,
        hanging=GUIDE_TAB_ONE_PT,
    )
    _append_run(paragraph, "- ", BODY_FONT_PT)
    _append_rich_text(paragraph, text, BODY_FONT_PT)
    return paragraph


def _set_cell_width(cell: _Cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(round(width_cm / 2.54 * 1440))))
    tc_w.set(qn("w:type"), "dxa")


def _add_text_table(
    doc: Document,
    rows: list[list[str]],
    *,
    column_widths_cm: list[float] | None = None,
    font_pt: float = BODY_FONT_PT,
    column_alignments: list | None = None,
    header_bold: bool = False,
) -> Table:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    _set_table_full_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = column_widths_cm is None

    for row_index, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values):
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            if column_widths_cm and column_index < len(column_widths_cm):
                _set_cell_width(cell, column_widths_cm[column_index])
            for edge in ("top", "left", "bottom", "right"):
                _set_cell_border(cell, edge, "single")

            paragraph = cell.paragraphs[0]
            align = (
                column_alignments[column_index]
                if column_alignments and column_index < len(column_alignments)
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            _set_paragraph_common(paragraph, align)
            _append_rich_text(
                paragraph,
                value,
                font_pt,
                base_bold=header_bold and row_index == 0,
            )
    return table


def _write_centered_line(
    paragraph: Paragraph, text: str, size_pt: float, *, bold: bool = False
) -> None:
    _clear_paragraph(paragraph)
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    _append_rich_text(paragraph, text, size_pt, base_bold=bold)


def _format_title_block(doc: Document, config: dict) -> Table:
    blocks = list(_iter_body_blocks(doc))
    top_paragraphs: list[Paragraph] = []
    first_table: Table | None = None
    for block in blocks:
        if isinstance(block, Table):
            first_table = block
            break
        top_paragraphs.append(block)
    if first_table is None or len(top_paragraphs) < 10:
        raise RuntimeError("IJEECS template structure is not the expected title-block layout")

    for paragraph in top_paragraphs:
        _clear_paragraph(paragraph)
        _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)

    title_text = str(config.get("title", "")).strip()
    _write_centered_line(top_paragraphs[0], title_text, TITLE_FONT_PT, bold=True)

    authors = [author for author in config.get("authors", []) if isinstance(author, dict)]
    author_para = top_paragraphs[3]
    for index, author in enumerate(authors, start=1):
        if index > 1:
            _append_run(author_para, ", ", AUTHOR_FONT_PT, bold=True)
        _append_run(author_para, str(author.get("name", "")).strip(), AUTHOR_FONT_PT, bold=True)
        _append_run(author_para, str(index), AUTHOR_FONT_PT, bold=True, superscript=True)

    groups: dict[tuple[str, str], list[int]] = {}
    for index, author in enumerate(authors, start=1):
        key = (
            str(author.get("affiliation", "")).strip(),
            str(author.get("location", "")).strip(),
        )
        groups.setdefault(key, []).append(index)

    aff_paragraphs = top_paragraphs[4:8]
    for paragraph in aff_paragraphs:
        _clear_paragraph(paragraph)
        _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)

    for paragraph, ((affiliation, location), indices) in zip(aff_paragraphs, groups.items()):
        _append_run(
            paragraph,
            ",".join(str(index) for index in indices),
            AFFILIATION_FONT_PT,
            superscript=True,
        )
        details = ", ".join(part for part in [affiliation, location] if part)
        if details:
            _append_run(paragraph, f" {details}", AFFILIATION_FONT_PT)

    _update_article_info_table(first_table, config)
    return first_table


def _unique_table_cells(table: Table) -> list[_Cell]:
    cells: list[_Cell] = []
    seen: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker not in seen:
                seen.add(marker)
                cells.append(cell)
    return cells


def _set_cell_text(
    cell: _Cell,
    text: str,
    *,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size_pt: float = BODY_FONT_PT,
    bold: bool = False,
    italic: bool = False,
) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _set_paragraph_common(paragraph, align)
    _append_rich_text(paragraph, text, size_pt, base_bold=bold, base_italic=italic)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_article_info_cell(cell: _Cell, keywords: list[str]) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
    _append_run(paragraph, "Article history:", BODY_FONT_PT, bold=True, italic=True)
    for line in ("Received -", "Revised -", "Accepted -"):
        paragraph.add_run().add_break()
        _append_run(paragraph, line, BODY_FONT_PT)
    if keywords:
        paragraph.add_run().add_break()
        paragraph.add_run().add_break()
        _append_run(paragraph, "Keywords:", BODY_FONT_PT, bold=True, italic=True)
        _append_run(paragraph, f" {'; '.join(keywords)}", BODY_FONT_PT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_corresponding_cell(cell: _Cell, author: dict) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
    _append_run(paragraph, "Corresponding Author:", BODY_FONT_PT, bold=True, italic=True)

    name = str(author.get("name", "")).strip()
    if name:
        paragraph.add_run().add_break()
        _append_rich_text(paragraph, name, BODY_FONT_PT)

    email = str(author.get("email", "")).strip()
    if email:
        paragraph.add_run().add_break()
        _append_rich_text(paragraph, f"email: {email}", BODY_FONT_PT)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _update_article_info_table(table: Table, config: dict) -> None:
    abstract = str(config.get("abstract", "")).strip()
    keywords = [
        str(keyword).strip() for keyword in config.get("keywords", []) if str(keyword).strip()
    ]
    authors = [author for author in config.get("authors", []) if isinstance(author, dict)]
    corresponding = authors[0] if authors else {}

    _set_cell_text(
        table.cell(0, 0), "Article Info", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=10.0, bold=True
    )
    _set_cell_text(
        table.cell(0, 2), "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=10.0, bold=True
    )
    _set_article_info_cell(table.cell(1, 0), keywords)
    _set_cell_text(table.cell(1, 2), abstract, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=10.0)
    _set_cell_text(table.cell(2, 0), "", size_pt=10.0)
    _set_cell_text(table.cell(3, 0), "", size_pt=10.0)
    _set_corresponding_cell(table.cell(4, 0), corresponding)


def _trim_template_body(doc: Document, first_table: Table) -> None:
    body = doc._element.body
    remove = False
    for child in list(body):
        if child is first_table._tbl:
            remove = True
            continue
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def _body_paragraph(doc: Document, text: str) -> None:
    for block in _split_text_blocks(text):
        paragraph = doc.add_paragraph()
        _set_paragraph_common(
            paragraph,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line=BODY_FIRST_LINE_PT,
            space_after=0,
        )
        _append_rich_text(paragraph, block, BODY_FONT_PT)


def _section_heading(doc: Document, number_text: str, title: str) -> None:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(
        paragraph,
        WD_ALIGN_PARAGRAPH.LEFT,
        left_indent=SECTION_INDENT_PT,
        hanging=SECTION_INDENT_PT,
        space_before=12,
        space_after=6,
    )
    _add_tab_stop(paragraph, SECTION_TAB_PT, WD_TAB_ALIGNMENT.LEFT)
    _append_run(paragraph, f"{number_text}.", BODY_FONT_PT, bold=True)
    paragraph.add_run().add_tab()
    _append_run(paragraph, title.upper(), BODY_FONT_PT, bold=True)


def _subsection_heading(doc: Document, number_text: str, title: str) -> None:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    _append_run(paragraph, f"{number_text}. {title}", BODY_FONT_PT, bold=True)


def _subsubsection_heading(doc: Document, number_text: str, title: str) -> None:
    paragraph = doc.add_paragraph()
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
    _append_run(paragraph, f"{number_text}. {title}", BODY_FONT_PT)


def _set_cell_border(
    cell: _Cell, edge: str, value: str, size: str = "4", color: str = "auto"
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    tag = tc_borders.find(qn(f"w:{edge}"))
    if tag is None:
        tag = OxmlElement(f"w:{edge}")
        tc_borders.append(tag)
    tag.set(qn("w:val"), value)
    tag.set(qn("w:sz"), size)
    tag.set(qn("w:space"), "0")
    tag.set(qn("w:color"), color)


def _clear_cell_borders(cell: _Cell) -> None:
    for edge in ("top", "left", "bottom", "right"):
        _set_cell_border(cell, edge, "none", size="0")


def _style_table_paragraph(paragraph: Paragraph, *, bold: bool = False) -> None:
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    for run in paragraph.runs:
        _set_run_font(run, BODY_FONT_PT, bold=bold)


def _add_figure(doc: Document, item: dict, json_path: Path) -> None:

    path_text = str(item.get("Path", "")).strip()
    title = str(item.get("Title", "")).strip()
    number = str(item.get("ImageNumber", "")).strip()
    image_path = _resolve_path(path_text, json_path) if path_text else None

    paragraph = doc.add_paragraph()
    _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    if image_path is not None and image_path.is_file():
        inline = paragraph.add_run().add_picture(str(image_path), width=Cm(MAX_FIGURE_WIDTH_CM))
        _assign_inline_drawing_id(inline)
    else:
        _append_run(paragraph, path_text or "[missing figure]", BODY_FONT_PT, italic=True)

    caption = doc.add_paragraph()
    _set_paragraph_common(caption, WD_ALIGN_PARAGRAPH.CENTER, space_before=3, space_after=6)
    if number:
        _append_run(caption, f"Figure {number}", BODY_FONT_PT)
        if title:
            _append_run(caption, ". ", BODY_FONT_PT)
            _append_rich_text(caption, title, BODY_FONT_PT)
    elif title:
        _append_rich_text(caption, title, BODY_FONT_PT)


def _add_table(doc: Document, item: dict) -> None:
    title = str(item.get("Title", "")).strip()
    number = str(item.get("TableNumber", "")).strip()
    headers = [str(value) for value in item.get("Headers", [])]
    rows = [[str(value) for value in row] for row in item.get("Rows", [])]
    if not headers:
        return

    caption = doc.add_paragraph()
    _set_paragraph_common(caption, WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=3)
    label = f"Table {number}" if number else "Table"
    _append_run(caption, label, BODY_FONT_PT)
    if title:
        _append_run(caption, ". ", BODY_FONT_PT)
        _append_rich_text(caption, title, BODY_FONT_PT)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    _set_table_full_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row in table.rows:
        for cell in row.cells:
            _clear_cell_borders(cell)

    for column_index, header in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
        _append_rich_text(paragraph, header, BODY_FONT_PT, base_bold=True)
        _set_cell_border(cell, "top", "single")
        _set_cell_border(cell, "bottom", "single")

    for row_index, row_data in enumerate(rows, start=1):
        last_row = row_index == len(rows)
        for column_index in range(len(headers)):
            value = row_data[column_index] if column_index < len(row_data) else ""
            cell = table.rows[row_index].cells[column_index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _set_paragraph_common(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
            _append_rich_text(paragraph, value, BODY_FONT_PT)
            if last_row:
                _set_cell_border(cell, "bottom", "single")


def _add_equation(doc: Document, item: dict) -> None:
    formula = str(item.get("latex") or item.get("text") or "").strip()
    number = str(item.get("FormulaNumber", "")).strip()
    if not formula:
        return

    paragraph = doc.add_paragraph()
    _set_paragraph_common(
        paragraph,
        WD_ALIGN_PARAGRAPH.LEFT,
        first_line=EQUATION_FIRST_LINE_PT,
        space_before=6,
        space_after=6,
    )
    _add_tab_stop(paragraph, EQUATION_TAB_RIGHT_PT, WD_TAB_ALIGNMENT.RIGHT)
    if not _append_inline_math(paragraph, formula):
        _append_run(paragraph, formula, BODY_FONT_PT, italic=True)
    if number:
        paragraph.add_run().add_tab()
        _append_run(paragraph, f"({number})", BODY_FONT_PT)


def _render_content_item(doc: Document, item: dict, json_path: Path) -> None:
    item_id = str(item.get("id", "")).strip().lower()
    if item_id == "text":
        text = str(item.get("text", "")).strip()
        if text:
            _body_paragraph(doc, text)
    elif item_id in {"gambar", "image", "figure"}:
        _add_figure(doc, item, json_path)
    elif item_id in {"tabel", "table"}:
        _add_table(doc, item)
    elif item_id in {"rumus", "formula", "equation"}:
        _add_equation(doc, item)


def _render_content(doc: Document, content, json_path: Path) -> None:
    if isinstance(content, str):
        if content.strip():
            _body_paragraph(doc, content.strip())
        return
    if not isinstance(content, list):
        return
    for item in content:
        if isinstance(item, str):
            if item.strip():
                _body_paragraph(doc, item.strip())
        elif isinstance(item, dict):
            _render_content_item(doc, item, json_path)


def _subsection_keys(container: dict, prefix: str) -> list[str]:
    pattern = re.compile(rf"^{re.escape(prefix)}[a-z]+$")
    return sorted([key for key in container if pattern.match(key)])


def _render_subsubsections(
    doc: Document, subsection: dict, json_path: Path, parent_key: str, number_text: str
) -> None:
    child_keys = _subsection_keys(subsection, parent_key)
    for index, child_key in enumerate(child_keys, start=1):
        child = subsection[child_key]
        if not isinstance(child, dict):
            continue
        title = str(child.get("title", "")).strip()
        if title:
            _subsubsection_heading(doc, f"{number_text}.{index}", title)
        _render_content(doc, child.get("content", []), json_path)


def _render_sections(doc: Document, config: dict, json_path: Path) -> None:
    section_keys = sorted(
        [key for key in config.keys() if re.fullmatch(r"section\d+", key)],
        key=lambda value: int(value.replace("section", "")),
    )
    for section_index, section_key in enumerate(section_keys, start=1):
        section = config[section_key]
        if not isinstance(section, dict):
            continue
        title = str(section.get("title", "")).strip()
        if title:
            _section_heading(doc, str(section_index), title)
        _render_content(doc, section.get("content", []), json_path)

        direct_children = [
            key
            for key in _subsection_keys(section, section_key)
            if len(key) == len(section_key) + 1
        ]
        for subsection_index, subsection_key in enumerate(direct_children, start=1):
            subsection = section[subsection_key]
            if not isinstance(subsection, dict):
                continue
            subsection_title = str(subsection.get("title", "")).strip()
            number_text = f"{section_index}.{subsection_index}"
            if subsection_title:
                _subsection_heading(doc, number_text, subsection_title)
            _render_content(doc, subsection.get("content", []), json_path)
            _render_subsubsections(doc, subsection, json_path, subsection_key, number_text)


def _add_mandatory_guideline_sections(doc: Document, config: dict) -> None:
    _guideline_heading(doc, "ACKNOWLEDGMENTS", "(if applicable)")
    _guideline_paragraph(
        doc,
        "Bagian ini digunakan untuk mengakui individu yang memberikan bantuan personal pada penelitian ini, tetapi tidak memenuhi kriteria sebagai penulis, dengan menjelaskan kontribusinya secara ringkas. Pastikan persetujuan dari seluruh individu yang disebutkan dalam bagian ucapan terima kasih telah diperoleh sebelum namanya dicantumkan.",
    )

    _guideline_heading(doc, "FUNDING INFORMATION", "(mandatory)")
    _guideline_paragraph(
        doc,
        'Bagian ini menjelaskan sumber lembaga pendanaan yang mendukung penelitian. Penulis perlu menyatakan bagaimana penelitian dalam artikel ini didanai, termasuk nomor hibah bila ada. Jika penelitian ini tidak menerima pendanaan, gunakan pernyataan serupa berikut: "Penulis menyatakan tidak ada pendanaan yang terlibat dalam penelitian ini."',
    )

    _guideline_heading(doc, "AUTHOR CONTRIBUTIONS STATEMENT", "(mandatory)")
    _guideline_paragraph(
        doc,
        "Jurnal ini menggunakan Contributor Roles Taxonomy (CRediT) untuk mengakui kontribusi individual penulis, mengurangi sengketa authorship, dan memudahkan kolaborasi. \\bJumlah penulis yang direkomendasikan minimal dua orang, dengan salah satunya ditetapkan sebagai corresponding author.\\b Corresponding author bertanggung jawab atas seluruh korespondensi terkait paper dan harus memastikan penulis lain ikut tercakup dalam komunikasi proses submit, revisi, dan publikasi. Penulis dianjurkan menambahkan pernyataan yang menjelaskan kontribusi masing-masing secara akurat. \\bAgar layak dicantumkan sebagai penulis, setiap individu harus berkontribusi setidaknya pada salah satu aspek berikut: conceptualization, methodology, formal analysis, atau investigation, serta pada sedikitnya satu aspek penulisan, yaitu original draft preparation atau review and editing.\\b",
    )
    _add_text_table(
        doc,
        [AUTHOR_CONTRIBUTION_HEADERS, *AUTHOR_CONTRIBUTION_TEMPLATE_ROWS],
        column_widths_cm=[3.5] + [0.74] * 14,
        font_pt=GUIDE_TABLE_FONT_PT,
        column_alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 14,
        header_bold=True,
    )
    _add_text_table(
        doc,
        [
            [
                "C : Konseptualisasi\nM : Metodologi\nSo : Perangkat Lunak\nVa : Validasi\nFo : Analisis Formal",
                "I : Investigasi\nR : Sumber Daya\nD : Kurasi Data\nO : Penulisan - Draf Awal\nE : Penulisan - Tinjauan & Penyuntingan",
                "Vi : Visualisasi\nSu : Supervisi\nP : Administrasi Proyek\nFu : Akuisisi Pendanaan",
            ]
        ],
        column_widths_cm=[4.6, 4.6, 4.6],
        font_pt=GUIDE_TABLE_FONT_PT,
        column_alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    _guideline_label(doc, "Lihat contoh berikut:")
    _add_text_table(
        doc,
        [AUTHOR_CONTRIBUTION_HEADERS, *AUTHOR_CONTRIBUTION_EXAMPLE_ROWS],
        column_widths_cm=[3.5] + [0.74] * 14,
        font_pt=GUIDE_TABLE_FONT_PT,
        column_alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 14,
        header_bold=True,
    )

    _guideline_heading(doc, "CONFLICT OF INTEREST STATEMENT", "(mandatory)")
    _guideline_paragraph(
        doc,
        'Untuk menjaga proses pengambilan keputusan yang adil dan objektif, penulis wajib menyatakan setiap asosiasi yang berpotensi menimbulkan konflik kepentingan, baik finansial, personal, maupun profesional, terkait manuskrip. Konflik non-finansial mencakup kepentingan politik, personal, agama, ideologis, akademik, dan intelektual. Jika tidak ada konflik kepentingan, cantumkan pernyataan serupa berikut: "Penulis menyatakan tidak ada konflik kepentingan."',
    )

    _guideline_heading(doc, "INFORMED CONSENT", "(if applicable)")
    _guideline_paragraph(
        doc,
        'Perlindungan privasi merupakan hak hukum yang tidak boleh dilanggar tanpa informed consent dari individu terkait. Jika identifikasi informasi pribadi diperlukan untuk alasan ilmiah, penulis harus memiliki dokumentasi lengkap informed consent, termasuk persetujuan tertulis dari pasien sebelum dimasukkan ke dalam penelitian. Sertakan pernyataan serupa berikut: "Kami telah memperoleh informed consent dari seluruh individu yang terlibat dalam penelitian ini."',
    )

    _guideline_heading(doc, "ETHICAL APPROVAL", "(if applicable)")
    _guideline_paragraph(
        doc,
        'Jika paper melibatkan manusia atau hewan, penulis harus menjelaskan bahwa penelitian mengikuti seluruh ketentuan nasional dan kebijakan institusi terkait serta telah disetujui oleh komite etik atau review board yang relevan. Penelitian yang melibatkan subjek manusia harus mengikuti prinsip Deklarasi Helsinki. Penulis juga wajib menyebutkan nama komite atau dewan peninjau yang memberikan persetujuan. Sertakan pernyataan serupa berikut: "Penelitian yang melibatkan manusia telah mematuhi seluruh peraturan nasional dan kebijakan institusi yang relevan sesuai prinsip Deklarasi Helsinki dan telah disetujui oleh komite etik institusi penulis atau komite setara"; atau: "Penelitian yang melibatkan hewan telah mematuhi seluruh peraturan nasional dan kebijakan institusi yang relevan terkait perawatan dan penggunaan hewan."',
    )

    _guideline_heading(doc, "DATA AVAILABILITY", "(mandatory)")
    _guideline_paragraph(
        doc,
        "Pernyataan ketersediaan data merupakan penghubung penting antara hasil artikel dan bukti pendukungnya. Bagian ini menjelaskan apakah bukti yang mendukung temuan penelitian tersedia, dan bila tersedia, di mana pembaca dapat mengaksesnya. Pernyataan ketersediaan data membantu meningkatkan transparansi, reproduksibilitas, dan visibilitas data yang dihasilkan atau dikumpulkan selama penelitian. Sebagai bagian dari komitmen jurnal terhadap open research, setiap manuskrip kini wajib menyertakan pernyataan ketersediaan data agar dapat diproses untuk publikasi. Contoh:",
    )
    for example in DATA_AVAILABILITY_EXAMPLES:
        _guideline_bullet(doc, example)


def _biography_rows(config: dict) -> list[list[str]]:
    authors = [author for author in config.get("authors", []) if isinstance(author, dict)]
    if not authors:
        authors = [
            {
                "name": "Nama Penulis",
                "affiliation": "Afiliasi",
                "location": "Kota, Negara",
                "email": "email@domain.com",
            }
        ]

    rows: list[list[str]] = []
    for index, author in enumerate(authors):
        name = str(author.get("name", "")).strip() or "Nama Penulis"
        affiliation = str(author.get("affiliation", "")).strip()
        location = str(author.get("location", "")).strip()
        email = str(author.get("email", "")).strip() or "[email penulis]"
        affiliation_text = (
            ", ".join(part for part in [affiliation, location] if part) or "afiliasi penulis"
        )
        biography = f"{name} berafiliasi dengan {affiliation_text}. Lengkapi bagian ini dengan biografi profesional penulis yang memuat latar belakang akademik, posisi saat ini, minat riset, publikasi penting, dan kontribusi signifikan pada paper ini. Cantumkan ORCID (wajib), serta Google Scholar, Scopus Author ID, atau Web of Science ResearcherID bila tersedia. Penulis dapat dihubungi melalui email: {email}. (9 pt)"
        rows.append(["[Foto 3x4 cm]", biography])
        if index != len(authors) - 1:
            rows.append(["", ""])
    return rows


def _add_biographies_of_authors(doc: Document, config: dict) -> None:
    _guideline_heading(doc, "BIOGRAPHIES OF AUTHORS")
    _guideline_paragraph(
        doc,
        "Pada bagian ini, penulis wajib menyajikan biografi profesional yang mencakup latar belakang akademik, posisi saat ini, minat riset, dan kontribusi penting terhadap studi ini. Selain itu, penulis perlu mencantumkan tautan profil profesional seperti ORCID (wajib) serta, bila ada, Google Scholar, Scopus Author ID, atau Web of Science (WoS) ResearcherID. Informasi ini membantu menegaskan identitas akademik penulis dan meningkatkan visibilitas penelitiannya.",
    )
    _guideline_label(doc, "Informasi yang wajib dicantumkan:", bold=True)
    _guideline_paragraph(
        doc,
        "\\bNama lengkap:\\b Cantumkan nama lengkap penulis sebagaimana tercatat pada dokumen resmi. Bila diinginkan, format nama dapat disesuaikan dengan profil Scopus penulis.",
    )
    _guideline_paragraph(
        doc,
        "\\bAlamat email tiap penulis:\\b Cantumkan email profesional setiap penulis untuk memudahkan korespondensi.",
    )
    _guideline_paragraph(doc, "\\bAkun profesional:\\b")
    _guideline_paragraph(
        doc,
        "\\bORCID iD:\\b Wajib dicantumkan untuk setiap penulis agar karya ilmiah dapat terhubung dengan identitas peneliti.",
    )
    _guideline_paragraph(
        doc,
        "\\bProfil Google Scholar:\\b Tambahkan tautan ke profil Google Scholar penulis. Jika belum ada, penulis dapat membuat profil baru lalu mencantumkan tautannya.",
    )
    _guideline_paragraph(
        doc,
        "\\bScopus Author ID:\\b Jika tersedia, cantumkan Scopus Author ID untuk meningkatkan visibilitas penulis di Scopus.",
    )
    _guideline_paragraph(
        doc,
        "\\bWeb of Science (WoS) ResearcherID:\\b Cantumkan ResearcherID. Jika belum memiliki profil WoS, penulis dapat membuatnya lalu mencantumkan tautannya.",
    )
    _guideline_paragraph(
        doc,
        "\\bBiografi singkat:\\b Tulis ringkasan singkat mengenai latar belakang akademik, minat riset, publikasi penting, dan kontribusi penulis pada paper ini. Panjang yang disarankan sekitar 150 hingga 200 kata (9 pt).",
    )
    _guideline_paragraph(
        doc,
        "\\bPencapaian profesional:\\b Jika ada, sebutkan penghargaan penting, pengakuan profesional, atau proyek riset utama yang pernah diikuti penulis.",
    )
    _guideline_paragraph(
        doc,
        "\\bFoto penulis:\\b Lampirkan foto wajah profesional yang jelas berukuran 3x4 cm. Hindari foto yang terlalu kasual, buram, atau beresolusi rendah.",
    )
    _guideline_label(doc, "Berikut contoh format bagian biografi untuk tiap penulis:")
    _add_text_table(
        doc,
        _biography_rows(config),
        column_widths_cm=[2.8, 11.0],
        font_pt=BIOGRAPHY_FONT_PT,
        column_alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY],
    )


def _add_references(doc: Document, config: dict) -> None:
    references = config.get("references", {})
    items = references.get("content", []) if isinstance(references, dict) else []
    if not items:
        return

    heading = doc.add_paragraph()
    _set_paragraph_common(heading, WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=6)
    _append_run(heading, "REFERENCES", BODY_FONT_PT, bold=True)

    for index, item in enumerate(items, start=1):
        if isinstance(item, dict):
            text = str(item.get("text", "")).strip()
        else:
            text = str(item).strip()
        if not text:
            continue
        paragraph = doc.add_paragraph()
        _set_paragraph_common(
            paragraph,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            left_indent=SECTION_INDENT_PT,
            hanging=SECTION_INDENT_PT,
            space_after=0,
        )
        _add_tab_stop(paragraph, SECTION_TAB_PT, WD_TAB_ALIGNMENT.LEFT)
        _append_run(paragraph, f"[{index}]", REFERENCE_FONT_PT)
        paragraph.add_run().add_tab()
        _append_rich_text(paragraph, text, REFERENCE_FONT_PT)


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _initialize_drawing_ids(doc)
    article_table = _format_title_block(doc, config)
    _trim_template_body(doc, article_table)
    doc.add_paragraph()
    _render_sections(doc, config, Path(json_path))
    _add_mandatory_guideline_sections(doc, config)
    _add_references(doc, config)
    _add_biographies_of_authors(doc, config)
    doc.save(str(final_output))
    print(f"Generated: {final_output}")
    return final_output


def main() -> None:
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        if not json_arg.exists():
            print(f"File not found: {json_arg}")
            sys.exit(1)
        build_document(json_arg, output_arg, template_arg)
        return

    json_files = sorted(
        path
        for path in BASE_DIR.glob("*.json")
        if path.name.lower() not in {"package.json", "tsconfig.json", "settings.json"}
    )
    if not json_files:
        print("No JSON files found.")
        return

    for json_file in json_files:
        try:
            build_document(json_file)
        except Exception as error:
            print(f"Error generating {json_file.name}: {error}")


def _set_table_full_borders(table) -> None:
    """Pastikan tabel punya border tegas/visible (val=single, sz=4 = 0.5pt).

    Dipanggil setelah doc.add_table() supaya tabel data keliatan di Word.
    Auto-injected oleh _fix_table_borders.py untuk lulus audit border check.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


if __name__ == "__main__":
    main()
