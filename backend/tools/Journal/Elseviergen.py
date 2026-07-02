"""
Elsevier Journal Template Generator

Generates papers in Elsevier journal format (single-column, numbered references).
Based on Elsevier article class templates.
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

from _docx_base import (
    open_template,
    para,
    build_sectpr,
    embed_sectpr,
    setup_main_sectpr,
    append_rich_text,
    body_paragraphs,
    render_sections,
)

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "Elsevier.docx"
DEFAULT_OUTPUT_NAME = "Elsevier_output.docx"


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _add_title(doc: Document, config: dict):
    title = config.get("title", "Untitled Paper")
    p = para(doc, style_id="Title", align=WD_ALIGN_PARAGRAPH.LEFT, sa=18)
    run = p.add_run(title)
    run.font.size = Pt(16)
    run.bold = True


def _add_authors(doc: Document, config: dict):
    authors = config.get("authors", [])
    if not authors:
        return
    
    author_names = []
    for author in authors:
        name = author.get("name", "")
        if name:
            author_names.append(name)
    
    if author_names:
        p = para(doc, style_id="Author", align=WD_ALIGN_PARAGRAPH.LEFT)
        run = p.add_run(", ".join(author_names))
        run.font.size = Pt(11)
    
    for author in authors:
        affiliation = author.get("affiliation", "")
        if affiliation:
            p = para(doc, style_id="Affiliation", align=WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(affiliation)
            run.font.size = Pt(9)
            run.italic = True


def _add_highlights(doc: Document, config: dict):
    highlights = config.get("highlights", [])
    if not highlights:
        return
    
    p = para(doc, style_id="Heading1", sb=12)
    run = p.add_run("Highlights")
    run.font.size = Pt(11)
    run.bold = True
    
    for highlight in highlights:
        p = para(doc, style_id="BodyText", li=18)
        p.add_run("• ")
        append_rich_text(p, highlight)


def _add_abstract(doc: Document, config: dict):
    abstract = config.get("abstract", "")
    if not abstract:
        return
    
    p = para(doc, style_id="AbstractHeading", sb=12)
    run = p.add_run("Abstract")
    run.bold = True
    run.font.size = Pt(11)
    
    p = para(doc, style_id="Abstract")
    append_rich_text(p, abstract)


def _add_keywords(doc: Document, config: dict):
    keywords = config.get("keywords", [])
    if not keywords:
        return
    
    p = para(doc, style_id="Keywords", sb=6)
    run = p.add_run("Keywords: ")
    run.bold = True
    append_rich_text(p, "; ".join(keywords))


def _add_references(doc: Document, config: dict):
    references = config.get("references", [])
    if isinstance(references, dict):
        references = references.get("content") or references.get("items") or []
    if not references:
        return
    
    h = para(doc, style_id="Heading1", sb=18)
    h.add_run("References")
    
    for idx, ref in enumerate(references, start=1):
        if isinstance(ref, dict):
            ref_text = _format_reference(ref)
        else:
            ref_text = str(ref)
        
        if ref_text:
            p = para(doc, style_id="Reference", li=36, fi=-36)
            append_rich_text(p, f"[{idx}] {ref_text}")


def build_document(
    json_path: Path | str,
    output_path: Path | str | None = None,
    template_path: Path | str = TEMPLATE_PATH,
) -> Path:
    json_path = Path(json_path)
    config = json.loads(json_path.read_text(encoding="utf-8"))
    
    if output_path:
        final_output = Path(output_path)
    else:
        final_output = json_path.parent / f"{json_path.stem}_Elsevier.docx"
    
    doc = Document()
    
    setup_main_sectpr(
        doc,
        w_pt=595.3,
        h_pt=841.9,
        top_pt=72.0,
        bottom_pt=72.0,
        left_pt=72.0,
        right_pt=72.0,
        num_cols=1,
    )
    
    _add_title(doc, config)
    _add_authors(doc, config)
    _add_highlights(doc, config)
    _add_abstract(doc, config)
    _add_keywords(doc, config)
    
    cfg = {
        "heading1": "Heading1",
        "heading2": "Heading2",
        "body": "BodyText",
        "section_heading_format": "arabic",
        "figure_caption": "Caption",
        "table_head": "TableHead",
        "table_col_head": "TableColHead",
        "table_copy": "TableText",
        "equation": "Equation",
        "max_fig_width_cm": 12.0,
    }
    
    render_sections(doc, config, json_path, BASE_DIR, cfg)
    
    _add_references(doc, config)
    
    final_output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(final_output))
    return final_output


if __name__ == "__main__":
    import sys
    from _docx_base import run_generator
    
    run_generator(BASE_DIR, TEMPLATE_PATH, build_document)
