"""
KKCKgen.py — Generator DOCX untuk jurnal Komunikacija i kultura (KKCK)
Menggunakan dokumen asli KKCK.docx sebagai base template (paste keep formatting).
Data diambil dari _template.json.
"""

import json
import re
import shutil
from pathlib import Path

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
TEMPLATE_DOCX = BASE / "KKCK.docx"
TEMPLATE_JSON = BASE / "_template.json"
OUTPUT_DOCX = BASE / "KKCK_output.docx"


CFG = {
    "page_width_tw": 12240,
    "page_height_tw": 15840,
    "margin_top_tw": 1440,
    "margin_bottom_tw": 1440,
    "margin_left_tw": 1800,
    "margin_right_tw": 1800,
    "header_distance_tw": 720,
    "footer_distance_tw": 720,
    "columns": 1,
    "font_body": "Times New Roman",
    "font_heading": "Times New Roman",
    "font_title": "Times New Roman",
    "size_body": 12,
    "size_heading": 12,
    "size_title": 12,
    "size_caption": 11,
    "size_reference": 12,
    "line_spacing_tw": 360,
    "line_spacing_rule": "auto",
    "first_line_indent_tw": 720,
    "ref_hanging_indent_tw": 720,
}


OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
INLINE_MATH_RE = re.compile(r"\$([^$]+)\$")


def _set_ai_prompt_color_red(doc):
    """Post-process output DOCX:
    1. Set warna text MERAH untuk paragraf prompt AI gambar.
    2. Set border tabel data tegas (single/sz=4) supaya keliatan di Word.
    Idempotent dan aman dipanggil sebelum doc.save()."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    RED = RGBColor(0xFF, 0x00, 0x00)

    def _color_prompt(p):
        text = p.text or ""
        if "[PROMPT UNTUK AI GAMBAR" in text or "[PROMPT AI GAMBAR" in text:
            for r in p.runs:
                try:
                    r.font.color.rgb = RED
                except Exception:
                    pass

    for p in doc.paragraphs:
        _color_prompt(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _color_prompt(p)

    # Set border tabel data (skip layout 1x1, 1xN equation)
    for t in doc.tables:
        rows = t.rows
        if len(rows) < 2 or len(rows[0].cells) < 2:
            continue
        header_text = "".join((c.text or "").strip() for c in rows[0].cells)
        if not header_text:
            continue
        tbl = t._element
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)
        borders = tblPr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = borders.find(qn(f"w:{side}"))
            if el is None:
                el = OxmlElement(f"w:{side}")
                borders.append(el)
            el.set(qn("w:val"), "single" if side in ("top", "bottom", "insideH") else "nil")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


def latex_to_omml_element(latex_src):
    mathml = latex2mathml.converter.convert(latex_src)
    omml_str = mathml2omml.convert(mathml)
    omml_str = re.sub(
        r"(<m:groupChrPr>.*?)</m:groupChr>(\s*<m:e>)",
        r"\1</m:groupChrPr>\2",
        omml_str,
    )
    if not omml_str.startswith("<m:oMath xmlns:m="):
        omml_str = omml_str.replace("<m:oMath>", f'<m:oMath xmlns:m="{OMML_NS}">', 1)
    return parse_xml(omml_str)


def load_json():
    with open(TEMPLATE_JSON, "r", encoding="utf-8") as f:
        return json.load(f)


def ensure_output_dir():
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)


def set_run_font(run, font_name=None, size_pt=None, bold=None, italic=None, color=None):
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    if font_name:
        run.font.name = font_name
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = parse_xml(
                f'<w:rFonts {nsdecls("w")} '
                f'w:ascii="{font_name}" w:hAnsi="{font_name}" '
                f'w:cs="{font_name}" w:eastAsia="{font_name}"/>'
            )
            rpr.insert(0, rfonts)
        else:
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(attr), font_name)


def set_paragraph_spacing(paragraph, before=None, after=None, line=None, line_rule=None):
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = parse_xml(f'<w:spacing {nsdecls("w")}/>')
        ppr.append(spacing)
    if before is not None:
        spacing.set(qn("w:before"), str(before))
    if after is not None:
        spacing.set(qn("w:after"), str(after))
    if line is not None:
        spacing.set(qn("w:line"), str(line))
    if line_rule is not None:
        spacing.set(qn("w:lineRule"), line_rule)


def set_paragraph_indent(paragraph, left=None, right=None, first_line=None, hanging=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = parse_xml(f'<w:ind {nsdecls("w")}/>')
        ppr.append(ind)
    if left is not None:
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:start"), str(left))
    if right is not None:
        ind.set(qn("w:right"), str(right))
        ind.set(qn("w:end"), str(right))
    if first_line is not None:
        ind.set(qn("w:firstLine"), str(first_line))
    if hanging is not None:
        ind.set(qn("w:hanging"), str(hanging))


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag != "sectPr":
            body.remove(child)


def _add_text_with_inline_math(p, text, font_name, size_pt, italic=False, bold=False):
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    pos = 0
    for m in INLINE_MATH_RE.finditer(text):
        before = text[pos : m.start()]
        if before:
            run = p.add_run(before)
            set_run_font(run, font_name, size_pt, bold=bold, italic=italic)
        try:
            omml = latex_to_omml_element(m.group(1))
            p._p.append(omml)
        except Exception:
            run = p.add_run(m.group(0))
            set_run_font(run, font_name, size_pt, bold=bold, italic=True)
        pos = m.end()
    tail = text[pos:]
    if tail:
        run = p.add_run(tail)
        set_run_font(run, font_name, size_pt, bold=bold, italic=italic)


def add_author_block(doc, data):
    authors = data.get("authors", []) or [
        {
            "name": "Author Name",
            "affiliation": "Department, University",
            "location": "City, Country",
            "email": "author@email.ac.id",
        }
    ]

    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p_name, before=0, after=0)
    run = p_name.add_run(", ".join(a.get("name", "Author Name") for a in authors))
    set_run_font(run, CFG["font_body"], CFG["size_body"], bold=True)

    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p_aff, before=0, after=0)
    affs = []
    for a in authors:
        aff_full = a.get("affiliation", "")
        loc = a.get("location", "")
        full = ", ".join(s for s in [aff_full, loc] if s)
        if full and full not in affs:
            affs.append(full)
    run = p_aff.add_run("; ".join(affs))
    set_run_font(run, CFG["font_body"], CFG["size_body"])

    p_em = doc.add_paragraph()
    p_em.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p_em, before=0, after=0)
    run = p_em.add_run(", ".join(a.get("email", "author@email.ac.id") for a in authors))
    set_run_font(run, CFG["font_body"], CFG["size_body"])


def add_blank(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(
        p, before=0, after=0, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )


def add_title(doc, data):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(
        p, before=0, after=0, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    run = p.add_run(data.get("title", "Paper Title Goes Here"))
    set_run_font(run, CFG["font_title"], CFG["size_title"], bold=True)


def add_abstract(doc, data):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=0)
    run_label = p.add_run("Summary. ")
    set_run_font(run_label, CFG["font_body"], CFG["size_body"], bold=True)
    abstract_text = data.get(
        "abstract",
        "Abstract text goes here. This section should contain 150-250 words "
        "summarizing the paper.",
    )
    _add_text_with_inline_math(p, abstract_text, CFG["font_body"], CFG["size_body"])


def add_keywords(doc, data):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, before=0, after=0)
    run_label = p.add_run("Key words: ")
    set_run_font(run_label, CFG["font_body"], CFG["size_body"], bold=True)
    keywords = data.get("keywords", ["keyword1", "keyword2", "keyword3"])
    run_kw = p.add_run(", ".join(keywords) + ".")
    set_run_font(run_kw, CFG["font_body"], CFG["size_body"])


def add_section_heading(doc, title, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(
        p, before=60, after=60, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    run = p.add_run(f"{number}. {title}")
    set_run_font(run, CFG["font_heading"], CFG["size_heading"], bold=True)
    return p


def add_subsection_heading(doc, title, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(
        p, before=60, after=60, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    run = p.add_run(f"{number}. {title}")
    set_run_font(run, CFG["font_heading"], CFG["size_heading"], bold=True)
    return p


def add_body_text(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(
        p, before=0, after=0, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    if indent:
        set_paragraph_indent(p, first_line=CFG["first_line_indent_tw"])
    _add_text_with_inline_math(p, text, CFG["font_body"], CFG["size_body"])
    return p


def add_figure(doc, fig_data):
    img_number = fig_data.get("ImageNumber", "1")
    title = fig_data.get("Title", "Figure Title")
    path = fig_data.get("Path", "")

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(
        p_img, before=60, after=60, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )

    img_full_path = BASE / path if path else None
    if img_full_path and img_full_path.exists():
        run = p_img.add_run()
        run.add_picture(str(img_full_path), width=Cm(12))
    else:
        prompt_text = (
            f"[PROMPT UNTUK AI GAMBAR: {title}. "
            f"Gaya ilustrasi teknis-akademik, latar bersih, resolusi tinggi, "
            f"komposisi terpusat, label elemen jelas dan terbaca.]"
        )
        run = p_img.add_run(prompt_text)
        set_run_font(run, CFG["font_body"], CFG["size_caption"], italic=True)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(
        p_cap, before=60, after=120, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    run_label = p_cap.add_run(f"Figure {img_number}. ")
    set_run_font(run_label, CFG["font_body"], CFG["size_caption"], bold=True)
    run_title = p_cap.add_run(title)
    set_run_font(run_title, CFG["font_body"], CFG["size_caption"])


def add_formula(doc, formula_data):
    formula_num = formula_data.get("FormulaNumber", "1")
    latex = formula_data.get("latex", "E = mc^2")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(
        p, before=60, after=60, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )

    try:
        omml = latex_to_omml_element(latex)
        p._p.append(omml)
    except Exception:
        run = p.add_run(f"  {latex}  ")
        set_run_font(run, CFG["font_body"], CFG["size_body"], italic=True)

    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(
        p_num, before=0, after=60, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )
    run_num = p_num.add_run(f"({formula_num})")
    set_run_font(run_num, CFG["font_body"], CFG["size_body"])


def add_table(doc, table_data):
    tbl_number = table_data.get("TableNumber", "I")
    title = table_data.get("Title", "Table Title")
    headers = table_data.get("Headers", [])
    rows = table_data.get("Rows", [])

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(
        p_title,
        before=120,
        after=60,
        line=CFG["line_spacing_tw"],
        line_rule=CFG["line_spacing_rule"],
    )
    run_label = p_title.add_run(f"Table {tbl_number}. ")
    set_run_font(run_label, CFG["font_body"], CFG["size_caption"], bold=True)
    run_title = p_title.add_run(title)
    set_run_font(run_title, CFG["font_body"], CFG["size_caption"])

    num_cols = len(headers) if headers else 1
    num_rows = len(rows) + 1
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    tbl_el = table._tbl
    tblPr = tbl_el.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl_el.insert(0, tblPr)

    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        "</w:tblBorders>"
    )
    tblPr.append(parse_xml(borders_xml))

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, before=0, after=0)
        run = p.add_run(h)
        set_run_font(run, CFG["font_body"], 11, bold=True)

    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, before=0, after=0)
            run = p.add_run(str(cell_text).replace("\\b", ""))
            set_run_font(run, CFG["font_body"], 11)

    p_after = doc.add_paragraph()
    set_paragraph_spacing(
        p_after, before=60, after=0, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
    )


def process_section_content(doc, content_list):
    if not content_list:
        return

    first_text = True
    for item in content_list:
        item_id = item.get("id", "")

        if item_id == "text":
            text = item.get("text", "")
            add_body_text(doc, text, indent=not first_text)
            first_text = False
        elif item_id == "gambar":
            add_figure(doc, item)
        elif item_id == "rumus":
            add_formula(doc, item)
        elif item_id == "tabel":
            add_table(doc, item)


def process_section(doc, section_key, section_data, level=1, parent_num=None):
    if not section_data:
        return

    title = section_data.get("title", "")

    sub_keys = []
    for key, val in section_data.items():
        if key in ("title", "content"):
            continue
        if isinstance(val, dict) and "title" in val:
            sub_keys.append(key)
    sub_keys.sort(key=lambda x: int(re.search(r"\d+$", x).group()) if re.search(r"\d+$", x) else 0)

    if level == 1:
        m = re.search(r"\d+", section_key)
        number = m.group() if m else "1"
        add_section_heading(doc, title, number)
    else:
        own_idx = re.findall(r"\d+", section_key)
        own_n = own_idx[-1] if own_idx else "1"
        number = f"{parent_num}.{own_n}" if parent_num else own_n
        add_subsection_heading(doc, title, number)

    content = section_data.get("content", [])
    if isinstance(content, list) and content:
        if isinstance(content[0], dict):
            process_section_content(doc, content)
        elif isinstance(content[0], str):
            for text in content:
                add_body_text(doc, text)

    own_idx = re.findall(r"\d+", section_key)
    own_n = own_idx[-1] if own_idx else None
    next_parent = (
        f"{parent_num}.{own_n}"
        if (level > 1 and parent_num and own_n)
        else (own_n if level == 1 else parent_num)
    )
    for sk in sub_keys:
        process_section(doc, sk, section_data[sk], level=level + 1, parent_num=next_parent)


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result


def add_references(doc, data):
    refs_data = data.get("references", {})
    if isinstance(refs_data, list):
        refs_data = {"title": "REFERENCES", "content": refs_data}
    refs_title = refs_data.get("title", "References")

    # Normalize to list of entries
    refs_content = refs_data.get("content", [])
    if not refs_content and isinstance(refs_data, dict):
        for key in ("items", "references"):
            candidate = refs_data.get(key)
            if isinstance(candidate, list):
                refs_content = candidate
                break

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(
        p_title,
        before=60,
        after=60,
        line=CFG["line_spacing_tw"],
        line_rule=CFG["line_spacing_rule"],
    )
    run = p_title.add_run(refs_title)
    set_run_font(run, CFG["font_body"], CFG["size_reference"], bold=True)

    if not refs_content:
        refs_content = ["Author, A. (Year). Title of work. Publisher."]

    for ref in refs_content:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_spacing(
            p, before=0, after=0, line=CFG["line_spacing_tw"], line_rule=CFG["line_spacing_rule"]
        )
        set_paragraph_indent(
            p, left=CFG["ref_hanging_indent_tw"], hanging=CFG["ref_hanging_indent_tw"]
        )
        ref_str = _format_reference(ref)
        run = p.add_run(ref_str)
        set_run_font(run, CFG["font_body"], CFG["size_reference"])


def generate():
    data = load_json()
    ensure_output_dir()

    shutil.copy2(TEMPLATE_DOCX, OUTPUT_DOCX)
    doc = Document(str(OUTPUT_DOCX))
    clear_body(doc)

    add_author_block(doc, data)
    add_blank(doc)
    add_blank(doc)
    add_title(doc, data)
    add_blank(doc)
    add_blank(doc)
    add_abstract(doc, data)
    add_blank(doc)
    add_blank(doc)
    add_keywords(doc, data)
    add_blank(doc)
    add_blank(doc)

    section_keys = [k for k in data.keys() if k.startswith("section") and isinstance(data[k], dict)]
    section_keys.sort(
        key=lambda x: int(re.search(r"\d+", x).group()) if re.search(r"\d+", x) else 0
    )

    for sk in section_keys:
        process_section(doc, sk, data[sk], level=1)

    add_blank(doc)
    add_references(doc, data)

    _set_ai_prompt_color_red(doc)
    doc.save(str(OUTPUT_DOCX))
    print(f"Generated: {OUTPUT_DOCX}")
    return str(OUTPUT_DOCX)


def build_document(json_path: str, output_path: str) -> str:
    """Wrapper for external calls."""
    global TEMPLATE_JSON, OUTPUT_DOCX
    import shutil
    shutil.copy(json_path, str(TEMPLATE_JSON))
    OUTPUT_DOCX = Path(output_path)
    return generate()


def build_pdf(json_path: Path, pdf_path: Path, template_path=None) -> Path:
    """Build a PDF for this journal template from a paper JSON.

    Calls build_document() to produce a .docx, then converts to .pdf
    via LibreOffice headless.  Final PDF is written to ``pdf_path``.
    """
    from ._render_pdf import build_pdf_from_builder
    return build_pdf_from_builder(build_document, json_path, pdf_path, template_path)

if __name__ == "__main__":
    generate()
