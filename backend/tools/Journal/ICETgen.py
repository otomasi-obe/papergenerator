"""
ICETgen.py - Generator DOCX gaya IEEE Conference (template ICET).

Template ICET.docx adalah template IEEE conference standard dengan
paragraph styles bawaan: papertitle, Author, Affiliation, Abstract,
Keywords, Heading1..5, BodyText, references, figurecaption, tablehead,
TABLE, Figure. Generator memakai style bawaan supaya formatting persis
seperti template asli.

Output: ICET_output.docx
"""

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

BASE = Path(__file__).resolve().parent
TEMPLATE_DOCX = BASE / "ICET.docx"
TEMPLATE_JSON = BASE / "_template.json"
OUTPUT_DOCX = BASE / "ICET_output.docx"

# ======================================================================
# CFG dari hasil _analyse.py ICET.docx
# ======================================================================
CFG = {
    # Section title page (1 kolom): pgMar top=540, bottom=1440, l/r=893
    "title_top_tw": 540,
    "title_bottom_tw": 1440,
    "title_left_tw": 893,
    "title_right_tw": 893,
    # Section author block (3 kolom): pgMar top=450, l/r=893, col space=720
    "author_top_tw": 450,
    "author_left_tw": 893,
    "author_right_tw": 893,
    "author_col_space_tw": 720,
    # Section body (2 kolom): pgMar top=1080, bottom=1440, l/r=907
    "body_top_tw": 1080,
    "body_bottom_tw": 1440,
    "body_left_tw": 907,
    "body_right_tw": 907,
    "page_w_tw": 11906,
    "page_h_tw": 16838,
    "header_dist_tw": 720,
    "footer_dist_tw": 720,
    # Body 2-column space
    "col_space_tw": 360,
    # Fonts
    "font_serif": "Times New Roman",
    "font_math": "Cambria Math",
    "font_sym": "Symbol",
    # Sizes
    "size_title": 24.0,
    "size_author": 11.0,
    "size_affil": 10.0,
    "size_abstract": 9.0,
    "size_keywords": 9.0,
    "size_heading1": 10.0,
    "size_heading2": 10.0,
    "size_body": 10.0,
    "size_caption": 8.0,
    "size_table": 8.0,
    "size_formula": 10.0,
    "size_ref": 8.0,
}

# Roman numerals
ROMAN = [
    "",
    "I",
    "II",
    "III",
    "IV",
    "V",
    "VI",
    "VII",
    "VIII",
    "IX",
    "X",
    "XI",
    "XII",
    "XIII",
    "XIV",
    "XV",
    "XVI",
    "XVII",
    "XVIII",
    "XIX",
    "XX",
]

ROMAN_UPPER_TBL = [
    "",
    "I",
    "II",
    "III",
    "IV",
    "V",
    "VI",
    "VII",
    "VIII",
    "IX",
    "X",
    "XI",
    "XII",
    "XIII",
    "XIV",
    "XV",
    "XVI",
    "XVII",
    "XVIII",
    "XIX",
    "XX",
]

LETTER = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

# ======================================================================
# HELPERS
# ======================================================================


def _set_ai_prompt_color_red(doc):
    """Post-process output DOCX:
    1. Set warna text MERAH untuk paragraf prompt AI gambar.
    2. Set border tabel data tegas (single/sz=4) supaya keliatan di Word.
    Idempotent dan aman dipanggil sebelum doc.save()."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    RED = RGBColor(0xFF, 0x00, 0x00)

    def _color_prompt(p):
        text = p.text or ""
        if "[PROMPT UNTUK AI GAMBAR" in text or "[PROMPT AI GAMBAR" in text:
            for r in p.runs:
                try:
                    r.font.color.rgb = RED
                except Exception:
                    pass

    for p in doc.paragraphs:
        _color_prompt(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _color_prompt(p)

    # Set border tabel data (skip layout 1x1, 1xN equation)
    for t in doc.tables:
        rows = t.rows
        if len(rows) < 2 or len(rows[0].cells) < 2:
            continue
        header_text = "".join((c.text or "").strip() for c in rows[0].cells)
        if not header_text:
            continue
        tbl = t._element
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)
        borders = tblPr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = borders.find(qn(f"w:{side}"))
            if el is None:
                el = OxmlElement(f"w:{side}")
                borders.append(el)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


def load_json():
    with open(TEMPLATE_JSON, encoding="utf-8") as f:
        return json.load(f)


def set_run_font(
    run,
    name=None,
    size_pt=None,
    bold=None,
    italic=None,
    color=None,
    superscript=False,
    subscript=False,
):
    if name is not None:
        run.font.name = name
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:ascii"), name)
        rFonts.set(qn("w:hAnsi"), name)
        rFonts.set(qn("w:cs"), name)
        rFonts.set(qn("w:eastAsia"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if superscript:
        run.font.superscript = True
    if subscript:
        run.font.subscript = True


def add_run(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_para_spacing(
    p,
    before_pt=None,
    after_pt=None,
    line_pt=None,
    alignment=None,
    left_tw=None,
    right_tw=None,
    first_line_tw=None,
    hanging_tw=None,
    keep_next=False,
):
    pf = p.paragraph_format
    if before_pt is not None:
        pf.space_before = Pt(before_pt)
    if after_pt is not None:
        pf.space_after = Pt(after_pt)
    if line_pt is not None:
        pf.line_spacing = Pt(line_pt)
    if alignment is not None:
        p.alignment = alignment
    if left_tw is not None:
        pf.left_indent = Twips(left_tw)
    if right_tw is not None:
        pf.right_indent = Twips(right_tw)
    if first_line_tw is not None:
        pf.first_line_indent = Twips(first_line_tw)
    if hanging_tw is not None:
        pf.first_line_indent = Twips(-hanging_tw)
    if keep_next:
        pf.keep_with_next = True


def remove_numpr(p):
    """Lepas numbering otomatis (numPr) dari pPr.

    Penting: kalau style inherit numPr dari styles.xml, hanya remove tag
    di paragraph tidak cukup - Word akan tetap apply numbering dari style.
    Solusi: insert `<w:numPr><w:numId w:val="0"/></w:numPr>` untuk OVERRIDE
    style numPr supaya tidak ada numbering di paragraf ini.
    """
    pPr = p._element.get_or_add_pPr()
    for numPr in pPr.findall(qn("w:numPr")):
        pPr.remove(numPr)
    # Override style numPr dengan numId=0 (no numbering)
    new_numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    new_numPr.append(ilvl)
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    new_numPr.append(numId)
    pPr.append(new_numPr)


def make_sectpr(
    top_tw,
    bottom_tw,
    left_tw,
    right_tw,
    cols=1,
    col_space_tw=720,
    sect_type=None,
    title_pg=False,
    footer_first_rid=None,
):
    sp = OxmlElement("w:sectPr")
    if footer_first_rid:
        fr = OxmlElement("w:footerReference")
        fr.set(qn("w:type"), "first")
        fr.set(qn("r:id"), footer_first_rid)
        sp.append(fr)
    if sect_type:
        t = OxmlElement("w:type")
        t.set(qn("w:val"), sect_type)
        sp.append(t)

    pg_sz = OxmlElement("w:pgSz")
    pg_sz.set(qn("w:w"), str(CFG["page_w_tw"]))
    pg_sz.set(qn("w:h"), str(CFG["page_h_tw"]))
    sp.append(pg_sz)

    pg_mar = OxmlElement("w:pgMar")
    pg_mar.set(qn("w:top"), str(top_tw))
    pg_mar.set(qn("w:right"), str(right_tw))
    pg_mar.set(qn("w:bottom"), str(bottom_tw))
    pg_mar.set(qn("w:left"), str(left_tw))
    pg_mar.set(qn("w:header"), str(CFG["header_dist_tw"]))
    pg_mar.set(qn("w:footer"), str(CFG["footer_dist_tw"]))
    pg_mar.set(qn("w:gutter"), "0")
    sp.append(pg_mar)

    cols_el = OxmlElement("w:cols")
    cols_el.set(qn("w:num"), str(cols))
    cols_el.set(qn("w:space"), str(col_space_tw))
    sp.append(cols_el)

    if title_pg:
        tp = OxmlElement("w:titlePg")
        sp.append(tp)

    return sp


def insert_section_break(
    doc,
    top_tw,
    bottom_tw,
    left_tw,
    right_tw,
    cols,
    col_space_tw=720,
    title_pg=False,
    footer_first_rid=None,
):
    """Sisipkan paragraf kosong dengan sectPr → continuous section break."""
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pPr.append(
        make_sectpr(
            top_tw,
            bottom_tw,
            left_tw,
            right_tw,
            cols=cols,
            col_space_tw=col_space_tw,
            sect_type="continuous",
            title_pg=title_pg,
            footer_first_rid=footer_first_rid,
        )
    )
    return p


def set_final_sectpr(doc):
    body = doc.element.body
    for old in body.findall(qn("w:sectPr")):
        body.remove(old)
    # Final sectPr menyamai original ICET.docx: top=1080, l/r=893, 1 kolom
    body.append(
        make_sectpr(
            CFG["body_top_tw"],
            CFG["body_bottom_tw"],
            CFG["title_left_tw"],
            CFG["title_right_tw"],
            cols=1,
            col_space_tw=CFG["col_space_tw"],
        )
    )


def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag in (qn("w:p"), qn("w:tbl")):
            body.remove(child)
    sp = body.find(qn("w:sectPr"))
    if sp is None:
        body.append(
            make_sectpr(
                CFG["title_top_tw"],
                CFG["title_bottom_tw"],
                CFG["title_left_tw"],
                CFG["title_right_tw"],
                cols=1,
            )
        )


def _resolve_style(doc, style_name):
    """Cari style berdasarkan style_id ATAU display name (suppress deprecation warning)."""
    for s in doc.styles:
        try:
            if s.style_id == style_name or s.name == style_name:
                return s
        except Exception:
            continue
    return None


def style_exists(doc, style_name):
    return _resolve_style(doc, style_name) is not None


# ======================================================================
# LATEX → UNICODE (sama seperti generator lain)
# ======================================================================

LATEX_SYMBOLS = {
    "\\alpha": "α",
    "\\beta": "β",
    "\\gamma": "γ",
    "\\delta": "δ",
    "\\epsilon": "ε",
    "\\theta": "θ",
    "\\lambda": "λ",
    "\\mu": "μ",
    "\\pi": "π",
    "\\sigma": "σ",
    "\\tau": "τ",
    "\\phi": "φ",
    "\\omega": "ω",
    "\\Theta": "Θ",
    "\\Sigma": "Σ",
    "\\Phi": "Φ",
    "\\Omega": "Ω",
    "\\sum": "∑",
    "\\prod": "∏",
    "\\int": "∫",
    "\\infty": "∞",
    "\\pm": "±",
    "\\times": "×",
    "\\cdot": "·",
    "\\leq": "≤",
    "\\geq": "≥",
    "\\neq": "≠",
    "\\approx": "≈",
    "\\rightarrow": "→",
    "\\leftarrow": "←",
    "\\Rightarrow": "⇒",
    "\\sqrt": "√",
    "\\partial": "∂",
    "\\nabla": "∇",
    "\\circ": "°",
    "\\degree": "°",
    "\\quad": "  ",
    "\\,": " ",
    "\\;": " ",
    "\\:": " ",
    "\\!": "",
    "\\left": "",
    "\\right": "",
    "\\cos": "cos",
    "\\sin": "sin",
    "\\tan": "tan",
    "\\log": "log",
    "\\ln": "ln",
    "\\exp": "exp",
    "\\max": "max",
    "\\min": "min",
    "\\arg": "arg",
    "\\to": "→",
    "\\dots": "…",
    "\\ldots": "…",
    "\\cdots": "⋯",
    "\\vec": "",
    "\\hat": "",
    "\\bar": "",
    "\\tilde": "",
    "\\%": "%",
    "\\&": "&",
    "\\#": "#",
    "\\$": "$",
}


def _find_balanced(s, start):
    assert s[start] == "{"
    depth = 1
    i = start + 1
    while i < len(s) and depth > 0:
        if s[i] == "{":
            depth += 1
        elif s[i] == "}":
            depth -= 1
        i += 1
    return i


def _expand_brace_command(s, cmd, transform):
    out = []
    i = 0
    needle = "\\" + cmd
    while i < len(s):
        idx = s.find(needle, i)
        if idx < 0:
            out.append(s[i:])
            break
        j = idx + len(needle)
        if j < len(s) and s[j].isalpha():
            out.append(s[i : idx + 1])
            i = idx + 1
            continue
        out.append(s[i:idx])
        while j < len(s) and s[j] in " \t":
            j += 1
        if j >= len(s) or s[j] != "{":
            out.append(needle)
            i = j
            continue
        end = _find_balanced(s, j)
        inner = s[j + 1 : end - 1]
        out.append(transform(inner))
        i = end
    return "".join(out)


def _expand_frac(s):
    while True:
        idx = s.find("\\frac")
        if idx < 0:
            return s
        j = idx + len("\\frac")
        if j < len(s) and s[j].isalpha():
            return s
        while j < len(s) and s[j] in " \t":
            j += 1
        if j >= len(s) or s[j] != "{":
            return s
        a_end = _find_balanced(s, j)
        a = s[j + 1 : a_end - 1]
        k = a_end
        while k < len(s) and s[k] in " \t":
            k += 1
        if k >= len(s) or s[k] != "{":
            return s
        b_end = _find_balanced(s, k)
        b = s[k + 1 : b_end - 1]
        a_r = latex_to_unicode(a)
        b_r = latex_to_unicode(b)
        s = s[:idx] + f"({a_r})/({b_r})" + s[b_end:]


_SUPER = {
    "0": "⁰",
    "1": "¹",
    "2": "²",
    "3": "³",
    "4": "⁴",
    "5": "⁵",
    "6": "⁶",
    "7": "⁷",
    "8": "⁸",
    "9": "⁹",
    "+": "⁺",
    "-": "⁻",
    "=": "⁼",
    "(": "⁽",
    ")": "⁾",
    "n": "ⁿ",
    "i": "ⁱ",
}
_SUB = {
    "0": "₀",
    "1": "₁",
    "2": "₂",
    "3": "₃",
    "4": "₄",
    "5": "₅",
    "6": "₆",
    "7": "₇",
    "8": "₈",
    "9": "₉",
    "+": "₊",
    "-": "₋",
    "=": "₌",
    "(": "₍",
    ")": "₎",
    "a": "ₐ",
    "e": "ₑ",
    "i": "ᵢ",
    "o": "ₒ",
    "n": "ₙ",
    "k": "ₖ",
    "j": "ⱼ",
    "m": "ₘ",
    "x": "ₓ",
}


def _to_super(s):
    if all(c in _SUPER for c in s):
        return "".join(_SUPER[c] for c in s)
    return "^(" + s + ")"


def _to_sub(s):
    if all(c in _SUB for c in s):
        return "".join(_SUB[c] for c in s)
    return "_" + s if len(s) == 1 else "_(" + s + ")"


def _expand_super_sub(s):
    out = []
    i = 0
    while i < len(s):
        c = s[i]
        if c in ("^", "_") and i + 1 < len(s):
            nxt = s[i + 1]
            if nxt == "{":
                end = _find_balanced(s, i + 1)
                inner = s[i + 2 : end - 1]
                inner = latex_to_unicode(inner)
                out.append(_to_super(inner) if c == "^" else _to_sub(inner))
                i = end
                continue
            elif nxt.isalnum() or nxt in "+-=()":
                out.append(_to_super(nxt) if c == "^" else _to_sub(nxt))
                i += 2
                continue
        out.append(c)
        i += 1
    return "".join(out)


def latex_to_unicode(s: str) -> str:
    if not s:
        return ""
    s = s.strip()
    if s.startswith("$") and s.endswith("$"):
        s = s[1:-1].strip()
    s = s.replace("\\\\", " ; ")
    s = re.sub(
        r"\\begin\{cases\}(.*?)\\end\{cases\}",
        lambda m: " { " + m.group(1).strip() + " }",
        s,
        flags=re.DOTALL,
    )
    for cmd in ("mathrm", "mathbf", "mathit", "mathcal", "text", "operatorname"):
        s = _expand_brace_command(s, cmd, lambda inner: inner)
    s = _expand_frac(s)
    for k in sorted(LATEX_SYMBOLS, key=len, reverse=True):
        s = s.replace(k, LATEX_SYMBOLS[k])
    s = _expand_super_sub(s)
    s = re.sub(r"\^(\w)", lambda m: _to_super(m.group(1)), s)
    s = re.sub(r"_(\w)", lambda m: _to_sub(m.group(1)), s)
    s = s.replace("$", "")
    s = re.sub(r"\\[a-zA-Z]+\*?", "", s)
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"[ \t]+", " ", s).strip()
    return s


# ======================================================================
# INLINE FORMATTING
# ======================================================================

_INLINE_RE = re.compile(
    r"\$([^$]+)\$" r"|\\b([^\\]*?)\\b" r"|\*\*([^*]+)\*\*" r"|\*([^*]+)\*",
    re.DOTALL,
)


def add_runs_with_inline(
    paragraph, text, base_font, base_size, base_bold=False, base_italic=False, base_color=None
):
    pos = 0
    text = text.replace(" ", " ")
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            chunk = text[pos : m.start()]
            if chunk:
                add_run(
                    paragraph,
                    chunk,
                    name=base_font,
                    size_pt=base_size,
                    bold=base_bold,
                    italic=base_italic,
                    color=base_color,
                )
        if m.group(1) is not None:
            content = latex_to_unicode(m.group(1))
            add_run(
                paragraph,
                content,
                name=CFG["font_math"],
                size_pt=base_size,
                italic=True,
                color=base_color,
            )
        elif m.group(2) is not None:
            add_run(
                paragraph,
                m.group(2),
                name=base_font,
                size_pt=base_size,
                bold=True,
                italic=base_italic,
                color=base_color,
            )
        elif m.group(3) is not None:
            add_run(
                paragraph,
                m.group(3),
                name=base_font,
                size_pt=base_size,
                bold=True,
                italic=base_italic,
                color=base_color,
            )
        elif m.group(4) is not None:
            add_run(
                paragraph,
                m.group(4),
                name=base_font,
                size_pt=base_size,
                bold=base_bold,
                italic=True,
                color=base_color,
            )
        pos = m.end()
    if pos < len(text):
        tail = text[pos:]
        if tail:
            add_run(
                paragraph,
                tail,
                name=base_font,
                size_pt=base_size,
                bold=base_bold,
                italic=base_italic,
                color=base_color,
            )


# ======================================================================
# CONTENT GENERATORS
# ======================================================================

_PARA_FIRST = {"first": True}


def reset_first_para():
    _PARA_FIRST["first"] = True


def _new_para(doc, style=None):
    """Buat paragraph baru. Bila style ada di template, pakai. Kalau tidak, fallback Normal."""
    if style:
        s = _resolve_style(doc, style)
        if s is not None:
            p = doc.add_paragraph(style=s)
            return p
    return doc.add_paragraph()


def add_title(doc, data):
    title = data.get("title", "Paper Title Goes Here")
    p = _new_para(doc, "papertitle")
    remove_numpr(p)
    set_para_spacing(p, before_pt=5, after_pt=6, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p, title, name=CFG["font_serif"], size_pt=CFG["size_title"], color=(0, 0, 0))


def add_author_notes(doc, data):
    """Dua paragraf catatan IEEE setelah title (Author style, size 8pt) — match
    paragraf #1 & #2 pada template asli ICET.docx."""
    notes = data.get("ieee_notes") or [
        "*Note: Sub-titles are not captured in Xplore and should not be used",
        "*Note: This paper version is for double blind review. Please remove "
        "all identifying information such as authors' names and affiliations, "
        "as well as acknowledgments before submission.",
    ]
    for note in notes:
        p = _new_para(doc, "Author")
        remove_numpr(p)
        set_para_spacing(p, before_pt=0, after_pt=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_run(p, note, name=CFG["font_serif"], size_pt=8.0, italic=True, color=(0, 0, 0))


def add_authors(doc, data):
    authors = data.get("authors", []) or [
        {
            "name": "Author Name",
            "affiliation": "Department, University",
            "location": "City, Country",
            "email": "author@email.ac.id",
        },
    ]

    # Author names dalam satu paragraf "Author"
    name_p = _new_para(doc, "Author")
    remove_numpr(name_p)
    set_para_spacing(name_p, before_pt=18, after_pt=2, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    for i, a in enumerate(authors):
        if i > 0:
            sep = ", and " if i == len(authors) - 1 else ", "
            add_run(
                name_p, sep, name=CFG["font_serif"], size_pt=CFG["size_author"], color=(0, 0, 0)
            )
        add_run(
            name_p,
            a.get("name", "Author Name"),
            name=CFG["font_serif"],
            size_pt=CFG["size_author"],
            color=(0, 0, 0),
        )
        add_run(
            name_p,
            str(i + 1),
            name=CFG["font_serif"],
            size_pt=CFG["size_author"] * 0.7,
            superscript=True,
            color=(0, 0, 0),
        )

    # Affiliations: tiap author punya satu paragraf Affiliation
    for i, a in enumerate(authors):
        ap = _new_para(doc, "Affiliation")
        remove_numpr(ap)
        set_para_spacing(ap, before_pt=2, after_pt=0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        add_run(
            ap,
            str(i + 1),
            name=CFG["font_serif"],
            size_pt=CFG["size_affil"] * 0.75,
            superscript=True,
            color=(0, 0, 0),
        )
        affil = a.get("affiliation", "Department, University")
        loc = a.get("location", "")
        full = affil + (", " + loc if loc else "")
        add_run(
            ap,
            " " + full,
            name=CFG["font_serif"],
            size_pt=CFG["size_affil"],
            italic=True,
            color=(0, 0, 0),
        )
        if a.get("email"):
            ap2 = _new_para(doc, "Affiliation")
            remove_numpr(ap2)
            set_para_spacing(ap2, before_pt=0, after_pt=2, alignment=WD_ALIGN_PARAGRAPH.CENTER)
            add_run(
                ap2, a["email"], name=CFG["font_serif"], size_pt=CFG["size_affil"], color=(0, 0, 0)
            )


def add_abstract(doc, data):
    abstract = data.get(
        "abstract",
        "Abstract text goes here. This section should contain "
        "150-250 words summarizing the paper.",
    )
    p = _new_para(doc, "Abstract")
    remove_numpr(p)
    set_para_spacing(
        p, before_pt=4, after_pt=8, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_tw=272
    )
    add_run(
        p,
        "Abstract",
        name=CFG["font_serif"],
        size_pt=CFG["size_abstract"],
        bold=True,
        italic=True,
        color=(0, 0, 0),
    )
    add_run(
        p, "—", name=CFG["font_serif"], size_pt=CFG["size_abstract"], bold=True, color=(0, 0, 0)
    )
    add_runs_with_inline(
        p, abstract, base_font=CFG["font_serif"], base_size=CFG["size_abstract"], base_bold=True
    )


def add_keywords(doc, data):
    kws = data.get("keywords", ["keyword1", "keyword2", "keyword3"])
    kw_str = ", ".join(kws) if isinstance(kws, list) else str(kws)
    p = _new_para(doc, "Keywords")
    remove_numpr(p)
    set_para_spacing(
        p, before_pt=4, after_pt=8, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_tw=274
    )
    add_run(
        p,
        "Keywords",
        name=CFG["font_serif"],
        size_pt=CFG["size_keywords"],
        bold=True,
        italic=True,
        color=(0, 0, 0),
    )
    add_run(
        p, "—", name=CFG["font_serif"], size_pt=CFG["size_keywords"], bold=True, color=(0, 0, 0)
    )
    add_run(
        p,
        kw_str,
        name=CFG["font_serif"],
        size_pt=CFG["size_keywords"],
        bold=True,
        italic=True,
        color=(0, 0, 0),
    )


def add_section_heading(doc, title, num):
    """IEEE Heading1: 'I. SECTION TITLE' uppercase center."""
    label = f"{ROMAN[num]}. {title.upper()}"
    p = _new_para(doc, "Heading1")
    remove_numpr(p)
    set_para_spacing(
        p,
        before_pt=8,
        after_pt=4,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_next=True,
        first_line_tw=0,
    )
    add_run(p, label, name=CFG["font_serif"], size_pt=CFG["size_heading1"], color=(0, 0, 0))
    reset_first_para()


def add_subsection_heading(doc, title, sub_idx):
    """IEEE Heading2: 'A. Title Case' italic left."""
    label = f"{LETTER[sub_idx - 1]}. {title}"
    p = _new_para(doc, "Heading2")
    remove_numpr(p)
    set_para_spacing(
        p,
        before_pt=6,
        after_pt=3,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        keep_next=True,
        first_line_tw=0,
    )
    add_run(
        p, label, name=CFG["font_serif"], size_pt=CFG["size_heading2"], italic=True, color=(0, 0, 0)
    )
    reset_first_para()


def add_body_text(doc, text):
    p = _new_para(doc, "BodyText")
    remove_numpr(p)
    first = 0 if _PARA_FIRST["first"] else 288
    _PARA_FIRST["first"] = False
    set_para_spacing(
        p,
        before_pt=0,
        after_pt=6,
        line_pt=11.4,
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        first_line_tw=first,
    )
    add_runs_with_inline(p, text, base_font=CFG["font_serif"], base_size=CFG["size_body"])


def add_figure(doc, fig):
    img_p = _new_para(doc, "Figure")
    remove_numpr(img_p)
    set_para_spacing(
        img_p, before_pt=4, after_pt=2, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_tw=0
    )
    img_path = fig.get("Path", "")
    full_path = BASE / img_path if img_path else None
    inserted = False
    if full_path and full_path.exists():
        try:
            run = img_p.add_run()
            run.add_picture(str(full_path), width=Cm(7.5))
            inserted = True
        except Exception:
            inserted = False
    if not inserted:
        prompt_text = (fig.get("Prompt") or "").strip()
        title_text = (fig.get("Title") or "untitled figure").strip()
        if not prompt_text:
            prompt_text = f"Generate a clean academic figure for: {title_text}"
        placeholder = f"[PROMPT UNTUK AI GAMBAR: {title_text}. {prompt_text}]"
        add_run(
            img_p,
            placeholder,
            name=CFG["font_serif"],
            size_pt=CFG["size_caption"],
            italic=True,
            color=(0xFF, 0x00, 0x00),
        )

    cap = _new_para(doc, "figurecaption")
    remove_numpr(cap)
    set_para_spacing(
        cap,
        before_pt=4,
        after_pt=10,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_tw=0,
        left_tw=0,
    )
    num = fig.get("ImageNumber", "?")
    title = fig.get("Title", "Figure")
    add_run(
        cap, f"Fig. {num}.", name=CFG["font_serif"], size_pt=CFG["size_caption"], color=(0, 0, 0)
    )
    add_run(
        cap, " " + title + ".", name=CFG["font_serif"], size_pt=CFG["size_caption"], color=(0, 0, 0)
    )
    reset_first_para()


def add_formula(doc, fm):
    latex = fm.get("latex", "")
    num = fm.get("FormulaNumber", "?")
    p = doc.add_paragraph()
    set_para_spacing(
        p, before_pt=4, after_pt=4, line_pt=12, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_tw=0
    )
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Twips(2200), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Twips(4400), WD_TAB_ALIGNMENT.RIGHT)
    add_run(p, "\t", name=CFG["font_serif"], size_pt=CFG["size_formula"])
    _omml_done = False
    try:
        from _math_omml import append_omml_math as _omml_fn
        _lx = (latex)
        _omml_done = bool(str(_lx or "").strip()) and _omml_fn(p, _lx)
    except Exception:
        _omml_done = False
    if not _omml_done:
        formula_text = latex_to_unicode(latex)
        add_run(
            p,
            formula_text,
            name=CFG["font_math"],
            size_pt=CFG["size_formula"],
            italic=True,
            color=(0, 0, 0),
        )
    add_run(
        p, "\t(" + num + ")", name=CFG["font_serif"], size_pt=CFG["size_formula"], color=(0, 0, 0)
    )
    reset_first_para()


def _set_cell_three_line(cell, top=None, bottom=None, header_bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is not None:
        tcPr.remove(tcBorders)
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if side == "top" and top:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(top))
        elif side == "bottom" and bottom:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(bottom))
        else:
            el.set(qn("w:val"), "nil")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_table(doc, tb):
    num = tb.get("TableNumber", "?")
    title = tb.get("Title", "Table")
    headers = tb.get("Headers", [])
    rows = tb.get("Rows", [])

    # IEEE table title format: "TABLE I" + newline + title (uppercase first), centered
    cap = _new_para(doc, "tablehead")
    remove_numpr(cap)
    set_para_spacing(
        cap,
        before_pt=8,
        after_pt=4,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_tw=0,
        left_tw=0,
        keep_next=True,
    )
    add_run(cap, f"TABLE {num}", name=CFG["font_serif"], size_pt=CFG["size_table"], color=(0, 0, 0))
    add_run(cap, "\n", name=CFG["font_serif"], size_pt=CFG["size_table"])
    add_runs_with_inline(
        cap,
        title.upper(),
        base_font=CFG["font_serif"],
        base_size=CFG["size_table"],
        base_italic=True,
    )

    if not headers and not rows:
        headers = ["Column 1", "Column 2"]
        rows = [["Data", "Data"]]
    if not headers and rows:
        headers = [f"Col {i+1}" for i in range(len(rows[0]))]

    n_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    last = len(rows) - 1
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.text = ""
        cp = cell.paragraphs[0]
        set_para_spacing(
            cp, before_pt=2, after_pt=2, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_tw=0
        )
        add_runs_with_inline(
            cp, str(h), base_font=CFG["font_serif"], base_size=CFG["size_table"], base_bold=True
        )
        _set_cell_three_line(cell, top=12, bottom=4)

    for i, row in enumerate(rows):
        is_last = i == last
        for j in range(n_cols):
            val = row[j] if j < len(row) else ""
            cell = table.cell(1 + i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            cp = cell.paragraphs[0]
            set_para_spacing(
                cp, before_pt=1, after_pt=1, alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_tw=0
            )
            add_runs_with_inline(
                cp, str(val), base_font=CFG["font_serif"], base_size=CFG["size_table"]
            )
            _set_cell_three_line(cell, top=None, bottom=12 if is_last else None)
    reset_first_para()


def process_section_content(doc, content):
    if not content:
        return
    for item in content:
        if isinstance(item, str):
            add_body_text(doc, item)
            continue
        kind = item.get("id", "")
        if kind == "text":
            add_body_text(doc, item.get("text", ""))
        elif kind == "gambar":
            add_figure(doc, item)
        elif kind == "rumus":
            add_formula(doc, item)
        elif kind == "tabel":
            add_table(doc, item)
        else:
            txt = item.get("text") or item.get("Title") or ""
            if txt:
                add_body_text(doc, txt)


def process_section(doc, sec, num):
    title = sec.get("title", f"Section {num}")
    add_section_heading(doc, title, num)
    if "content" in sec:
        process_section_content(doc, sec["content"])

    sub_keys = sorted(
        [k for k in sec.keys() if k not in ("title", "content") and isinstance(sec[k], dict)]
    )
    sub_idx = 0
    for k in sub_keys:
        sub_idx += 1
        sub = sec[k]
        sub_title = sub.get("title", k)
        add_subsection_heading(doc, sub_title, sub_idx)
        if "content" in sub:
            process_section_content(doc, sub["content"])


def add_references(doc, data):
    refs = data.get("references", {})
    title = "References"
    items = (
        refs.get("content", [])
        if isinstance(refs, dict)
        else (refs if isinstance(refs, list) else [])
    )

    head = _new_para(doc, "Heading1")
    remove_numpr(head)
    set_para_spacing(
        head,
        before_pt=8,
        after_pt=4,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        keep_next=True,
        first_line_tw=0,
    )
    add_run(head, title, name=CFG["font_serif"], size_pt=CFG["size_heading1"], color=(0, 0, 0))

    if not items:
        items = ["[1] Author, Title, Journal, Year."]
    for i, ref in enumerate(items, 1):
        p = _new_para(doc, "references")
        remove_numpr(p)
        set_para_spacing(
            p,
            before_pt=0,
            after_pt=2.5,
            line_pt=9,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            left_tw=288,
            hanging_tw=288,
        )
        add_run(p, f"[{i}] ", name=CFG["font_serif"], size_pt=CFG["size_ref"], color=(0, 0, 0))
        add_runs_with_inline(p, (str(ref.get("text") or ref.get("Text") or "").strip() if isinstance(ref, dict) else str(ref)), base_font=CFG["font_serif"], base_size=CFG["size_ref"])


# ======================================================================
# MAIN
# ======================================================================


def generate():
    if not TEMPLATE_DOCX.exists():
        raise FileNotFoundError(f"Template tidak ditemukan: {TEMPLATE_DOCX}")
    if not TEMPLATE_JSON.exists():
        raise FileNotFoundError(f"JSON tidak ditemukan: {TEMPLATE_JSON}")

    data = load_json()

    shutil.copy2(TEMPLATE_DOCX, OUTPUT_DOCX)
    doc = Document(str(OUTPUT_DOCX))
    clear_body(doc)

    # Original ICET.docx layout: 7 sections.
    #   S1: title (1col, top=540, titlePg, footer rId8)
    #   S2: author block #1 (3col, top=450)
    #   S3: author block #2 (3col, top=450)
    #   S4: abstract+keywords+early body (2col, top=1080)
    #   S5: mid body wide (1col, top=1080) — IEEE typically uses for wide tables/figures
    #   S6: body+references (2col, top=1080)
    #   S7: final (1col, top=1080, l/r=893)

    # ── S1: title page (1 kolom)
    add_title(doc, data)
    add_author_notes(doc, data)
    insert_section_break(
        doc,
        CFG["title_top_tw"],
        CFG["title_bottom_tw"],
        CFG["title_left_tw"],
        CFG["title_right_tw"],
        cols=1,
        col_space_tw=720,
        title_pg=True,
        footer_first_rid="rId8",
    )

    # ── S2: author block (3 kolom) — gabungan author/affiliation
    add_authors(doc, data)
    insert_section_break(
        doc,
        CFG["author_top_tw"],
        CFG["body_bottom_tw"],
        CFG["author_left_tw"],
        CFG["author_right_tw"],
        cols=3,
        col_space_tw=CFG["author_col_space_tw"],
    )

    # ── S3: author block lanjutan (3 kolom) — placeholder paragraph kosong
    insert_section_break(
        doc,
        CFG["author_top_tw"],
        CFG["body_bottom_tw"],
        CFG["author_left_tw"],
        CFG["author_right_tw"],
        cols=3,
        col_space_tw=CFG["author_col_space_tw"],
    )

    # ── S4: abstract + keywords + paruh awal body (2 kolom)
    add_abstract(doc, data)
    add_keywords(doc, data)

    section_keys = sorted(
        [k for k in data.keys() if re.fullmatch(r"section\d+", k)],
        key=lambda s: int(s.replace("section", "")),
    )

    # Split sections: first half → S4, then S5 (wide 1col), remainder → S6
    half = max(1, len(section_keys) // 2)
    for k in section_keys[:half]:
        process_section(doc, data[k], int(k.replace("section", "")))
    insert_section_break(
        doc,
        CFG["body_top_tw"],
        CFG["body_bottom_tw"],
        CFG["body_left_tw"],
        CFG["body_right_tw"],
        cols=2,
        col_space_tw=CFG["col_space_tw"],
    )

    # ── S5: wide 1-kolom (mid body) — empty placeholder paragraph
    insert_section_break(
        doc,
        CFG["body_top_tw"],
        CFG["body_bottom_tw"],
        CFG["body_left_tw"],
        CFG["body_right_tw"],
        cols=1,
        col_space_tw=CFG["col_space_tw"],
    )

    # ── S6: paruh kedua body + references (2 kolom)
    for k in section_keys[half:]:
        process_section(doc, data[k], int(k.replace("section", "")))
    add_references(doc, data)
    insert_section_break(
        doc,
        CFG["body_top_tw"],
        CFG["body_bottom_tw"],
        CFG["body_left_tw"],
        CFG["body_right_tw"],
        cols=2,
        col_space_tw=CFG["col_space_tw"],
    )

    # ── S7: final 1-kolom (margin title l/r=893)
    set_final_sectpr(doc)
    _set_ai_prompt_color_red(doc)
    doc.save(str(OUTPUT_DOCX))
    print(f"Generated: {OUTPUT_DOCX}")
    return str(OUTPUT_DOCX)


if __name__ == "__main__":
    generate()
