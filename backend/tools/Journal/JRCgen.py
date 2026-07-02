"""
JRCgen.py — Generator dokumen JRC (Journal of Robotics and Control) dari JSON.

Menggunakan JRC.docx sebagai template asli (paste keep formatting).
Semua style, numbering, header, footer, dan tema diambil langsung dari template.

Numbering otomatis dari template:
  - Heading1 → I. II. III. (upperRoman, numId=1)
  - Heading2 → A. B. C. (upperLetter, numId=1 lvl1)
  - figurecaption → Fig. 1. Fig. 2. (decimal, numId=5)
  - tablehead → TABLE I. TABLE II. (upperRoman, numId=3)
  - references → [1] [2] [3] (decimal, numId=2)

Semua konten diambil dari _PLC-MediapipeID.json.
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

# ── Paths ─────────────────────────────────────────────────────────
BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "JRC.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

# ── Namespace constants ───────────────────────────────────────────
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

NS_MAP_STRICT = {
    b"http://purl.oclc.org/ooxml/wordprocessingml/main": b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing": b"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    b"http://purl.oclc.org/ooxml/officeDocument/math": b"http://schemas.openxmlformats.org/officeDocument/2006/math",
}

# ── Page/margin dimensions (dari analisis JRC.docx) ──────────────
PAGE_W_PT = 595.3  # A4 width
PAGE_H_PT = 841.9  # A4 height
MARGIN_TOP_FIRST = 27.0  # top margin, title/first page
MARGIN_TOP_BODY = 54.0  # top margin, body pages
MARGIN_BOTTOM = 72.0  # bottom margin
MARGIN_LEFT = 44.65  # left/right margin title section (893tw)
MARGIN_RIGHT = 44.65
MARGIN_LEFT_BODY = 45.35  # left/right margin body section (907tw)
MARGIN_RIGHT_BODY = 45.35
HEADER_DIST = 36.0  # header distance (720tw)
FOOTER_DIST = 36.0  # footer distance (720tw)
COL_SPACE = 36.0  # column gap (720tw)
COL_WIDTH = 243.25  # each column width in 2-col layout (4865tw)
COL1_SPACE_PT = 18.0  # custom child-column spacing in JRC template (360tw)
BODY_COLUMN_WIDTH_PT = COL_WIDTH  # used for equation tab stops
MAX_FIGURE_WIDTH_CM = 8.4  # max figure width per column

# ── XSL for math conversion ───────────────────────────────────────
XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None

# ── Style XML IDs (sesuai JRC.docx styles.xml) ───────────────────
# Key = nama style python-docx, Value = styleId dalam XML
STYLE_XML_ID = {
    "paper title": "papertitle",  # 24pt, Times New Roman
    "Author": "Author",  # author names line
    "Affiliation": "Affiliation",  # affiliation/department
    "Abstract": "Abstract",  # 9pt, bold, firstLine indent
    "Keywords": "Keywords",  # 9pt, bold+italic
    "heading 1": "Heading1",  # Roman auto-numbered, 10pt, center
    "heading 2": "Heading2",  # Letter auto-numbered, 10pt, italic
    "heading 3": "Heading3",  # 1) auto-numbered, 10pt, italic
    "heading 4": "Heading4",  # 10pt, italic
    "heading 5": "Heading5",  # 10pt, spBefore=160
    "Body Text": "BodyText",  # 10pt, justified, firstLine 288tw
    "bullet list": "bulletlist",  # based on BodyText, ind_left 576
    "figure caption": "figurecaption",  # 8pt, auto-Fig.N. (numId=5)
    "table head": "tablehead",  # 8pt, auto-TABLE N. (numId=3)
    "table col head": "tablecolhead",  # 8pt, bold
    "table col subhead": "tablecolsubhead",  # 7.5pt, italic
    "table copy": "tablecopy",  # 8pt, justified
    "table footnote": "tablefootnote",  # 6pt
    "equation": "equation",  # Symbol font, spBefore/After=240
    "references": "references",  # 8pt, auto-[N] (numId=2)
    "sponsors": "sponsors",  # 8pt, firstLine 288
}


# ═══════════════════════════════════════════════════════════════════
# § 1  Utility helpers
# ═══════════════════════════════════════════════════════════════════


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _pt2tw(pt: float) -> int:
    """Convert points to twips (1pt = 20tw)."""
    return int(round(pt * 20))


def _wq(tag: str) -> str:
    return f"{{{NS_W}}}{tag}"


def _rq(tag: str) -> str:
    return f"{{{NS_R}}}{tag}"


def _set_para_style(paragraph, style_name: str) -> None:
    """Set paragraph style by style XML ID (bypasses python-docx name lookup)."""
    xml_id = STYLE_XML_ID.get(style_name, style_name)
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), xml_id)


def _para(
    doc: Document,
    style_id: str | None = None,
    align=None,
    sb: float | None = None,
    sa: float | None = None,
    fi: float | None = None,
    li: float | None = None,
) -> object:
    """Add a paragraph with optional style and formatting."""
    paragraph = doc.add_paragraph()
    if style_id:
        _set_para_style(paragraph, style_id)
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    if sb is not None:
        pf.space_before = Pt(sb)
    if sa is not None:
        pf.space_after = Pt(sa)
    if fi is not None:
        pf.first_line_indent = Pt(fi)
    if li is not None:
        pf.left_indent = Pt(li)
    return paragraph


def _clear_document_body(doc: Document) -> None:
    """Remove all body children except the final sectPr."""
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


# ═══════════════════════════════════════════════════════════════════
# § 2  Math / OMML conversion
# ═══════════════════════════════════════════════════════════════════


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


# ═══════════════════════════════════════════════════════════════════
# § 3  Rich-text helpers
# ═══════════════════════════════════════════════════════════════════


def _append_text_run(
    paragraph, text: str, bold: bool = False, italic: bool = False, underline: bool = False
):
    if not text:
        return
    run = paragraph.add_run(text)
    if bold:
        run.bold = bold
    if italic:
        run.italic = italic
    if underline:
        run.underline = underline


def _append_line_break(paragraph):
    paragraph.add_run().add_break()


def _normalize_text_commands(text: str) -> str:
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    """Parse markup: \\b bold, \\i italic, \\u underline, $...$ math, \\n newline."""
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = italic = underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            cmd = normalized[index + 1]
            if cmd == "\\":
                buffer.append("\\")
                index += 2
                continue
            if cmd == "b":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\b")
                    index += 2
                    continue
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if cmd == "i":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\i")
                    index += 2
                    continue
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if cmd == "u":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\u")
                    index += 2
                    continue
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def _append_rich_text(paragraph, text: str):
    """Append rich-text (with markup) to paragraph."""
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            _append_line_break(paragraph)
        elif token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = paragraph.add_run(token["value"])
                run.italic = True
        else:
            _append_text_run(
                paragraph,
                token["value"],
                bold=token["bold"],
                italic=token["italic"],
                underline=token["underline"],
            )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [p for p in re.split(r"\n\s*\n", normalized) if p.strip()]
    return parts or [""]


def _body_paragraphs(doc: Document, text: str, style_id: str = "Body Text"):
    """Add one or more body paragraphs (split on blank lines)."""
    for block in _split_body_blocks(text):
        paragraph = _para(doc, style_id=style_id)
        _append_rich_text(paragraph, block)


# ═══════════════════════════════════════════════════════════════════
# § 4  Section properties builders
# ═══════════════════════════════════════════════════════════════════


def _get_hf_rids(template_path: Path) -> dict:
    """
    Baca rId header/footer dari word/_rels/document.xml.rels template.
    Mapping berdasarkan analisis JRC.docx:
            header1.xml = default header, header2.xml = first-page header
            footer1.xml = default footer, footer2.xml = first-page footer
    """
    result = {
        "header_default": None,
        "header_first": None,
        "footer_default": None,
        "footer_first": None,
    }
    try:
        with zipfile.ZipFile(str(template_path)) as zf:
            if "word/_rels/document.xml.rels" not in zf.namelist():
                return result
            raw = zf.read("word/_rels/document.xml.rels")
        root = etree.fromstring(raw)
        hf: dict[str, str] = {}
        for rel in root:
            rid = rel.get("Id", "")
            rtype = rel.get("Type", "").lower()
            target = rel.get("Target", "").lower()
            if "header" in rtype:
                if "header1" in target:
                    hf["header1"] = rid
                elif "header2" in target:
                    hf["header2"] = rid
            elif "footer" in rtype:
                if "footer1" in target:
                    hf["footer1"] = rid
                elif "footer2" in target:
                    hf["footer2"] = rid
        # JRC template uses:
        #   default header  -> header1.xml (rId8)
        #   first-page head -> header2.xml (rId10)
        #   default footer  -> footer1.xml (rId9)
        #   first-page foot -> footer2.xml (rId11)
        result["header_default"] = hf.get("header1")
        result["header_first"] = hf.get("header2")
        result["footer_default"] = hf.get("footer1")
        result["footer_first"] = hf.get("footer2")
    except Exception:
        pass
    return result


def _build_title_sectpr(hf_rids: dict) -> etree._Element:
    """
    1-kolom sectPr untuk area judul (halaman pertama).
    Matches: inline sectPr #0 dari analisis JRC.docx.
    top=27pt, margins=44.65pt, 1-col, titlePg, headerRef/footerRef.
    """
    sp = etree.Element(_wq("sectPr"))

    # Match JRC template ordering:
    # header default, footer default, header first, footer first
    rel_order = [
        ("headerReference", "default", "header_default"),
        ("footerReference", "default", "footer_default"),
        ("headerReference", "first", "header_first"),
        ("footerReference", "first", "footer_first"),
    ]
    for tag_name, hf_type, key in rel_order:
        rid = hf_rids.get(key)
        if rid:
            el = etree.SubElement(sp, _wq(tag_name))
            el.set(_wq("type"), hf_type)
            el.set(_rq("id"), rid)

    # pgSz
    pgsz = etree.SubElement(sp, _wq("pgSz"))
    pgsz.set(_wq("w"), str(_pt2tw(PAGE_W_PT)))
    pgsz.set(_wq("h"), str(_pt2tw(PAGE_H_PT)))

    # pgMar (top=27pt untuk area judul)
    pgmar = etree.SubElement(sp, _wq("pgMar"))
    pgmar.set(_wq("top"), str(_pt2tw(MARGIN_TOP_FIRST)))
    pgmar.set(_wq("right"), str(_pt2tw(MARGIN_RIGHT)))
    pgmar.set(_wq("bottom"), str(_pt2tw(MARGIN_BOTTOM)))
    pgmar.set(_wq("left"), str(_pt2tw(MARGIN_LEFT)))
    pgmar.set(_wq("header"), str(_pt2tw(HEADER_DIST)))
    pgmar.set(_wq("footer"), str(_pt2tw(FOOTER_DIST)))
    pgmar.set(_wq("gutter"), "0")

    # cols (1 column)
    cols = etree.SubElement(sp, _wq("cols"))
    cols.set(_wq("space"), str(_pt2tw(COL_SPACE)))

    # titlePg (first page has different header/footer)
    etree.SubElement(sp, _wq("titlePg"))

    # pgNumType (mulai dari halaman 1)
    pgnum = etree.SubElement(sp, _wq("pgNumType"))
    pgnum.set(_wq("start"), "1")

    return sp


def _build_body_sectpr() -> etree._Element:
    """
    2-kolom continuous sectPr untuk area isi (abstract s.d. referensi).
    Matches: inline sectPr #1 dari analisis JRC.docx.
    top=54pt, left/right=45.35pt, 2-col equal 243.25pt each, continuous.
    """
    sp = etree.Element(_wq("sectPr"))

    # Section type: continuous (tidak ganti halaman)
    stype = etree.SubElement(sp, _wq("type"))
    stype.set(_wq("val"), "continuous")

    # pgSz
    pgsz = etree.SubElement(sp, _wq("pgSz"))
    pgsz.set(_wq("w"), str(_pt2tw(PAGE_W_PT)))
    pgsz.set(_wq("h"), str(_pt2tw(PAGE_H_PT)))

    # pgMar
    pgmar = etree.SubElement(sp, _wq("pgMar"))
    pgmar.set(_wq("top"), str(_pt2tw(MARGIN_TOP_BODY)))
    pgmar.set(_wq("right"), str(_pt2tw(MARGIN_RIGHT_BODY)))
    pgmar.set(_wq("bottom"), str(_pt2tw(MARGIN_BOTTOM)))
    pgmar.set(_wq("left"), str(_pt2tw(MARGIN_LEFT_BODY)))
    pgmar.set(_wq("header"), str(_pt2tw(HEADER_DIST)))
    pgmar.set(_wq("footer"), str(_pt2tw(FOOTER_DIST)))
    pgmar.set(_wq("gutter"), "0")

    # cols: 2 kolom dengan lebar eksplisit (equalWidth=0, individual col elements)
    cols = etree.SubElement(sp, _wq("cols"))
    cols.set(_wq("num"), "2")
    cols.set(_wq("space"), str(_pt2tw(COL_SPACE)))  # 720tw = 36pt
    cols.set(_wq("equalWidth"), "0")
    col1 = etree.SubElement(cols, _wq("col"))
    col1.set(_wq("w"), str(_pt2tw(COL_WIDTH)))  # 4865tw = 243.25pt
    col1.set(_wq("space"), str(_pt2tw(COL1_SPACE_PT)))  # 360tw = 18pt
    col2 = etree.SubElement(cols, _wq("col"))
    col2.set(_wq("w"), str(_pt2tw(COL_WIDTH)))
    col2.set(_wq("space"), "0")

    return sp


def _embed_sectpr(
    doc: Document,
    sectpr_el,
    style_id: str | None = None,
    align=None,
    sb: float | None = None,
    sa: float | None = None,
    li: float | None = None,
    fi: float | None = None,
) -> object:
    """Tambah paragraf dengan sectPr yang disematkan di pPr-nya."""
    paragraph = doc.add_paragraph()
    if style_id:
        _set_para_style(paragraph, style_id)
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    if sb is not None:
        pf.space_before = Pt(sb)
    if sa is not None:
        pf.space_after = Pt(sa)
    if li is not None:
        pf.left_indent = Pt(li)
    if fi is not None:
        pf.first_line_indent = Pt(fi)
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(sectpr_el)
    return paragraph


# ═══════════════════════════════════════════════════════════════════
# § 5  Title block
# ═══════════════════════════════════════════════════════════════════


def _add_title(doc: Document, config: dict):
    """
    Judul paper. Style: papertitle (24pt, Times New Roman).
    Tidak bold, tidak italic — sesuai analisis JRC.docx.
    """
    title = config.get("title", "") or config.get("TitleBlock", {}).get("Title", "")
    paragraph = _para(doc, style_id="paper title")
    _append_rich_text(paragraph, title)


def _parse_author_entries(config: dict) -> list[dict]:
    """
    Parse author entries dari format JSON baru (authors array).
    Returns list of author dicts dengan field inti yang dinormalisasi.
    """
    authors = config.get("authors", [])
    if isinstance(authors, list) and authors:
        parsed_entries: list[dict] = []
        for author in authors:
            if not isinstance(author, dict):
                continue
            entry = dict(author)
            entry["name"] = str(author.get("name", "") or "")
            entry["affiliation"] = str(author.get("affiliation", "") or "")
            entry["location"] = str(author.get("location", "") or "")
            entry["email"] = str(author.get("email", "") or "")
            parsed_entries.append(entry)
        return parsed_entries
    return []


_CORRESPONDING_FLAG_KEYS = {
    "corresponding",
    "correspondingauthor",
    "correspondingname",
    "correspondingemail",
    "iscorresponding",
    "coresponding",
    "corespondingauthor",
    "corespondingname",
    "corespondingemail",
    "iscoresponding",
}

_CORRESPONDING_TEXT_KEYS = {
    "role",
    "roles",
    "type",
    "types",
    "label",
    "labels",
    "note",
    "notes",
    "authortype",
}


def _normalize_meta_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _has_config_value(value: object) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        normalized = value.strip().lower()
        return bool(normalized) and normalized not in {"0", "false", "no", "none", "null", "-"}
    if isinstance(value, dict):
        return any(_has_config_value(item) for item in value.values())
    if isinstance(value, (list, tuple, set)):
        return any(_has_config_value(item) for item in value)
    return bool(value)


def _text_mentions_corresponding(value: object) -> bool:
    if isinstance(value, dict):
        return any(_text_mentions_corresponding(item) for item in value.values())
    if isinstance(value, (list, tuple, set)):
        return any(_text_mentions_corresponding(item) for item in value)
    if not isinstance(value, str):
        return False
    normalized = re.sub(r"[^a-z0-9]+", " ", value.lower())
    return "corresponding" in normalized or "coresponding" in normalized


def _has_corresponding_author(config: dict, entries: list[dict]) -> bool:
    sources: list[dict] = [config]
    title_block = config.get("TitleBlock")
    if isinstance(title_block, dict):
        sources.append(title_block)

    for source in [*sources, *entries]:
        if not isinstance(source, dict):
            continue
        for key, value in source.items():
            normalized_key = _normalize_meta_key(key)
            if normalized_key in _CORRESPONDING_FLAG_KEYS and _has_config_value(value):
                return True
            if normalized_key in _CORRESPONDING_TEXT_KEYS and _text_mentions_corresponding(value):
                return True
    return False


def _add_superscript_run(paragraph, text: str) -> object:
    """Add a superscript run to paragraph."""
    run = paragraph.add_run(text)
    rpr = run._r.get_or_add_rPr()
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "superscript")
    rpr.append(va)
    return run


def _add_authors(doc: Document, config: dict):
    """
    Tambah author, affiliation, email sesuai format JRC template.
    Analisis template JRC:
      [001] Author style: "Author^1, Author^2, ..."
      [002] Affiliation style: "^1,2 Dept, Univ, City, Country"
      [003] Normal: "Email: ^1 email1, ^2 email2"
      [004] Normal: "*Corresponding Author"
    """
    entries = _parse_author_entries(config)
    if not entries:
        return

    # ── Kelompokkan author berdasarkan afiliasi unik ──────────────
    # Affiliation key = (affiliation, location)
    aff_map: dict[tuple, int] = {}  # (affiliation, location) → aff_number
    author_aff_nums: list[int] = []  # aff number untuk setiap author
    for entry in entries:
        key = (entry["affiliation"], entry["location"])
        if key not in aff_map:
            aff_map[key] = len(aff_map) + 1
        author_aff_nums.append(aff_map[key])

    # JRC selalu gunakan nomor posisi author (1,2,3...) sebagai superscript
    # bahkan jika semua author berbagi satu afiliasi (sesuai template JRC)

    # ── [001] Author names line ───────────────────────────────────
    # Format: "Rofiq^1, Ahmad^2, Munadi^3"
    author_para = _para(doc, style_id="Author")
    for idx, entry in enumerate(entries):
        if idx > 0:
            author_para.add_run(", ")
        author_para.add_run(entry["name"])
        _add_superscript_run(author_para, str(idx + 1))  # nomor posisi author

    # ── [002] Affiliation lines ───────────────────────────────────
    # Satu Affiliation paragraph per afiliasi unik, superscript = nomor posisi author
    # Contoh: "^1,2 Automation Engineering, Diponegoro University, Semarang, Indonesia"
    # Contoh: "^3 Mechanical Engineering, Diponegoro University, Semarang, Indonesia"
    aff_list: list[tuple] = sorted(aff_map.items(), key=lambda x: x[1])
    for (affiliation, location), aff_num in aff_list:
        # Nomor POSISI author yang punya afiliasi ini (bukan nomor grup)
        author_pos_for_aff = [str(i + 1) for i, an in enumerate(author_aff_nums) if an == aff_num]
        aff_para = _para(doc, style_id="Affiliation")
        _add_superscript_run(aff_para, ", ".join(author_pos_for_aff))
        aff_para.add_run(" ")
        parts = [p for p in [affiliation, location] if p]
        aff_para.add_run(", ".join(parts))

    # ── [003] Email line ──────────────────────────────────────────
    # Format: "Email: ^1 email1, ^2 email2, ^3 email3"
    email_entries = [(i + 1, e["email"]) for i, e in enumerate(entries) if e["email"]]
    if email_entries:
        email_para = doc.add_paragraph()
        email_para.add_run("Email: ")
        for j, (pos_num, email) in enumerate(email_entries):
            if j > 0:
                email_para.add_run(", ")
            _add_superscript_run(email_para, str(pos_num))
            email_para.add_run(f" {email}")

    if _has_corresponding_author(config, entries):
        corresponding_para = _para(doc)
        corresponding_para.add_run("*Corresponding Author")
    _para(doc)


# ═══════════════════════════════════════════════════════════════════
# § 6  Abstract dan Keywords
# ═══════════════════════════════════════════════════════════════════


def _add_abstracts(doc: Document, config: dict):
    """
    Abstract dan Keywords sesuai style JRC:
      Abstract style: 9pt, bold — "Abstract" (italic) + em-dash + isi
      Keywords style: 9pt, bold+italic — "Keywords—kw1; kw2; ..."
    """
    abstract = config.get("abstract", "")
    keywords = config.get("keywords", [])

    # Fallback ke format lama
    if not abstract:
        ab = config.get("Abstract", {})
        abstract = str(ab.get("English") or ab.get("Indonesian") or "").strip()
    if not keywords:
        ab = config.get("Abstract", {})
        kw_str = str(ab.get("KeywordsEnglish") or ab.get("KeywordsIndonesian") or "")
        if kw_str:
            keywords = [k.strip() for k in kw_str.split(",")]

    # Abstract paragraph
    if abstract:
        para = _para(doc, style_id="Abstract")
        # "Abstract" dengan italic (bold dari style)
        r_ab = para.add_run("Abstract")
        r_ab.italic = True
        # Em-dash + isi (bold dari style, tidak italic)
        para.add_run("\u2014")
        _append_rich_text(para, abstract)

    # Keywords paragraph
    if keywords:
        kw_text = "; ".join(str(k).strip() for k in keywords)
        para = _para(doc, style_id="Keywords")
        # Seluruh teks bold+italic dari style
        para.add_run(f"Keywords\u2014{kw_text}")


# ═══════════════════════════════════════════════════════════════════
# § 7  Section headings
# ═══════════════════════════════════════════════════════════════════


def _add_section_heading(doc: Document, text: str):
    """
    Heading1: auto-numbered Roman (I. II. III.) dari numbering.xml JRC.
    Teks UPPERCASE, tidak perlu prefix angka.
    Format: center, 10pt, bold, spBefore=160, spAfter=80.
    """
    para = _para(doc, style_id="heading 1")
    _append_rich_text(para, text.upper() if text == text.lower() else text)


def _add_subsection_heading(doc: Document, text: str):
    """
    Heading2: auto-numbered Letter (A. B. C.) dari numbering.xml JRC.
    Teks normal case, italic dari style.
    Format: left, italic, 10pt, spBefore=120, spAfter=60.
    """
    para = _para(doc, style_id="heading 2")
    _append_rich_text(para, text)


def _add_subsubsection_heading(doc: Document, text: str):
    """
    Heading3: auto-numbered decimal "1)" dari numbering.xml JRC.
    Format: justified, italic, 10pt, line=240 exact.
    """
    para = _para(doc, style_id="heading 3")
    _append_rich_text(para, text)


# ═══════════════════════════════════════════════════════════════════
# § 8  Figures
# ═══════════════════════════════════════════════════════════════════


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _set_full_cell_borders(cell):
    """Set single-line borders semua sisi cell (untuk placeholder box)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "bottom", "left", "right"):
        el = tc_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def _add_prompt_box(doc: Document, text: str):
    """Tambah kotak placeholder (bordered table 1x1) untuk gambar yang belum ada."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Normal Table"
    cell = table.cell(0, 0)
    _set_full_cell_borders(cell)
    para = cell.paragraphs[0]
    _set_para_style(para, "Body Text")
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _append_rich_text(para, text)


def _add_figure(doc: Document, item: dict, json_path: Path):

    """
    Tambah gambar + caption JRC:
      - Gambar: centered, width max 8.4cm (1 column)
      - Caption: figurecaption style (auto-numbered "Fig. N." dari numId=5)
      - Tidak perlu prefix "Fig. X." — template auto-numbering
    """
    path_text = str(item.get("Path", "")).strip()
    prompt = str(item.get("Prompt", "")).strip()
    title = str(item.get("Title", "")).strip()

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except (TypeError, ValueError):
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    image_path = _resolve_path(path_text, json_path) if path_text else None

    # ── Image / placeholder ─────────────────────────────────────
    if image_path is not None and image_path.is_file():
        img_para = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, sb=6.0, sa=2.0)
        img_para.add_run().add_picture(str(image_path), width=Cm(width_cm))
    else:
        fallback = (
            prompt
            if prompt
            else (f"[Gambar {item.get('ImageNumber', '')}] {path_text or 'tidak ditemukan'}")
        )
        _add_prompt_box(doc, fallback)

    # ── Caption (tanpa prefix "Fig. N." — auto dari style numPr) ──
    if title:
        caption = _para(doc, style_id="figure caption")
        _append_rich_text(caption, title)


# ═══════════════════════════════════════════════════════════════════
# § 9  Tables
# ═══════════════════════════════════════════════════════════════════


def _set_table_full_borders(table):
    """Set full borders (top, left, bottom, right, insideH, insideV) sesuai JRC."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def _set_table_style_xml(table, style_id: str):
    """Set table style by XML styleId, including hidden table styles like 'a'."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_style = tbl_pr.find(qn("w:tblStyle"))
    if tbl_style is None:
        tbl_style = OxmlElement("w:tblStyle")
        tbl_pr.insert(0, tbl_style)
    tbl_style.set(qn("w:val"), style_id)


def _set_table_width_fixed(table, total_width_tw: int, column_widths_tw: list[int]):
    """Configure table width/grid to fixed dxa widths."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width_tw))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in column_widths_tw:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)


def _set_cell_width(cell, width_tw: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_tw))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, top=100, left=100, bottom=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = tc_mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def _set_cell_vertical_align(cell, align: str = "center"):
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = tc_pr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tc_pr.append(v_align)
    v_align.set(qn("w:val"), align)


def _configure_equation_table(table):
    """Match JRC equation table layout: fixed 4860tw wide, 4230tw + 630tw columns."""
    _set_table_style_xml(table, "a")
    _set_table_width_fixed(table, 4860, [4230, 630])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    _set_cell_width(left_cell, 4230)
    _set_cell_width(right_cell, 630)
    _set_cell_margins(left_cell, 100, 100, 100, 100)
    _set_cell_margins(right_cell, 100, 100, 100, 100)
    _set_cell_vertical_align(left_cell, "center")
    _set_cell_vertical_align(right_cell, "center")


def _add_table(doc: Document, item: dict):
    """
    Tambah tabel JRC:
      - Caption ATAS: tablehead style (auto-numbered "TABLE I." dari numId=3)
      - Full borders (top+left+bottom+right+insideH+insideV)
      - Header row: tablecolhead style (8pt, bold, center)
      - Data rows: tablecopy style (8pt, justified)
    """
    title = str(item.get("Title") or item.get("title", "")).strip()
    headers = list(item.get("Headers", []) or item.get("headers", []))
    rows = list(item.get("Rows", []) or item.get("rows", []))

    if not headers:
        return

    # Caption ATAS (auto-numbered "TABLE I." dari style numPr, jadi tidak perlu prefix)
    caption = _para(doc, style_id="table head")
    if title:
        _append_rich_text(caption, title)

    # Buat tabel
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_full_borders(table)

    # Header row (tablecolhead: 8pt, bold)
    for col_idx, value in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.text = ""
        para = cell.paragraphs[0]
        _set_para_style(para, "table col head")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(str(value))
        run.bold = True

    # Data rows (tablecopy: 8pt, justified)
    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_data):
            if col_idx >= len(headers):
                break
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            _set_para_style(para, "table copy")
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _append_rich_text(para, str(value))

    # Spasi setelah tabel
    _para(doc, sa=4.0)


# ═══════════════════════════════════════════════════════════════════
# § 10  Equations
# ═══════════════════════════════════════════════════════════════════


def _add_equation_line(doc: Document, formula: str, number: str | None = None):
    """
    Satu baris persamaan JRC:
      - Format asli JRC: tabel 1x2 fixed layout tanpa border
      - Kolom kiri: rumus terpusat dalam lebar 4230tw
      - Kolom kanan: nomor persamaan rata kanan dalam lebar 630tw
    """
    table = doc.add_table(rows=1, cols=2)
    _configure_equation_table(table)

    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)

    left_cell.text = ""
    left_para = left_cell.paragraphs[0]
    left_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    left_pf = left_para.paragraph_format
    left_pf.space_after = Pt(6)
    left_pf.first_line_indent = Pt(14.4)
    left_pf.tab_stops.add_tab_stop(Pt(14.4), WD_TAB_ALIGNMENT.LEFT)
    if not _append_inline_math(left_para, formula):
        run = left_para.add_run(formula)
        run.italic = True

    right_cell.text = ""
    right_para = right_cell.paragraphs[0]
    right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_pf = right_para.paragraph_format
    right_pf.space_after = Pt(6)
    right_pf.right_indent = Pt(-2.95)
    if number is not None:
        right_para.add_run(f"({number})")


def _add_equation_group(doc: Document, item: dict):
    """Tambah satu atau beberapa baris persamaan."""
    # Support format baru: "latex" atau "text" key
    single_latex = str(item.get("latex", "") or item.get("text", "")).strip()
    formulas = [str(f).strip() for f in item.get("Lines", []) if str(f).strip()]
    if not formulas and single_latex:
        formulas = [single_latex]
    if not formulas:
        return

    number = str(item.get("FormulaNumber", "") or item.get("NumberiOrLetter", "")).strip() or None

    for formula in formulas[:-1]:
        _add_equation_line(doc, formula)
    _add_equation_line(doc, formulas[-1], number=number)


# ═══════════════════════════════════════════════════════════════════
# § 11  Bullet/Number lists
# ═══════════════════════════════════════════════════════════════════


def _iter_point_entries(item: dict):
    items = item.get("Items")
    if isinstance(items, list) and items:
        for entry in items:
            if isinstance(entry, str):
                yield entry, None
            elif isinstance(entry, dict):
                yield entry.get("Text", ""), entry.get("Label")
            else:
                yield str(entry), None
        return
    text = item.get("Text", "")
    if text:
        yield text, item.get("Label")


def _add_point_list(doc: Document, item: dict):
    """Tambah daftar bullet atau bernomor."""
    list_type = str(
        item.get("ListType") or ("number" if item.get("Numbered") else "bullet")
    ).lower()
    start_at = int(item.get("StartAt", 1))
    for index, (text, custom_label) in enumerate(_iter_point_entries(item), start=start_at):
        para = _para(doc, style_id="bullet list")
        prefix = custom_label
        if prefix is None:
            prefix = f"{index}. " if list_type in {"number", "numbering", "ordered"} else "\u2022 "
        para.add_run(prefix)
        _append_rich_text(para, text)


# ═══════════════════════════════════════════════════════════════════
# § 12  References
# ═══════════════════════════════════════════════════════════════════


def _add_references(doc: Document, config: dict):
    """
    Tambah section REFERENCES + daftar referensi.
    Heading: Heading1 style (auto-numbered Roman, uppercase)
    Items: references style (auto-numbered "[N]" dari numId=2, 8pt, justified)
    TIDAK perlu prefix "[N]" — auto dari template numbering.
    """
    references: list | None = None

    # Cari referensi dari berbagai format JSON
    if "references" in config and isinstance(config["references"], dict):
        references = config["references"].get("content") or config["references"].get("items") or []
    elif "section_references" in config:
        references = config["section_references"].get("content") or config["section_references"].get("items") or []
    elif "references" in config and isinstance(config["references"], list):
        references = config["references"]
    else:
        for section in config.get("sections", []):
            if section.get("title", "").upper() == "REFERENCES":
                references = section.get("content", [])
                break
    if not references:
        references = config.get("References", [])
    if not references:
        return

    # REFERENCES heading
    _add_section_heading(doc, "REFERENCES")

    # Referensi items
    if references and isinstance(references[0], dict):
        # Format baru: [{id: "1", text: "..."} ...]
        for ref in references:
            ref_text = ref.get("text", ref.get("Text", "")).strip()
            if ref_text:
                para = _para(doc, style_id="references")
                _append_rich_text(para, ref_text)
    else:
        # Format lama: ["[1] G. Eason, ..." ...]
        for ref in references:
            ref_text = re.sub(r"^\s*\[\d+\]\s*", "", str(ref)).strip()
            if ref_text:
                para = _para(doc, style_id="references")
                _append_rich_text(para, ref_text)


# ═══════════════════════════════════════════════════════════════════
# § 13  Content item rendering
# ═══════════════════════════════════════════════════════════════════


def _render_content_item(doc: Document, item: dict, json_path: Path):
    """Render satu item konten dari array content JSON."""
    item_id = str(item.get("id", "")).lower().strip()

    if item_id == "text":
        text = str(item.get("text", ""))
        if text.strip():
            _body_paragraphs(doc, text)

    elif item_id in ("gambar", "image", "figure"):
        _add_figure(doc, item, json_path)

    elif item_id in ("rumus", "formula", "equation"):
        _add_equation_group(doc, item)

    elif item_id in ("tabel", "table"):
        _add_table(doc, item)

    elif item_id in ("poin", "list", "bullet"):
        _add_point_list(doc, item)


def _render_subsection(doc: Document, subsection: dict, json_path: Path, sub_key: str = ""):
    """Render subsection (Heading2 + konten)."""
    title = str(subsection.get("title", "")).strip()

    # Heading2 (auto-lettered A. B. C. dari style)
    if title:
        _add_subsection_heading(doc, title)

    # Konten
    content = subsection.get("content", [])
    if isinstance(content, list):
        for item in content:
            if isinstance(item, dict):
                _render_content_item(doc, item, json_path)
            elif isinstance(item, str) and item.strip():
                _body_paragraphs(doc, item.strip())
    elif isinstance(content, str) and content.strip():
        _body_paragraphs(doc, content.strip())

    # Sub-subsections (key pattern: section{N}{letters}+, misal section2a -> section2ab)
    # Atau juga pola lama: sub{N}{letters}
    if sub_key:
        # Cari sub-keys di dalam subsection yang mengindikasikan sub-subsection
        supp_keys = [
            k
            for k in subsection.keys()
            if (
                re.match(r"^section\d+[a-z]+[a-z]+$", k)
                or (k.startswith("sub") and len(k) > 3 and k[3:].isalnum())
            )
        ]
        for sk in supp_keys:
            subsubsec = subsection[sk]
            if isinstance(subsubsec, dict):
                _render_subsubsection(doc, subsubsec, json_path)


def _render_subsubsection(doc: Document, subsubsec: dict, json_path: Path):
    """Render sub-subsection (Heading3 + konten)."""
    title = str(subsubsec.get("title", "")).strip()
    if title:
        _add_subsubsection_heading(doc, title)
    content = subsubsec.get("content", [])
    if isinstance(content, list):
        for item in content:
            if isinstance(item, dict):
                _render_content_item(doc, item, json_path)
            elif isinstance(item, str) and item.strip():
                _body_paragraphs(doc, item.strip())
    elif isinstance(content, str) and content.strip():
        _body_paragraphs(doc, content.strip())


def _render_sections(doc: Document, config: dict, json_path: Path):
    """
    Render semua section (section1, section2, ...) dari config JSON.
    Setiap section:
      - Heading1 (judul UPPERCASE, auto-Roman dari style)
      - Konten teks/gambar/rumus/tabel langsung
      - Subsections: keys section{N}[a-z]+ atau sub{N}[a-z]+
    """
    section_keys = sorted(
        [k for k in config.keys() if re.match(r"^section\d+$", k)],
        key=lambda x: int(x.replace("section", "")),
    )

    for section_key in section_keys:
        section = config[section_key]
        section_num = section_key.replace("section", "")
        section_title = str(section.get("title", "")).strip()

        # Heading1 — auto-numbered dari style, teks judul UPPERCASE
        _add_section_heading(doc, section_title)

        # Konten langsung di level section
        content = section.get("content", "")
        if isinstance(content, str) and content.strip():
            _body_paragraphs(doc, content.strip())
        elif isinstance(content, list):
            for item in content:
                if isinstance(item, str) and item.strip():
                    _body_paragraphs(doc, item.strip())
                elif isinstance(item, dict):
                    _render_content_item(doc, item, json_path)

        # Subsections: section{N}[a-z]+ (misal section2a, section2b)
        # atau sub{N}[a-z]+ (format lama)
        subsection_keys = sorted(
            [
                k
                for k in section.keys()
                if re.match(rf"^section{section_num}[a-z]+$", k)
                or (k.startswith("sub") and len(k) > 3 and k[3:].isalnum())
            ]
        )
        for sub_key in subsection_keys:
            subval = section[sub_key]
            if isinstance(subval, dict):
                _render_subsection(doc, subval, json_path, sub_key=sub_key)

    # Fallback: format lama dengan "sections" array
    if not section_keys and "sections" in config:
        for section in config.get("sections", []):
            str(section.get("number", "")).strip()
            sec_title = str(section.get("title", "")).strip()
            heading = sec_title.upper()
            _add_section_heading(doc, heading)
            content = section.get("content", "")
            if isinstance(content, str) and content.strip():
                _body_paragraphs(doc, content.strip())
            elif isinstance(content, list):
                for item in content:
                    if isinstance(item, str):
                        _body_paragraphs(doc, item.strip())
                    elif isinstance(item, dict):
                        _render_content_item(doc, item, json_path)
            for sub in section.get("subsections", []):
                _render_subsection(doc, sub, json_path)

    # Format paling lama: "Items" array (JTM format)
    if not section_keys and "sections" not in config and "Items" in config:
        for item in config.get("Items", []):
            _render_item_legacy(doc, item, json_path)


def _render_item_legacy(doc: Document, item: dict, json_path: Path):
    """Render item format lama JTM (ID-based)."""
    item_id = str(item.get("ID") or "").strip()
    str(item.get("NumberiOrLetter") or "").strip()
    text_val = str(item.get("Text") or "").strip()

    if item_id in {"Bab", "Bab utama"}:
        _add_section_heading(doc, text_val.upper())
    elif item_id in {"SubBab", "Bab 1.1"}:
        _add_subsection_heading(doc, text_val)
    elif item_id == "SubSubBab":
        _add_subsubsection_heading(doc, text_val)
    elif item_id == "Paragraf":
        _body_paragraphs(doc, text_val)
    elif item_id == "Poin":
        _add_point_list(doc, item)
    elif item_id == "Gambar":
        _add_figure(doc, item, json_path)
    elif item_id == "Tabel":
        _add_table(doc, item)
    elif item_id == "Rumus":
        _add_equation_group(doc, item)


# ═══════════════════════════════════════════════════════════════════
# § 14  Main build function
# ═══════════════════════════════════════════════════════════════════


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    """
    Generate JRC DOCX dari JSON.

    Alur:
      1. Buka JRC.docx template langsung (preserve styles, numbering, header/footer)
      2. Kosongkan body (biarkan final sectPr)
      3. Tambah title block → inline sectPr 1-col (dengan headerRef/footerRef)
      4. Tambah abstract, keywords, sections, references
      5. Trailing paragraph → inline sectPr 2-col continuous
      6. Final sectPr dari template (1-col) tersisa sebagai penutup dokumen
    """
    import shutil

    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    # ── Buka template langsung (paste keep formatting) ────────────
    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _clear_document_body(doc)

    # Ambil rId header/footer dari template
    hf_rids = _get_hf_rids(template_path)

    # ═══ SECTION 1: Title area (1 kolom, halaman pertama) ════════
    _add_title(doc, config)
    _add_authors(doc, config)
    # Paragraf pemisah dengan inline sectPr 1-kolom
    # Spacing sesuai analisis template: spBefore=18pt, spAfter=14pt
    _embed_sectpr(
        doc,
        _build_title_sectpr(hf_rids),
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        sb=18.0,
        sa=14.0,
    )

    # ═══ SECTION 2: Body content (2 kolom, continuous) ═══════════
    _add_abstracts(doc, config)
    _render_sections(doc, config, Path(json_path))
    _add_references(doc, config)

    # Trailing paragraph untuk menutup 2-kolom section
    # spAfter=2.5pt sesuai template trailing para di akhir references
    _embed_sectpr(
        doc,
        _build_body_sectpr(),
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        sa=2.5,
        li=17.7,
        fi=-17.7,
    )

    # Final sectPr dari template sudah tersisa di body → 1-kolom penutup

    doc.save(str(final_output))
    print(f"[OK] Generated: {final_output}")
    return final_output


# ═══════════════════════════════════════════════════════════════════
# § 15  CLI entry point
# ═══════════════════════════════════════════════════════════════════


def main():
    base_dir = Path(__file__).parent

    if len(sys.argv) >= 2:
        # Mode single file
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH

        if not json_arg.exists():
            print(f"[ERR] File tidak ditemukan: {json_arg}")
            sys.exit(1)

        result = build_document(json_arg, output_arg, template_arg)
        print(f"[OK] Selesai: {result}")

    else:
        # Mode batch: semua JSON di folder
        print(f"Generating JRC DOCX untuk semua JSON di {base_dir}...")
        json_files = sorted(base_dir.glob("*.json"))
        if not json_files:
            print("[ERR] Tidak ada file JSON ditemukan.")
            return

        ok = err = 0
        for jf in json_files:
            if jf.name.lower() in {"package.json", "tsconfig.json", "settings.json"}:
                continue
            try:
                data = json.loads(jf.read_text(encoding="utf-8"))
                if not isinstance(data, dict):
                    print(f"[ERR] Skip: {jf.name} (bukan dict JSON)")
                    continue
                build_document(jf)
                ok += 1
            except json.JSONDecodeError as e:
                print(f"[ERR] Skip: {jf.name} — JSON error: {e}")
                err += 1
            except Exception as e:
                print(f"[ERR] Error: {jf.name} — {e}")
                err += 1

        print(f"\nSelesai: {ok} berhasil, {err} gagal.")


if __name__ == "__main__":
    main()
