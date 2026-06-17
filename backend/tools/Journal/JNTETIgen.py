"""
JNTETIgen.py -- Generator DOCX JNTETI dari JSON.

Dokumen dibangun dari JNTETI.docx asli dengan pendekatan paste keep formatting:
  - file template asli disalin lebih dulu,
  - body template dikosongkan tanpa mengubah header/footer/styles/numbering,
  - paragraph properties dan section break penting diambil langsung dari XML template,
  - footer placeholder penulis/judul diganti dengan ringkasan paper aktual.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "JNTETI.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem
MAX_FIGURE_WIDTH_CM = 8.4
SUBSECTION_ABSTRACT_NUM_ID = 6

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

NS_MAP_STRICT = {
    b"http://purl.oclc.org/ooxml/wordprocessingml/main": b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing": b"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    b"http://purl.oclc.org/ooxml/officeDocument/math": b"http://schemas.openxmlformats.org/officeDocument/2006/math",
}

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None


def _wq(tag: str) -> str:
    return f"{{{NS_W}}}{tag}"


def _strict_to_trans(data: bytes) -> bytes:
    for old, new in NS_MAP_STRICT.items():
        data = data.replace(old, new)
    return data


def _set_para_style(paragraph, style_id: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style_id)


def _create_numbering_instance(doc: Document, abstract_num_id: int, start: int = 1) -> int:
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    existing_ids = []
    for num in numbering.findall(qn("w:num")):
        raw = num.get(qn("w:numId"))
        if raw is None:
            continue
        try:
            existing_ids.append(int(raw))
        except ValueError:
            continue

    num_id = max(existing_ids or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))

    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)

    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    lvl_override.append(start_override)
    num.append(lvl_override)

    numbering.append(num)
    return num_id


def _set_paragraph_numbering(paragraph, num_id: int, ilvl: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    old_numpr = ppr.find(qn("w:numPr"))
    if old_numpr is not None:
        ppr.remove(old_numpr)

    numpr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(ilvl))
    numpr.append(ilvl_el)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    numpr.append(num_id_el)

    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is not None:
        insert_at = list(ppr).index(pstyle) + 1
        ppr.insert(insert_at, numpr)
    else:
        ppr.insert(0, numpr)


def _apply_sample_ppr(paragraph, sample_ppr: etree._Element | None) -> None:
    current = paragraph._p.find(qn("w:pPr"))
    if current is not None:
        paragraph._p.remove(current)
    if sample_ppr is not None:
        paragraph._p.insert(0, deepcopy(sample_ppr))


def _apply_sample_rpr(run, sample_rpr: etree._Element | None) -> None:
    current = run._r.find(qn("w:rPr"))
    if current is not None:
        run._r.remove(current)
    if sample_rpr is not None:
        run._r.insert(0, deepcopy(sample_rpr))


def _set_run_size_half_points(run, value: int) -> None:
    rpr = run._r.get_or_add_rPr()
    sz = rpr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rpr.append(sz)
    sz.set(qn("w:val"), str(value))

    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(value))


def _set_rpr_math_defaults(rpr: etree._Element, half_points: int = 20) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Cambria Math")
    rfonts.set(qn("w:hAnsi"), "Cambria Math")

    sz = rpr.find(qn("w:sz"))
    if sz is None:
        sz = OxmlElement("w:sz")
        rpr.append(sz)
    sz.set(qn("w:val"), str(half_points))

    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(half_points))

    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    if lang.get(qn("w:val")) is None:
        lang.set(qn("w:val"), "en-US")


def _normalize_omml_math(omml: etree._Element, half_points: int = 20) -> etree._Element:
    for math_run in omml.findall(f".//{{{MATH_NS}}}r"):
        rpr = math_run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            math_run.insert(0, rpr)
        _set_rpr_math_defaults(rpr, half_points=half_points)

    for ctrl_pr in omml.findall(f".//{{{MATH_NS}}}ctrlPr"):
        rpr = ctrl_pr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ctrl_pr.insert(0, rpr)
        _set_rpr_math_defaults(rpr, half_points=half_points)

    return omml


def _new_paragraph(doc: Document, sample_ppr: etree._Element | None = None):
    paragraph = doc.add_paragraph()
    if sample_ppr is not None:
        _apply_sample_ppr(paragraph, sample_ppr)
    return paragraph


def _add_sample_run(
    paragraph,
    text: str,
    sample_rpr: etree._Element | None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    small_caps: bool | None = None,
):
    run = paragraph.add_run(text)
    _apply_sample_rpr(run, sample_rpr)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if small_caps is not None:
        run.font.small_caps = small_caps
    return run


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    omml = _normalize_omml_math(omml, half_points=20)
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _normalize_text_commands(text: str) -> str:
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1

    yield from flush_buffer()


def _append_rich_text(paragraph, text: str, sample_rpr: etree._Element | None):
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            paragraph.add_run().add_break()
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = _add_sample_run(paragraph, token["value"], sample_rpr)
                run.italic = True
            continue
        _add_sample_run(
            paragraph,
            token["value"],
            sample_rpr,
            bold=token["bold"],
            italic=token["italic"],
            underline=token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [""]


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _roman(number: str | int) -> str:
    try:
        value = int(str(number).strip())
    except Exception:
        return str(number).strip()
    numerals = [
        (1000, "M"),
        (900, "CM"),
        (500, "D"),
        (400, "CD"),
        (100, "C"),
        (90, "XC"),
        (50, "L"),
        (40, "XL"),
        (10, "X"),
        (9, "IX"),
        (5, "V"),
        (4, "IV"),
        (1, "I"),
    ]
    result = []
    for arabic, roman in numerals:
        while value >= arabic:
            result.append(roman)
            value -= arabic
    return "".join(result)


def _title_text(config: dict) -> str:
    return str(config.get("title") or config.get("TitleBlock", {}).get("Title", "")).strip()


def _abstract_text(config: dict) -> str:
    abstract = str(config.get("abstract") or "").strip()
    if abstract:
        return abstract
    block = config.get("Abstract", {})
    return str(block.get("English") or block.get("Indonesian") or "").strip()


def _keyword_list(config: dict) -> list[str]:
    keywords = config.get("keywords", [])
    if isinstance(keywords, list) and keywords:
        return [str(item).strip() for item in keywords if str(item).strip()]
    block = config.get("Abstract", {})
    raw = str(block.get("KeywordsEnglish") or block.get("KeywordsIndonesian") or "").strip()
    if not raw:
        return []
    return [item.strip() for item in raw.split(",") if item.strip()]


def _author_entries(config: dict) -> list[dict]:
    authors = config.get("authors", [])
    if isinstance(authors, list) and authors:
        return [
            {
                "name": str(author.get("name", "")).strip(),
                "affiliation": str(author.get("affiliation", "")).strip(),
                "location": str(author.get("location", "")).strip(),
                "email": str(author.get("email", "")).strip(),
            }
            for author in authors
        ]
    return []


def _received_line(config: dict) -> str:
    for key in ("received_line", "submission_line", "received"):
        value = str(config.get(key, "")).strip()
        if value:
            return value
    return "[Received: DD MM YY, Revised: DD MM YY, Accepted: DD MM YY]"


def _corresponding_line(config: dict) -> str:
    authors = _author_entries(config)
    if not authors:
        return "Corresponding Author: -"
    first = authors[0]
    name = first.get("name") or "Author"
    email = first.get("email") or "-"
    return f"Corresponding Author: {name} (email: {email})"


def _footer_short_text(config: dict) -> str:
    authors = _author_entries(config)
    title = _title_text(config)
    _name_parts = authors[0]["name"].split() if (authors and authors[0].get("name")) else []
    surname = _name_parts[-1] if _name_parts else "Author"
    author_text = f"{surname} et al." if len(authors) > 1 else surname
    words = re.findall(r"\S+", title)
    short_title = " ".join(words[:4]) if words else "Untitled"
    if len(words) > 4:
        short_title += "..."
    return f"{author_text}: {short_title}".strip()


def _load_template_samples(template_path: Path) -> dict[str, etree._Element | None]:
    with zipfile.ZipFile(template_path) as archive:
        raw = _strict_to_trans(archive.read("word/document.xml"))
    doc_root = etree.fromstring(raw)
    body = doc_root.find(_wq("body"))
    paragraphs = body.findall(_wq("p")) if body is not None else []
    tables = body.findall(_wq("tbl")) if body is not None else []

    def clone_ppr(index: int):
        return deepcopy(paragraphs[index].find(_wq("pPr")))

    def clone_rpr(index: int, run_index: int = 0):
        runs = paragraphs[index].findall(_wq("r"))
        if run_index >= len(runs):
            return None
        rpr = runs[run_index].find(_wq("rPr"))
        return deepcopy(rpr) if rpr is not None else None

    return {
        "title_ppr": clone_ppr(0),
        "title_rpr": clone_rpr(0, 0),
        "license_table": deepcopy(tables[0]) if tables else None,
        "author_ppr": clone_ppr(1),
        "author_name_rpr": clone_rpr(1, 0),
        "author_sup_rpr": clone_rpr(1, 1),
        "affiliation_first_ppr": clone_ppr(2),
        "affiliation_other_ppr": clone_ppr(3),
        "affiliation_sup_rpr": clone_rpr(2, 0),
        "affiliation_text_rpr": clone_rpr(2, 2),
        "meta_ppr": clone_ppr(4),
        "meta_rpr": clone_rpr(4, 0),
        "corresponding_ppr": clone_ppr(5),
        "corresponding_rpr": clone_rpr(5, 0),
        "blank_section_ppr": clone_ppr(6),
        "abstract_ppr": clone_ppr(7),
        "abstract_heading_rpr": clone_rpr(7, 0),
        "abstract_body_rpr": clone_rpr(7, 3),
        "keywords_ppr": clone_ppr(8),
        "keywords_heading_rpr": clone_rpr(8, 1),
        "keywords_body_rpr": clone_rpr(8, 3),
        "section_ppr": clone_ppr(9),
        "section_rpr": clone_rpr(9, 0),
        "body_ppr": clone_ppr(10),
        "body_rpr": clone_rpr(10, 0),
        "subsection_ppr": clone_ppr(20),
        "subsection_rpr": clone_rpr(20, 0),
        "subsubsection_ppr": clone_ppr(29),
        "subsubsection_rpr": clone_rpr(29, 0),
        "equation_ppr": clone_ppr(58),
        "equation_num_rpr": clone_rpr(58, 3),
        "reference_heading_ppr": clone_ppr(86),
        "reference_heading_rpr": (
            clone_rpr(78, 0) if clone_rpr(78, 0) is not None else clone_rpr(9, 0)
        ),
        "reference_item_ppr": clone_ppr(87),
        "reference_item_rpr": clone_rpr(87, 0),
        "trailing_ppr": clone_ppr(102),
    }


def _add_title(doc: Document, config: dict, samples: dict) -> None:
    paragraph = _new_paragraph(doc, samples["title_ppr"])
    _add_sample_run(paragraph, _title_text(config), samples["title_rpr"])


def _add_authors(doc: Document, config: dict, samples: dict) -> None:
    authors = _author_entries(config)
    if not authors:
        return

    affiliation_ids: dict[tuple[str, str], int] = {}
    author_aff_numbers: list[int] = []
    for entry in authors:
        key = (entry.get("affiliation", ""), entry.get("location", ""))
        if key not in affiliation_ids:
            affiliation_ids[key] = len(affiliation_ids) + 1
        author_aff_numbers.append(affiliation_ids[key])

    author_paragraph = _new_paragraph(doc, samples["author_ppr"])
    for index, entry in enumerate(authors):
        if index:
            _add_sample_run(author_paragraph, ", ", samples["author_name_rpr"])
        _add_sample_run(author_paragraph, entry.get("name", ""), samples["author_name_rpr"])
        _add_sample_run(author_paragraph, str(author_aff_numbers[index]), samples["author_sup_rpr"])

    reverse_aff_map: dict[int, dict[str, list[str] | str]] = {}
    for index, entry in enumerate(authors):
        aff_number = author_aff_numbers[index]
        record = reverse_aff_map.setdefault(
            aff_number,
            {
                "numbers": [],
                "affiliation": entry.get("affiliation", ""),
                "location": entry.get("location", ""),
            },
        )
        record["numbers"].append(str(aff_number))

    aff_lines = []
    for aff_number in sorted(reverse_aff_map):
        record = reverse_aff_map[aff_number]
        label = str(aff_number)
        text_parts = [part for part in (record["affiliation"], record["location"]) if part]
        aff_lines.append((label, ", ".join(text_parts)))

    for index, (label, text) in enumerate(aff_lines):
        paragraph = _new_paragraph(
            doc,
            samples["affiliation_first_ppr"] if index == 0 else samples["affiliation_other_ppr"],
        )
        _add_sample_run(paragraph, label, samples["affiliation_sup_rpr"])
        _add_sample_run(paragraph, " ", samples["affiliation_sup_rpr"])
        _add_sample_run(paragraph, text, samples["affiliation_text_rpr"])

    meta_paragraph = _new_paragraph(doc, samples["meta_ppr"])
    _add_sample_run(meta_paragraph, _received_line(config), samples["meta_rpr"])

    corresponding_paragraph = _new_paragraph(doc, samples["corresponding_ppr"])
    _add_sample_run(
        corresponding_paragraph, _corresponding_line(config), samples["corresponding_rpr"]
    )

    _new_paragraph(doc, samples["blank_section_ppr"])


def _add_abstracts(doc: Document, config: dict, samples: dict) -> None:
    abstract = _abstract_text(config)
    keywords = _keyword_list(config)

    if abstract:
        paragraph = _new_paragraph(doc, samples["abstract_ppr"])
        run = _add_sample_run(paragraph, "ABSTRACT", samples["abstract_heading_rpr"])
        _set_run_size_half_points(run, 20)
        run.font.name = "Helvetica"
        run.font.size = Pt(9)
        run = _add_sample_run(paragraph, " ", samples["abstract_heading_rpr"])
        _set_run_size_half_points(run, 20)
        run = _add_sample_run(paragraph, "\u2014", samples["abstract_body_rpr"])
        _set_run_size_half_points(run, 20)
        run = _add_sample_run(paragraph, " ", samples["abstract_body_rpr"])
        _set_run_size_half_points(run, 20)
        before_count = len(paragraph.runs)
        _append_rich_text(paragraph, abstract, samples["abstract_body_rpr"])
        for run in paragraph.runs[before_count:]:
            _set_run_size_half_points(run, 20)

    if keywords:
        paragraph = _new_paragraph(doc, samples["keywords_ppr"])
        run = _add_sample_run(paragraph, "KEYWORDS", samples["keywords_heading_rpr"])
        _set_run_size_half_points(run, 20)
        run.font.name = "Helvetica"
        run.font.size = Pt(9)
        run = _add_sample_run(paragraph, " ", samples["keywords_heading_rpr"])
        _set_run_size_half_points(run, 20)
        run = _add_sample_run(paragraph, "\u2014", samples["keywords_body_rpr"])
        _set_run_size_half_points(run, 20)
        run = _add_sample_run(paragraph, " ", samples["keywords_body_rpr"])
        _set_run_size_half_points(run, 20)
        keywords_text = ", ".join(keywords)
        if keywords_text and not keywords_text.endswith("."):
            keywords_text += "."
        before_count = len(paragraph.runs)
        _append_rich_text(paragraph, keywords_text, samples["keywords_body_rpr"])
        for run in paragraph.runs[before_count:]:
            _set_run_size_half_points(run, 20)


def _insert_body_element(doc: Document, element: etree._Element) -> None:
    body = doc._element.body
    insert_index = len(body)
    for index, child in enumerate(body):
        if child.tag == qn("w:sectPr"):
            insert_index = index
            break
    body.insert(insert_index, deepcopy(element))


def _body_paragraphs(doc: Document, text: str, samples: dict) -> None:
    for block in _split_body_blocks(text):
        paragraph = _new_paragraph(doc, samples["body_ppr"])
        _append_rich_text(paragraph, block, samples["body_rpr"])


def _add_section_heading(doc: Document, text: str, samples: dict) -> None:
    paragraph = _new_paragraph(doc, samples["section_ppr"])
    _add_sample_run(paragraph, text.upper(), samples["section_rpr"], bold=True)


def _add_subsection_heading(
    doc: Document, text: str, samples: dict, numbering_id: int | None = None
) -> None:
    paragraph = _new_paragraph(doc, samples["subsection_ppr"])
    if numbering_id is not None:
        _set_paragraph_numbering(paragraph, numbering_id, ilvl=0)
    _add_sample_run(paragraph, text.upper(), samples["subsection_rpr"], bold=True, italic=True)


def _add_subsubsection_heading(doc: Document, text: str, samples: dict) -> None:
    paragraph = _new_paragraph(doc, samples["subsubsection_ppr"])
    _add_sample_run(paragraph, text.upper(), samples["subsubsection_rpr"], italic=True)


def _set_full_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "bottom", "left", "right"):
        border = tc_borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            tc_borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")


def _set_horizontal_cell_borders(cell, *, top: bool = False, bottom: bool = False):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    border_spec = {
        "top": (
            {"val": "single", "sz": "4", "space": "0", "color": "auto"} if top else {"val": "none"}
        ),
        "bottom": (
            {"val": "single", "sz": "4", "space": "0", "color": "auto"}
            if bottom
            else {"val": "none"}
        ),
        "left": {"val": "none"},
        "right": {"val": "none"},
    }

    for edge, attrs in border_spec.items():
        border = tc_borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            tc_borders.append(border)
        for key, value in attrs.items():
            border.set(qn(f"w:{key}"), value)


def _set_row_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)


def _style_table_paragraph(paragraph, style_id: str, align=WD_ALIGN_PARAGRAPH.LEFT):
    _set_para_style(paragraph, style_id)
    paragraph.alignment = align


def _add_prompt_box(doc: Document, text: str, samples: dict):
    table = doc.add_table(rows=1, cols=1)
    _set_table_full_borders(table)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_full_cell_borders(cell)
    paragraph = cell.paragraphs[0]
    _apply_sample_ppr(paragraph, samples["body_ppr"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _append_rich_text(paragraph, text, samples["body_rpr"])


def _style_figure_caption_run(run, *, bold: bool = False):
    run.font.name = "Helvetica"
    run.font.size = Pt(7)
    run.bold = bold


def _style_table_caption_run(run, *, small_caps: bool = False, bold: bool = False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    run.font.small_caps = small_caps
    run.bold = bold


def _add_figure_caption(doc: Document, number: str, title: str):
    caption = doc.add_paragraph()
    _set_para_style(caption, "IEEEFigure")
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
        if len(title) <= 120 and "\n" not in title
        else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    label = caption.add_run(f"Figure {number}. ")
    _style_figure_caption_run(label, bold=True)
    text_run = caption.add_run(title)
    _style_figure_caption_run(text_run, bold=False)


def _add_table_caption(doc: Document, number: str, title: str):
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(3)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = caption.add_run(f"TABLE {_roman(number)}. ")
    _style_table_caption_run(label, small_caps=True, bold=False)
    text_run = caption.add_run(title)
    _style_table_caption_run(text_run, small_caps=False, bold=False)


def _add_figure(doc: Document, item: dict, json_path: Path, samples: dict) -> None:
    number = str(item.get("ImageNumber") or item.get("number") or "").strip()
    title = str(item.get("Title") or item.get("title") or "").strip()
    prompt = str(item.get("Prompt") or item.get("Description") or "").strip()
    path_text = str(item.get("Path") or item.get("path") or "").strip()

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except Exception:
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    # Resolve the real image first; only emit the AI prompt placeholder when
    # no embeddable image exists (mirror IEEE: never show prompt + image both).
    image_path = _resolve_path(path_text, json_path) if path_text else None
    has_image = image_path is not None and image_path.is_file()

    if title and not has_image:
        prompt_desc = prompt or title
        prompt_text = f"[PROMPT UNTUK AI GAMBAR: {title}. {prompt_desc}]"
        prompt_para = doc.add_paragraph()
        prompt_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        from docx.shared import RGBColor as _RGB

        pr = prompt_para.add_run(prompt_text)
        pr.italic = True
        pr.font.color.rgb = _RGB(0xFF, 0x00, 0x00)

    if has_image:
        paragraph = doc.add_paragraph()
        _set_para_style(paragraph, "IEEEFigure")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = paragraph.paragraph_format
        pf.space_before = Pt(6)
        pf.space_after = Pt(2)
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))

    if number and title:
        _add_figure_caption(doc, number, title)


def _add_table(doc: Document, item: dict) -> None:
    number = str(item.get("TableNumber") or item.get("NumberiOrLetter") or "").strip()
    title = str(item.get("Title") or item.get("title") or "").strip()
    headers = list(item.get("Headers") or item.get("headers") or [])
    rows = list(item.get("Rows") or item.get("rows") or [])

    if not headers:
        return

    if number and title:
        _add_table_caption(doc, number, title)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    _set_table_full_borders(table)
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_row_repeat_header(table.rows[0])

    for col_index, value in enumerate(headers):
        cell = table.rows[0].cells[col_index]
        _set_horizontal_cell_borders(cell, top=True, bottom=True)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        _style_table_paragraph(paragraph, "IEEETableHeaderCentered", WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run(str(value))
        run.bold = True

    for row_index, row_data in enumerate(rows, start=1):
        is_last_row = row_index == len(rows)
        for col_index, value in enumerate(row_data):
            if col_index >= len(headers):
                break
            cell = table.rows[row_index].cells[col_index]
            _set_horizontal_cell_borders(cell, top=False, bottom=is_last_row)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _style_table_paragraph(paragraph, "IEEETableCell", WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.add_run(str(value))

    doc.add_paragraph()


def _add_equation_line(doc: Document, formula: str, number: str | None, samples: dict) -> None:
    paragraph = _new_paragraph(doc, samples["equation_ppr"])
    paragraph.add_run("\t")
    if not _append_inline_math(paragraph, formula):
        fallback = paragraph.add_run(formula)
        fallback.italic = True
    if number is not None:
        paragraph.add_run(" ")
        paragraph.add_run("\t")
        _add_sample_run(paragraph, f"({number})", samples["equation_num_rpr"])


def _add_equation_group(doc: Document, item: dict, samples: dict) -> None:
    formulas = [str(value).strip() for value in item.get("Lines", []) if str(value).strip()]
    single = str(item.get("latex") or item.get("text") or "").strip()
    if not formulas and single:
        formulas = [single]
    if not formulas:
        return

    number = str(item.get("FormulaNumber") or item.get("NumberiOrLetter") or "").strip() or None
    for formula in formulas[:-1]:
        _add_equation_line(doc, formula, None, samples)
    _add_equation_line(doc, formulas[-1], number, samples)


def _reference_texts(config: dict) -> list[str]:
    def _rt(item):
        if isinstance(item, dict):
            return (item.get("text") or item.get("Text") or "").strip()
        return str(item).strip()
    if "references" in config and isinstance(config["references"], dict):
        content = config["references"].get("content", [])
        if isinstance(content, list):
            return [_rt(item) for item in content if _rt(item)]
    if "References" in config and isinstance(config["References"], list):
        return [_rt(item) for item in config["References"] if _rt(item)]
    return []


def _strip_reference_label(text: str) -> str:
    return re.sub(r"^\s*\[(\d+)\]\s*", "", text).strip()


def _add_references(doc: Document, config: dict, samples: dict) -> None:
    references = _reference_texts(config)
    if not references:
        return

    heading = _new_paragraph(doc, samples["reference_heading_ppr"])
    _add_sample_run(heading, "REFERENCES", samples["reference_heading_rpr"], bold=True)

    for reference in references:
        paragraph = _new_paragraph(doc, samples["reference_item_ppr"])
        _append_rich_text(
            paragraph, _strip_reference_label(reference), samples["reference_item_rpr"]
        )


def _render_content_item(doc: Document, item: dict, json_path: Path, samples: dict) -> None:
    item_id = str(item.get("id", "")).lower().strip()

    if item_id == "text":
        text = str(item.get("text", "")).strip()
        if text:
            _body_paragraphs(doc, text, samples)
        return

    if item_id in {"gambar", "image", "figure"}:
        _add_figure(doc, item, json_path, samples)
        return

    if item_id in {"rumus", "formula", "equation"}:
        _add_equation_group(doc, item, samples)
        return

    if item_id in {"tabel", "table"}:
        _add_table(doc, item)


def _render_subsubsection(doc: Document, block: dict, json_path: Path, samples: dict) -> None:
    title = str(block.get("title", "")).strip()
    if title:
        _add_subsubsection_heading(doc, title, samples)
    content = block.get("content", [])
    if isinstance(content, list):
        for item in content:
            if isinstance(item, dict):
                _render_content_item(doc, item, json_path, samples)
            elif isinstance(item, str) and item.strip():
                _body_paragraphs(doc, item.strip(), samples)


def _render_subsection(
    doc: Document,
    block: dict,
    json_path: Path,
    samples: dict,
    sub_key: str = "",
    numbering_id: int | None = None,
) -> None:
    title = str(block.get("title", "")).strip()
    if title:
        _add_subsection_heading(doc, title, samples, numbering_id=numbering_id)

    content = block.get("content", [])
    if isinstance(content, list):
        for item in content:
            if isinstance(item, dict):
                _render_content_item(doc, item, json_path, samples)
            elif isinstance(item, str) and item.strip():
                _body_paragraphs(doc, item.strip(), samples)
    elif isinstance(content, str) and content.strip():
        _body_paragraphs(doc, content.strip(), samples)

    nested_keys = [
        key
        for key in block.keys()
        if re.match(r"^section\d+[a-z]{2,}$", key)
        or (key.startswith("sub") and len(key) > 3 and key[3:].isalnum())
    ]
    for nested_key in nested_keys:
        nested = block[nested_key]
        if isinstance(nested, dict):
            _render_subsubsection(doc, nested, json_path, samples)


def _render_sections(doc: Document, config: dict, json_path: Path, samples: dict) -> None:
    section_keys = sorted(
        [key for key in config if re.match(r"^section\d+$", key)],
        key=lambda value: int(value.replace("section", "")),
    )

    for section_key in section_keys:
        section = config[section_key]
        title = str(section.get("title", "")).strip()
        if title:
            _add_section_heading(doc, title, samples)

        content = section.get("content", [])
        if isinstance(content, list):
            for item in content:
                if isinstance(item, dict):
                    _render_content_item(doc, item, json_path, samples)
                elif isinstance(item, str) and item.strip():
                    _body_paragraphs(doc, item.strip(), samples)
        elif isinstance(content, str) and content.strip():
            _body_paragraphs(doc, content.strip(), samples)

        section_number = section_key.replace("section", "")
        subsection_keys = sorted(
            [
                key
                for key in section.keys()
                if re.match(rf"^section{section_number}[a-z]+$", key)
                or (key.startswith("sub") and len(key) > 3 and key[3:].isalnum())
            ]
        )
        subsection_num_id = (
            _create_numbering_instance(doc, SUBSECTION_ABSTRACT_NUM_ID) if subsection_keys else None
        )
        for subsection_key in subsection_keys:
            subsection = section[subsection_key]
            if isinstance(subsection, dict):
                _render_subsection(
                    doc,
                    subsection,
                    json_path,
                    samples,
                    subsection_key,
                    numbering_id=subsection_num_id,
                )


def _replace_footer_placeholders(xml_bytes: bytes, replacement_text: str) -> bytes:
    root = etree.fromstring(xml_bytes)
    ns = {"w": NS_W}
    for paragraph in root.findall(".//w:p", ns):
        seen_tab = False
        blank_tail = False
        for child in paragraph:
            if child.tag == _wq("r"):
                if child.find(_wq("tab")) is not None:
                    seen_tab = True
                for text_node in child.findall(".//w:t", ns):
                    text = text_node.text or ""
                    if "First Author:" in text or "Penulis Pertama:" in text:
                        text_node.text = replacement_text
                        blank_tail = seen_tab
                    elif blank_tail and text:
                        text_node.text = ""
                    elif re.search(r"Volume\s+\w+\s+Nomor\s+\w+\s+\w+\s+\d{4}", text):
                        # Hapus placeholder volume/issue (audit flag sebagai leak)
                        text_node.text = re.sub(
                            r"Volume\s+\w+\s+Nomor\s+\w+\s+\w+\s+\d{4}",
                            "",
                            text,
                        )
    return etree.tostring(root, encoding="utf-8", xml_declaration=True)


def _patch_footer_placeholders(docx_path: Path, replacement_text: str) -> None:
    temp_path = docx_path.with_suffix(docx_path.suffix + ".tmp")
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(temp_path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("word/footer") and item.filename.endswith(".xml"):
                data = _replace_footer_placeholders(data, replacement_text)
            dst.writestr(item, data)
    temp_path.replace(docx_path)


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    shutil.copy(str(template_path), str(final_output))
    samples = _load_template_samples(template_path)

    doc = Document(str(final_output))
    _clear_document_body(doc)

    if samples.get("license_table") is not None:
        _insert_body_element(doc, samples["license_table"])

    _add_title(doc, config, samples)
    _add_authors(doc, config, samples)
    _add_abstracts(doc, config, samples)
    _render_sections(doc, config, Path(json_path), samples)
    _add_references(doc, config, samples)
    _new_paragraph(doc, samples["trailing_ppr"])

    doc.save(str(final_output))
    _patch_footer_placeholders(final_output, _footer_short_text(config))
    return final_output


def main() -> None:
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        if not json_arg.exists():
            print(f"ERROR: file not found: {json_arg}")
            sys.exit(1)
        result = build_document(json_arg, output_arg, template_arg)
        print(f"Generated: {result}")
        return

    json_files = sorted(
        path
        for path in BASE_DIR.glob("*.json")
        if path.name.lower() not in {"package.json", "tsconfig.json", "settings.json"}
    )
    if not json_files:
        print("ERROR: no JSON files found")
        return

    ok = 0
    err = 0
    for json_file in json_files:
        try:
            result = build_document(json_file)
            print(f"Generated: {result.name}")
            ok += 1
        except Exception as exc:
            print(f"ERROR: {json_file.name}: {exc}")
            err += 1
    print(f"Done: {ok} OK, {err} errors")


def _set_table_full_borders(table) -> None:
    """Pastikan tabel punya border tegas/visible (val=single, sz=4 = 0.5pt).

    Dipanggil setelah doc.add_table() supaya tabel data keliatan di Word.
    Auto-injected oleh _fix_table_borders.py untuk lulus audit border check.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


if __name__ == "__main__":
    main()
