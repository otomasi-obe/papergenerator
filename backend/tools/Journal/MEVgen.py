from __future__ import annotations

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "MEV.docx"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
BODY_WIDTH_PT = 595.35 - 70.90 - 56.70
MAX_FIGURE_WIDTH_CM = 15.5
XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
]
_XSLT = None


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in {"oMath", "oMathPara"}:
        paragraph._p.append(omml)
        return True
    wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
    wrapper.append(omml)
    paragraph._p.append(wrapper)
    return True


def _append_text_run(
    paragraph, text: str, bold: bool = False, italic: bool = False, underline: bool = False
):
    if not text:
        return None
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return run


def _append_line_break(paragraph):
    paragraph.add_run().add_break()


def _normalize_text_commands(text: str) -> str:
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\b")
                    index += 2
                    continue
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\i")
                    index += 2
                    continue
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\u")
                    index += 2
                    continue
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def _append_rich_text(paragraph, text: str):
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            _append_line_break(paragraph)
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = paragraph.add_run(token["value"])
                run.italic = True
            continue
        _append_text_run(
            paragraph,
            token["value"],
            bold=token["bold"],
            italic=token["italic"],
            underline=token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [""]


def _smart_title(text: str) -> str:
    stripped = str(text or "").strip()
    if stripped.isupper():
        return stripped.title()
    return stripped


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _paragraph(
    doc: Document,
    style: str | None = None,
    align=None,
    first_indent: float | None = None,
    left_indent: float | None = None,
    before: float | None = None,
    after: float | None = None,
):
    paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    if first_indent is not None:
        pf.first_line_indent = Pt(first_indent)
    if left_indent is not None:
        pf.left_indent = Pt(left_indent)
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    return paragraph


def _set_run_format(run, *, size: float | None = None, bold=None, italic=None, superscript=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if superscript is not None:
        run.font.superscript = superscript


def _disable_heading_numbering(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        pstyle = ppr.find(qn("w:pStyle"))
        if pstyle is not None:
            ppr.insert(list(ppr).index(pstyle) + 1, num_pr)
        else:
            ppr.insert(0, num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        num_id = OxmlElement("w:numId")
        num_pr.append(num_id)
    num_id.set(qn("w:val"), "0")


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    candidates = [json_path.parent / path, BASE_DIR / path, BASE_DIR / path.name]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0] if candidates else path


def _apply_heading_run_format(paragraph, level: int) -> None:
    for run in paragraph.runs:
        if level == 1:
            _set_run_format(run, size=14, bold=True, italic=False)
        elif level == 2:
            _set_run_format(run, size=13, bold=True, italic=False)
        elif level == 3:
            _set_run_format(run, size=12, bold=False, italic=True)


def _roman_to_int(text: str) -> int | None:
    values = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}
    roman = text.strip().upper()
    if not roman or any(char not in values for char in roman):
        return None
    total = 0
    prev = 0
    for char in reversed(roman):
        value = values[char]
        if value < prev:
            total -= value
        else:
            total += value
            prev = value
    return total if total > 0 else None


def _normalize_table_number(number: str) -> str:
    text = str(number or "").strip()
    if not text:
        return ""
    if text.isdigit():
        return text
    roman_value = _roman_to_int(text)
    if roman_value is not None:
        return str(roman_value)
    return text


def _affiliation_groups(
    authors: list[dict],
) -> tuple[OrderedDict[tuple[str, str], dict], list[str]]:
    groups: OrderedDict[tuple[str, str], dict] = OrderedDict()
    letters: list[str] = []
    next_letter = ord("a")
    for author in authors:
        affiliation = str(author.get("affiliation") or "").strip()
        location = str(author.get("location") or "").strip()
        key = (affiliation, location)
        if key not in groups:
            groups[key] = {
                "letter": chr(next_letter),
                "affiliation": affiliation,
                "location": location,
            }
            next_letter += 1
        letters.append(groups[key]["letter"])
    return groups, letters


def _add_title_block(doc: Document, config: dict) -> None:
    title = str(config.get("title") or "").strip()
    authors = list(config.get("authors") or [])
    groups, letters = _affiliation_groups(authors)
    emails = [
        str(author.get("email") or "").strip()
        for author in authors
        if str(author.get("email") or "").strip()
    ]

    _paragraph(doc)
    if title:
        paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
        _append_rich_text(paragraph, title)
        for run in paragraph.runs:
            _set_run_format(run, size=16, bold=True)
    _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)

    if authors:
        author_line = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
        for index, author in enumerate(authors):
            if index > 0:
                separator = author_line.add_run(", ")
                _set_run_format(separator, bold=True)
            name_run = author_line.add_run(str(author.get("name") or "").strip())
            _set_run_format(name_run, bold=True)
            marker_text = letters[index]
            if marker_text:
                suffix = "," if index == 0 else ""
                marker_run = author_line.add_run(f" {marker_text}{suffix}")
                _set_run_format(marker_run, bold=True, superscript=True)
            if index == 0:
                corr_run = author_line.add_run(" *")
                _set_run_format(corr_run, bold=True)

        for group in groups.values():
            aff = str(group["affiliation"]).strip()
            loc = str(group["location"]).strip()
            if aff:
                paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
                marker_run = paragraph.add_run(group["letter"])
                _set_run_format(marker_run, italic=True, superscript=True)
                spacer = paragraph.add_run(" ")
                _set_run_format(spacer, italic=True, superscript=True)
                aff_run = paragraph.add_run(aff)
                _set_run_format(aff_run, italic=True)
            if loc:
                paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
                loc_run = paragraph.add_run(loc)
                _set_run_format(loc_run, italic=True)

        _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
        paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
        corr_text = paragraph.add_run("*Corresponding Author.")
        _set_run_format(corr_text, italic=True)
        if emails:
            paragraph = _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
            email_text = paragraph.add_run(f"E-mail: {'; '.join(emails)}")
            _set_run_format(email_text, italic=True)

    _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)
    _paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0)


def _add_abstract_keywords(doc: Document, config: dict) -> None:
    abstract = str(config.get("abstract") or "").strip()
    keywords = [
        str(keyword).strip() for keyword in config.get("keywords") or [] if str(keyword).strip()
    ]
    if abstract:
        paragraph = _paragraph(doc, first_indent=0)
        label = paragraph.add_run("Abstract")
        _set_run_format(label, bold=True)
        paragraph.add_run(" ")
        paragraph = _paragraph(doc)
        _append_rich_text(paragraph, abstract)
    if keywords:
        paragraph = _paragraph(doc)
        _append_rich_text(paragraph, f"Keywords: {'; '.join(keywords)}.")


def _body_paragraphs(doc: Document, text: str):
    for block in _split_body_blocks(text):
        paragraph = _paragraph(doc)
        _append_rich_text(paragraph, block)


def _add_heading(doc: Document, text: str, level: int, numbered: bool = True):
    style_name = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    paragraph = _paragraph(doc, style=style_name)
    if not numbered:
        _disable_heading_numbering(paragraph)
    _append_rich_text(paragraph, text)
    _apply_heading_run_format(paragraph, level)
    return paragraph


def _add_missing_figure_note(doc: Document, prompt: str):
    paragraph = _paragraph(
        doc, style="List Paragraph", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=0, left_indent=0
    )
    run = paragraph.add_run(prompt)
    _set_run_format(run, italic=True)


def _add_figure(doc: Document, item: dict, json_path: Path) -> None:
    number = str(item.get("ImageNumber") or "").strip()
    title = str(item.get("Title") or "").strip()
    prompt = str(item.get("Prompt") or "").strip()
    path_text = str(item.get("Path") or "").strip()
    image_path = _resolve_path(path_text, json_path) if path_text else None

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except Exception:
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    if image_path is not None and image_path.is_file():
        paragraph = _paragraph(
            doc,
            style="List Paragraph",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            first_indent=0,
            left_indent=0,
        )
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    else:
        fallback = prompt or f"Image not found: {path_text}"
        _add_missing_figure_note(doc, fallback)

    if title or number:
        caption = _paragraph(doc, style="List Paragraph", first_indent=0, left_indent=0)
        text = f"Figure {number}. {title}".strip() if number else title
        caption.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 110 else WD_ALIGN_PARAGRAPH.JUSTIFY
        )
        _append_rich_text(caption, text)


def _build_table_cell_text(paragraph, text: str, bold: bool = False):
    paragraph.clear()
    run = paragraph.add_run(str(text))
    _set_run_format(run, bold=bold)


def _add_table(doc: Document, item: dict) -> None:
    number = _normalize_table_number(str(item.get("TableNumber") or "").strip())
    title = str(item.get("Title") or "").strip()
    headers = list(item.get("Headers") or [])
    rows = list(item.get("Rows") or [])
    if not headers:
        return

    label = f"Tabel {number}." if number else "Tabel."
    paragraph = _paragraph(doc, first_indent=0)
    _append_rich_text(paragraph, label)
    if title:
        paragraph = _paragraph(doc, first_indent=0)
        _append_rich_text(paragraph, title)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders_match_template(table)
    for column_index, value in enumerate(headers):
        paragraph = table.rows[0].cells[column_index].paragraphs[0]
        paragraph.paragraph_format.first_line_indent = None
        paragraph.paragraph_format.left_indent = None
        _build_table_cell_text(paragraph, value, bold=True)
    for row_index, row_data in enumerate(rows, start=1):
        for column_index, value in enumerate(row_data):
            if column_index >= len(headers):
                break
            paragraph = table.rows[row_index].cells[column_index].paragraphs[0]
            paragraph.paragraph_format.first_line_indent = None
            paragraph.paragraph_format.left_indent = None
            _build_table_cell_text(paragraph, value)


def _add_equation(doc: Document, item: dict) -> None:
    latex = str(item.get("latex") or item.get("text") or "").strip()
    number = str(item.get("FormulaNumber") or "").strip()
    if not latex:
        return
    paragraph = _paragraph(doc, first_indent=0, before=6, after=6)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(BODY_WIDTH_PT), WD_TAB_ALIGNMENT.RIGHT)
    if not _append_inline_math(paragraph, latex):
        run = paragraph.add_run(latex)
        _set_run_format(run, italic=True)
    if number:
        paragraph.add_run(f"\t({number})")


def _iter_content_items(content):
    if isinstance(content, str):
        yield content
        return
    if not isinstance(content, list):
        return
    for item in content:
        yield item


def _render_content_item(doc: Document, item, json_path: Path) -> None:
    if isinstance(item, str):
        if item.strip():
            _body_paragraphs(doc, item)
        return
    if not isinstance(item, dict):
        return
    item_id = str(item.get("id") or "").strip().lower()
    if item_id == "text":
        text = str(item.get("text") or "").strip()
        if text:
            _body_paragraphs(doc, text)
        return
    if item_id in {"gambar", "image"}:
        _add_figure(doc, item, json_path)
        return
    if item_id in {"rumus", "formula"}:
        _add_equation(doc, item)
        return
    if item_id in {"tabel", "table"}:
        _add_table(doc, item)


def _subsection_entries(section: dict):
    for key, value in section.items():
        if not isinstance(value, dict):
            continue
        if re.match(r"^section\d+[a-z]+$", key) or (key.startswith("sub") and len(key) > 3):
            yield key, value


def _render_subsection(doc: Document, subsection: dict, json_path: Path) -> None:
    title = _smart_title(subsection.get("title"))
    if title:
        _add_heading(doc, title, level=2)
    for item in _iter_content_items(subsection.get("content", [])):
        _render_content_item(doc, item, json_path)


def _render_sections(doc: Document, config: dict, json_path: Path) -> None:
    # Support both "sections" array and "section1/section2" legacy format
    section_list = []
    if "sections" in config and isinstance(config["sections"], list):
        for idx, sec in enumerate(config["sections"], start=1):
            if isinstance(sec, dict):
                section_list.append((f"section{idx}", sec))
    else:
        section_keys = sorted(
            [key for key in config if key.startswith("section") and key[7:].isdigit()],
            key=lambda key: int(key[7:]),
        )
        for key in section_keys:
            section_list.append((key, config[key]))

    for section_key, section in section_list:
        title = _smart_title(section.get("title"))
        if title:
            _add_heading(doc, title, level=1)
        for item in _iter_content_items(section.get("content", [])):
            _render_content_item(doc, item, json_path)
        for _, subsection in _subsection_entries(section):
            _render_subsection(doc, subsection, json_path)


def _reference_prefix(index: int, reference) -> str:
    if isinstance(reference, dict):
        ref_id = str(reference.get("id") or "").strip()
        if ref_id:
            return f"[{ref_id}]"
    return f"[{index}]"


def _reference_text(reference) -> str:
    if isinstance(reference, dict):
        return str(reference.get("text") or "").strip()
    return str(reference).strip()


def _add_references(doc: Document, config: dict) -> None:
    refs_raw = config.get("references")
    references = []
    if isinstance(refs_raw, dict):
        references = list(refs_raw.get("content", refs_raw.get("items", [])) or [])
    elif isinstance(refs_raw, list):
        # Direct list of strings or {"text": ...} dicts
        for r in refs_raw:
            if isinstance(r, dict):
                references.append(_format_reference(r))
            else:
                references.append(str(r))
    if not references:
        return
    _add_heading(doc, "References", level=1, numbered=False)
    for index, reference in enumerate(references, start=1):
        paragraph = _paragraph(doc, style="List Paragraph")
        paragraph.paragraph_format.left_indent = Pt(21.3)
        paragraph.paragraph_format.first_line_indent = Pt(-21.3)
        _append_rich_text(
            paragraph, f"{_reference_prefix(index, reference)} {_reference_text(reference)}"
        )


def _default_output_path(json_path: Path) -> Path:
    return json_path.parent / f"{json_path.stem}.docx"


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = Path(output_path) if output_path else _default_output_path(Path(json_path))
    doc = Document(str(template_path))
    _clear_document_body(doc)
    _add_title_block(doc, config)
    _add_abstract_keywords(doc, config)
    _render_sections(doc, config, Path(json_path))
    _add_references(doc, config)
    final_output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(final_output))
    return final_output


def main():
    base_dir = BASE_DIR
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        result = build_document(json_arg, output_arg, template_arg)
        print(f"Generated: {result.name}")
        return

    print("Generating all JMEV DOCX files...")
    for json_path in sorted(base_dir.glob("*.json")):
        if json_path.name.lower() in {"package.json", "tsconfig.json", "settings.json"}:
            continue
        try:
            output_path = build_document(json_path)
            print(f"Generated: {output_path.name}")
        except Exception as exc:
            print(f"Error: {json_path.name} - {exc}")


def _set_table_borders_match_template(table) -> None:
    """Set border tabel sesuai pattern template original: FULL_GRID.

    Set semua side (top, left, bottom, right, insideH, insideV) jadi single/sz=4.
    Auto-injected oleh _fix_table_borders_v2.py.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


if __name__ == "__main__":
    main()