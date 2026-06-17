"""
ICIMECEgen.py — Generator DOCX untuk Web of Conferences (ICIMECE) format.

Strategi:
1. Copy ICIMECE.docx asli → ICIMECE_output.docx (preserve styles, headers, footers, sectPr)
2. Hapus body content (paragraf + tabel), pertahankan sectPr final
3. Generate konten dari _template.json menggunakan styles asli template
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(__file__).resolve().parent
TEMPLATE_DOCX = BASE / "ICIMECE.docx"
TEMPLATE_JSON = BASE / "_template.json"
OUTPUT_DOCX = BASE / "ICIMECE_output.docx"

# Style names dari template asli
STYLE_TITLE = "Titre"
STYLE_AUTHOR = "AuthorLastName"
STYLE_AFFIL = "Affiliation"
STYLE_ABSTRACT = "Abstractbody"
STYLE_SECTION = "Section"
STYLE_SUBSECTION = "Subsection"
STYLE_SUBSUBSECTION = "Subsubsection"
STYLE_PARA_FIRST = "Paragraphfirst"
STYLE_PARA = "Paragraph"
STYLE_TBL_CAPTION = "TableCaption"
STYLE_FIG_CAPTION = "StyleFigureCaption"
STYLE_REF_BODY = "ReferencesBody"
STYLE_REF_HEAD = "Section0"


def _strip_latex(text: str) -> str:
    """Strip common LaTeX commands sehingga tidak terdeteksi sebagai LaTeX leak.

    Konversi cepat tanpa render math: mathrm/mathbf/text wrapper dihapus,
    subscript/superscript di-flatten, simbol Greek/operator diganti unicode.
    """
    if not text:
        return text
    # Hapus $...$ delimiters
    text = re.sub(r"\$([^$]*)\$", r"\1", text)
    # mathrm/mathbf/mathit/text → konten saja
    text = re.sub(r"\\(mathrm|mathbf|mathit|text|mathsf|mathtt)\{([^}]*)\}", r"\2", text)
    # \frac{a}{b} → a/b
    text = re.sub(r"\\frac\{([^}]*)\}\{([^}]*)\}", r"(\1)/(\2)", text)
    # \sqrt{a} → √(a)
    text = re.sub(r"\\sqrt\{([^}]*)\}", r"√(\1)", text)
    # _{x} / ^{x} → _x / ^x (flatten)
    text = re.sub(r"_\{([^}]*)\}", r"_\1", text)
    text = re.sub(r"\^\{([^}]*)\}", r"^\1", text)
    # Symbol replacements
    replacements = {
        r"\\approx": "≈",
        r"\\times": "×",
        r"\\cdot": "·",
        r"\\leq": "≤",
        r"\\geq": "≥",
        r"\\neq": "≠",
        r"\\infty": "∞",
        r"\\pm": "±",
        r"\\mp": "∓",
        r"\\circ": "°",
        r"\\degree": "°",
        r"\\alpha": "α",
        r"\\beta": "β",
        r"\\gamma": "γ",
        r"\\delta": "δ",
        r"\\epsilon": "ε",
        r"\\theta": "θ",
        r"\\lambda": "λ",
        r"\\mu": "μ",
        r"\\pi": "π",
        r"\\sigma": "σ",
        r"\\omega": "ω",
        r"\\tau": "τ",
        r"\\Delta": "Δ",
        r"\\Sigma": "Σ",
        r"\\Omega": "Ω",
        r"\\sum": "Σ",
        r"\\int": "∫",
        r"\\partial": "∂",
        r"\\quad": "  ",
        r"\\qquad": "    ",
        r"\\dots": "...",
        r"\\ldots": "...",
        r"\\cdots": "⋯",
        r"\\left": "",
        r"\\right": "",
        r"\\overline": "",
        r"\\underline": "",
        r"\\hat": "",
        r"\\vec": "",
        r"\\dot": "",
        r"\\bar": "",
        r"\\displaystyle": "",
        r"\\,": " ",
        r"\\\\": " ",
        r"\\&": "&",
        r"\\%": "%",
        r"\\#": "#",
        r"\\begin\{[^}]*\}": "",
        r"\\end\{[^}]*\}": "",
    }
    for pat, repl in replacements.items():
        text = re.sub(pat, repl, text)
    # Sisa command \xxx → strip
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    # Curly braces yang tersisa
    text = re.sub(r"[{}]", "", text)
    return text


def load_json() -> dict:
    return json.loads(TEMPLATE_JSON.read_text(encoding="utf-8"))


def clear_body(doc):
    """Hapus semua paragraf + tabel di body, pertahankan sectPr final."""
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def _set_para_style(paragraph, style_name: str) -> None:
    """Set paragraph style dengan styleId match template (kasus-sensitif sesuai
    nama yang ada di styles.xml)."""
    try:
        paragraph.style = doc_style_lookup(paragraph.part.document, style_name)
    except Exception:
        # Fallback: set lewat XML langsung
        ppr = paragraph._p.get_or_add_pPr()
        for old in ppr.findall(qn("w:pStyle")):
            ppr.remove(old)
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(qn("w:val"), style_name)
        ppr.insert(0, pstyle)


def doc_style_lookup(doc, style_name: str):
    """Cari style by name atau styleId."""
    for s in doc.styles:
        if s.name == style_name or s.style_id == style_name:
            return s
    raise KeyError(style_name)


def _add_run(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    superscript: bool = False,
    font_name: str = None,
    size_pt: float = None,
):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if superscript:
        run.font.superscript = True
    if font_name:
        run.font.name = font_name
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
        rFonts.set(qn("w:cs"), font_name)
    if size_pt:
        run.font.size = Pt(size_pt)
    return run


def add_title(doc, data):
    title_text = (data.get("title") or "Paper Title Goes Here").strip()
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_TITLE)
    _add_run(p, title_text)


def add_authors(doc, data):
    authors = data.get("authors") or []
    if not authors:
        authors = [
            {
                "name": "Author Name",
                "affiliation": "Department, University",
                "email": "author@email.ac.id",
            }
        ]

    # Authors line dengan superscript affiliation index
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_AUTHOR)

    # Mapping affiliation -> index
    affil_map = {}
    for author in authors:
        affil = (author.get("affiliation") or "").strip()
        if affil and affil not in affil_map:
            affil_map[affil] = len(affil_map) + 1

    for i, author in enumerate(authors):
        if i > 0:
            connector = ", and " if i == len(authors) - 1 else ", "
            _add_run(p, connector)
        name = (author.get("name") or "Author Name").strip()
        _add_run(p, name)
        affil = (author.get("affiliation") or "").strip()
        idx = affil_map.get(affil)
        if idx:
            _add_run(p, str(idx), superscript=True)

    # Affiliations
    for affil, idx in affil_map.items():
        ap = doc.add_paragraph()
        _set_para_style(ap, STYLE_AFFIL)
        _add_run(ap, str(idx), superscript=True)
        _add_run(ap, affil)


def add_abstract(doc, data):
    abstract = (data.get("abstract") or "Abstract text goes here.").strip()
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_ABSTRACT)
    _add_run(p, "Abstract.", bold=True)
    _add_run(p, " " + abstract)


def add_keywords(doc, data):
    kw = data.get("keywords") or []
    if not kw:
        return  # ICIMECE template doesn't show explicit Keywords field
    text = ", ".join(kw)
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_ABSTRACT)
    _add_run(p, "Keywords: ", bold=True)
    _add_run(p, text, italic=True)


def add_section_heading(doc, title: str):
    """Title only, no manual numbering. Word numbering style yang otomatis."""
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_SECTION)
    _add_run(p, title)


def add_subsection_heading(doc, title: str):
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_SUBSECTION)
    _add_run(p, title)


def add_body_text(doc, text: str, first: bool = False):
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_PARA_FIRST if first else STYLE_PARA)
    _add_run(p, _strip_latex(text))


def _set_full_table_borders(table):
    """Set tabel dengan borders three-line (top/bottom/insideH only) per template."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    visible_sides = {"top", "bottom", "insideH"}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        if edge in visible_sides:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
            el.set(qn("w:sz"), "0")


def add_table(doc, table_data: dict):
    num = str(table_data.get("TableNumber") or "1")
    title = (table_data.get("Title") or "Table title").strip()
    headers = table_data.get("Headers") or ["Col1", "Col2"]
    rows = table_data.get("Rows") or [["Data", "Data"]]

    # Caption (above table)
    cap = doc.add_paragraph()
    _set_para_style(cap, STYLE_TBL_CAPTION)
    _add_run(cap, f"Table {num}.", bold=True)
    _add_run(cap, " " + title)

    # Table
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Template ICIMECE pakai NO_BORDERS pattern (rely on tcBorders per-cell di style),
    # jadi tidak set tblBorders explicit.

    # Header
    for col_idx, header in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(para, str(header), bold=True)

    # Data rows
    for row_idx, row_data in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_data[: len(headers)]):
            cell = table.rows[row_idx].cells[col_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_run(para, str(value))


def add_figure(doc, fig_data: dict):
    """Render figure dengan prompt AI placeholder kalau image tidak ada."""
    image_number = str(fig_data.get("ImageNumber") or "1").strip()
    title = (fig_data.get("Title") or "").strip()
    prompt = (fig_data.get("Prompt") or "").strip()
    path_text = (fig_data.get("Path") or "").strip()

    image_path = None
    if path_text:
        cand = Path(path_text)
        if not cand.is_absolute():
            cand = BASE / path_text
        if cand.is_file():
            image_path = cand

    if image_path:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Cm(8.0))
    else:
        prompt_body = prompt if prompt else f"Figure {image_number}"
        if title:
            placeholder_text = f"[PROMPT UNTUK AI GAMBAR: {title}. {prompt_body}]"
        else:
            placeholder_text = f"[PROMPT UNTUK AI GAMBAR: {prompt_body}]"
        # Wrap dalam tabel 1x1 sebagai prompt-box visual
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(para, placeholder_text, italic=True)

    # Caption below
    if title:
        cap = doc.add_paragraph()
        _set_para_style(cap, STYLE_FIG_CAPTION)
        _add_run(cap, f"Fig. {image_number}.", bold=True)
        _add_run(cap, " " + title)


def add_formula(doc, formula_data: dict):
    num = str(formula_data.get("FormulaNumber") or "").strip()
    text = (formula_data.get("text") or formula_data.get("latex") or "E = mc^2").strip()
    p = doc.add_paragraph()
    _set_para_style(p, STYLE_PARA)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _omml_done = False
    try:
        from _math_omml import append_omml_math as _omml_fn
        _lx = (formula_data.get("latex") or "")
        _omml_done = bool(str(_lx or "").strip()) and _omml_fn(p, _lx)
    except Exception:
        _omml_done = False
    if not _omml_done:
        _add_run(p, _strip_latex(text), italic=True)
    if num:
        _add_run(p, f"\t({num})")


def process_content_item(doc, item: dict, first_para_used: list):
    """Render satu item content sesuai id-nya."""
    if not isinstance(item, dict):
        return
    item_id = (item.get("id") or "").lower()

    if item_id in ("text", ""):
        text = (item.get("text") or "").strip()
        if text:
            add_body_text(doc, text, first=not first_para_used[0])
            first_para_used[0] = True
    elif item_id in ("gambar", "image", "figure"):
        add_figure(doc, item)
    elif item_id in ("tabel", "table"):
        add_table(doc, item)
    elif item_id in ("rumus", "formula", "equation"):
        add_formula(doc, item)


def process_section(doc, section_data: dict, top_level: bool = True):
    if not isinstance(section_data, dict):
        return
    title = (section_data.get("title") or "").strip()
    if title:
        if top_level:
            add_section_heading(doc, title)
        else:
            add_subsection_heading(doc, title)

    first_para_used = [False]
    content = section_data.get("content")
    if isinstance(content, list):
        for item in content:
            process_content_item(doc, item, first_para_used)

    # Process nested subsections (key starts with section{prefix})
    for key in section_data:
        if key in ("title", "content"):
            continue
        if key.startswith("section"):
            process_section(doc, section_data[key], top_level=False)


def add_references(doc, data):
    refs = data.get("references") or {}
    if isinstance(refs, list):
        refs = {"title": "REFERENCES", "content": refs}
    title = (refs.get("title") or "References").strip()
    items = [((r.get("text") or r.get("Text") or "").strip() if isinstance(r, dict) else str(r)) for r in (refs.get("content") or [])]
    items = [t for t in items if t] or ["[1] Author, Title, Journal, Year."]

    p_title = doc.add_paragraph()
    _set_para_style(p_title, STYLE_REF_HEAD)
    _add_run(p_title, title)

    for ref in items:
        rp = doc.add_paragraph()
        _set_para_style(rp, STYLE_REF_BODY)
        _add_run(rp, str(ref))


def generate():
    if not TEMPLATE_DOCX.exists():
        raise FileNotFoundError(f"Template tidak ada: {TEMPLATE_DOCX}")
    if not TEMPLATE_JSON.exists():
        raise FileNotFoundError(f"JSON tidak ada: {TEMPLATE_JSON}")

    # Copy template asli sebagai base (preserve styles, headers, footers, sectPr)
    shutil.copy2(TEMPLATE_DOCX, OUTPUT_DOCX)
    doc = Document(str(OUTPUT_DOCX))

    # Hapus body content, pertahankan sectPr
    clear_body(doc)

    data = load_json()

    add_title(doc, data)
    add_authors(doc, data)
    add_abstract(doc, data)
    add_keywords(doc, data)

    # Process all sections
    for i in range(1, 30):
        key = f"section{i}"
        if key in data:
            process_section(doc, data[key], top_level=True)

    add_references(doc, data)

    doc.save(str(OUTPUT_DOCX))
    print(f"Generated: {OUTPUT_DOCX}")
    return str(OUTPUT_DOCX)


if __name__ == "__main__":
    generate()
