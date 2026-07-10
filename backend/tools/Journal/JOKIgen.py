"""
JOKIgen.py — Generator dokumen JOKI dari JSON.

Dokumen dibangun dengan menyalin JOKI.docx asli lalu mengisi ulang body.
Header, footer, styles, numbering, theme, dan section properties tetap diambil
langsung dari template asli agar hasilnya mengikuti XML dokumen sumber.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "JOKI.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem
MAX_FIGURE_WIDTH_CM = 15.5

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]
_XSLT = None


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _style_id(paragraph) -> str | None:
    ppr = paragraph._p.pPr
    if ppr is None:
        return None
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        return None
    return pstyle.get(qn("w:val"))


def _paragraph_text(paragraph) -> str:
    return "".join(node.text or "" for node in paragraph._p.iter(qn("w:t")))


def _has_numpr(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def _clone_ppr(paragraph):
    ppr = paragraph._p.pPr
    return deepcopy(ppr) if ppr is not None else None


def _find_paragraph(paragraphs, label: str, predicate):
    for paragraph in paragraphs:
        if predicate(paragraph):
            return paragraph
    raise RuntimeError(f"Prototype paragraph not found: {label}")


def _find_table(doc: Document, label: str, predicate):
    for table in doc.tables:
        if predicate(table):
            return table
    raise RuntimeError(f"Prototype table not found: {label}")


def _collect_prototypes(doc: Document) -> dict[str, object]:
    paragraphs = list(doc.paragraphs)
    title_paragraphs = [p for p in paragraphs if _style_id(p) == "TitleJOKI"]
    abstract_paragraphs = [p for p in paragraphs if _style_id(p) == "AbstractJOKI"]

    if len(title_paragraphs) < 2:
        raise RuntimeError("TitleJOKI paragraphs not found in template")
    if len(abstract_paragraphs) < 2:
        raise RuntimeError("AbstractJOKI paragraphs not found in template")

    equation_table = _find_table(
        doc,
        "equation_table",
        lambda table: bool(
            table._tbl.xpath(".//*[local-name()='oMath' or local-name()='oMathPara']")
        ),
    )

    return {
        "title_primary": _clone_ppr(title_paragraphs[0]),
        "title_secondary": _clone_ppr(title_paragraphs[1]),
        "author": _clone_ppr(
            _find_paragraph(paragraphs, "author", lambda p: _style_id(p) == "AuthorJOKI")
        ),
        "address": _clone_ppr(
            _find_paragraph(paragraphs, "address", lambda p: _style_id(p) == "AddressJOKI")
        ),
        "email": _clone_ppr(
            _find_paragraph(paragraphs, "email", lambda p: _style_id(p) == "Authoremail")
        ),
        "abstract_heading_en": _clone_ppr(
            _find_paragraph(
                paragraphs,
                "abstract_heading_en",
                lambda p: _style_id(p) == "AbstractheadingJOKI"
                and _paragraph_text(p).strip().lower().startswith("abstract"),
            )
        ),
        "abstract_heading_id": _clone_ppr(
            _find_paragraph(
                paragraphs,
                "abstract_heading_id",
                lambda p: _style_id(p) == "AbstractheadingJOKI"
                and _paragraph_text(p).strip().lower().startswith("abstrak"),
            )
        ),
        "abstract_body_en": _clone_ppr(abstract_paragraphs[0]),
        "abstract_body_id": _clone_ppr(abstract_paragraphs[1]),
        "keywords_en": _clone_ppr(
            _find_paragraph(
                paragraphs,
                "keywords_en",
                lambda p: _paragraph_text(p).strip().lower().startswith("keywords:"),
            )
        ),
        "keywords_id": _clone_ppr(
            _find_paragraph(
                paragraphs,
                "keywords_id",
                lambda p: _paragraph_text(p).strip().lower().startswith("kata kunci:"),
            )
        ),
        "heading1": _clone_ppr(
            _find_paragraph(paragraphs, "heading1", lambda p: _style_id(p) == "Heading1")
        ),
        "heading2": _clone_ppr(
            _find_paragraph(paragraphs, "heading2", lambda p: _style_id(p) == "Heading2")
        ),
        "heading3": _clone_ppr(
            _find_paragraph(paragraphs, "heading3", lambda p: _style_id(p) == "Heading3")
        ),
        "body": _clone_ppr(
            _find_paragraph(paragraphs, "body", lambda p: _style_id(p) == "bodytextJOKI")
        ),
        "figure": _clone_ppr(
            _find_paragraph(paragraphs, "figure", lambda p: _style_id(p) == "figureJOKI")
        ),
        "figure_caption": _clone_ppr(
            _find_paragraph(
                paragraphs, "figure_caption", lambda p: _style_id(p) == "CaptionFigureJOKI"
            )
        ),
        "table_caption": _clone_ppr(
            _find_paragraph(
                paragraphs, "table_caption", lambda p: _style_id(p) == "CaptionTableJOKI"
            )
        ),
        "reference": _clone_ppr(
            _find_paragraph(
                paragraphs,
                "reference",
                lambda p: _style_id(p) == "reference" and _has_numpr(p),
            )
        ),
        "equation_table": deepcopy(equation_table._tbl),
    }


def _apply_ppr(paragraph, prototype_ppr) -> None:
    existing = paragraph._p.pPr
    if existing is not None:
        paragraph._p.remove(existing)
    if prototype_ppr is not None:
        paragraph._p.insert(0, deepcopy(prototype_ppr))


def _set_para_style(paragraph, style_id: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style_id)


def _add_proto_paragraph(doc: Document, prototype_ppr, style_id: str | None = None):
    paragraph = doc.add_paragraph()
    _apply_ppr(paragraph, prototype_ppr)
    if style_id:
        _set_para_style(paragraph, style_id)
    return paragraph


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _insert_block_before_sectpr(doc: Document, element) -> None:
    body = doc._element.body
    sectpr = body.find(qn("w:sectPr"))
    if sectpr is None:
        body.append(element)
        return
    index = list(body).index(sectpr)
    body.insert(index, element)


def _append_cloned_table(doc: Document, prototype_tbl) -> Table:
    cloned = deepcopy(prototype_tbl)
    _insert_block_before_sectpr(doc, cloned)
    return Table(cloned, doc._body)


def _clear_paragraph_contents(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _clear_cell(cell) -> object:
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        return cell.add_paragraph()
    paragraph = paragraphs[0]
    _clear_paragraph_contents(paragraph)
    for extra in paragraphs[1:]:
        cell._tc.remove(extra._p)
    return paragraph


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in {"oMath", "oMathPara"}:
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _append_text_run(
    paragraph, text: str, bold: bool = False, italic: bool = False, underline: bool = False
):
    if not text:
        return None
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if underline:
        run.underline = True
    return run


def _append_line_break(paragraph):
    paragraph.add_run().add_break()


def _normalize_text_commands(text: str) -> str:
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\b")
                    index += 2
                    continue
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\i")
                    index += 2
                    continue
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\u")
                    index += 2
                    continue
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def _append_rich_text(paragraph, text: str):
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            _append_line_break(paragraph)
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                run = paragraph.add_run(token["value"])
                run.italic = True
            continue
        _append_text_run(
            paragraph,
            token["value"],
            bold=token["bold"],
            italic=token["italic"],
            underline=token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [""]


def _body_paragraphs(doc: Document, prototypes: dict[str, object], text: str):
    for block in _split_body_blocks(text):
        paragraph = _add_proto_paragraph(doc, prototypes["body"])
        _append_rich_text(paragraph, block)


def _add_superscript_run(paragraph, text: str):
    run = paragraph.add_run(text)
    rpr = run._r.get_or_add_rPr()
    valign = OxmlElement("w:vertAlign")
    valign.set(qn("w:val"), "superscript")
    rpr.append(valign)
    return run


def _keyword_paragraph(
    doc: Document, prototypes: dict[str, object], text: str, local: bool = False
):
    prototype_key = "keywords_id" if local else "keywords_en"
    paragraph = _add_proto_paragraph(doc, prototypes[prototype_key], style_id="keywordJOKI")
    run = paragraph.add_run(text)
    run.italic = True
    return paragraph


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _is_numericish(text: str) -> bool:
    compact = text.strip()
    if not compact:
        return False
    normalized = compact
    for old in [" ", ",", ".", "%", "(", ")", "[", "]", "±", "~", "°"]:
        normalized = normalized.replace(old, "")
    normalized = normalized.replace("–", "-")
    if re.fullmatch(r"[-+]?\d+(?:-\d+)?", normalized):
        return True
    if re.fullmatch(r"[A-Z]?\d+(?:-[A-Z]?\d+)?", compact):
        return True
    return False


def _table_style_for_cell(text: str, is_header: bool) -> tuple[str, int]:
    if is_header:
        return "tablecenter", WD_ALIGN_PARAGRAPH.CENTER
    if _is_numericish(text):
        return "tablecenter", WD_ALIGN_PARAGRAPH.CENTER
    if len(text.strip()) <= 12 and " " not in text.strip():
        return "tablecenter", WD_ALIGN_PARAGRAPH.CENTER
    return "tableleft", WD_ALIGN_PARAGRAPH.LEFT


def _configure_data_table(table: Table) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    tbl_pr = table._tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tbl_pr)

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "4944")
    tbl_w.set(qn("w:type"), "pct")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "108")
    tbl_ind.set(qn("w:type"), "dxa")


def _add_title_block(doc: Document, prototypes: dict[str, object], config: dict) -> None:
    title_primary = str(
        config.get("title_en")
        or config.get("title")
        or config.get("TitleBlock", {}).get("Title", "")
    ).strip()
    title_secondary = str(
        config.get("title_id") or config.get("judul") or config.get("title_local") or title_primary
    ).strip()

    paragraph = _add_proto_paragraph(doc, prototypes["title_primary"])
    _append_rich_text(paragraph, title_primary)

    paragraph = _add_proto_paragraph(doc, prototypes["title_secondary"])
    _append_rich_text(paragraph, title_secondary)


def _parse_author_entries(config: dict) -> list[dict[str, str]]:
    authors = config.get("authors", [])
    if not isinstance(authors, list):
        return []
    entries: list[dict[str, str]] = []
    for author in authors:
        if not isinstance(author, dict):
            continue
        entries.append(
            {
                "name": str(author.get("name", "")).strip(),
                "affiliation": str(author.get("affiliation", "")).strip(),
                "location": str(author.get("location", "")).strip(),
                "email": str(author.get("email", "")).strip(),
            }
        )
    return [entry for entry in entries if entry["name"]]


def _add_authors(doc: Document, prototypes: dict[str, object], config: dict) -> None:
    entries = _parse_author_entries(config)
    if not entries:
        return

    aff_map: dict[tuple[str, str], int] = {}
    aff_numbers: list[int] = []
    for entry in entries:
        key = (entry["affiliation"], entry["location"])
        if key not in aff_map:
            aff_map[key] = len(aff_map) + 1
        aff_numbers.append(aff_map[key])

    author_paragraph = _add_proto_paragraph(doc, prototypes["author"])
    for index, entry in enumerate(entries):
        if index > 0:
            author_paragraph.add_run(", " if index < len(entries) - 1 else " and ")
        _add_superscript_run(author_paragraph, str(aff_numbers[index]))
        author_paragraph.add_run(entry["name"])
        if index == 0 and any(author["email"] for author in entries):
            author_paragraph.add_run(" ")
            _add_superscript_run(author_paragraph, "*)")

    grouped_affiliations: list[tuple[int, tuple[str, str]]] = []
    for key, number in aff_map.items():
        grouped_affiliations.append((number, key))
    grouped_affiliations.sort(key=lambda item: item[0])

    for number, (affiliation, location) in grouped_affiliations:
        paragraph = _add_proto_paragraph(doc, prototypes["address"])
        _add_superscript_run(paragraph, str(number))
        paragraph.add_run(", ".join(part for part in [affiliation, location] if part))

    emails = [entry["email"] for entry in entries if entry["email"]]
    if emails:
        paragraph = _add_proto_paragraph(doc, prototypes["email"])
        _add_superscript_run(paragraph, "*) ")
        paragraph.add_run(f"corresponding email: {emails[0]}")
        if len(emails) > 1:
            paragraph.add_run("; ")
            paragraph.add_run("; ".join(emails[1:]))


def _keywords_text(config: dict, local: bool = False) -> str:
    if local:
        values = (
            config.get("keywords_id")
            or config.get("kata_kunci")
            or config.get("keywords_local")
            or config.get("keywords")
            or []
        )
    else:
        values = config.get("keywords_en") or config.get("keywords") or []

    if isinstance(values, str):
        parts = [part.strip() for part in re.split(r"[,;]", values) if part.strip()]
    else:
        parts = [str(part).strip() for part in values if str(part).strip()]
    return ", ".join(parts)


def _add_abstracts(doc: Document, prototypes: dict[str, object], config: dict) -> None:
    abstract_en = str(config.get("abstract_en") or config.get("abstract") or "").strip()
    abstract_id = str(
        config.get("abstract_id")
        or config.get("abstrak")
        or config.get("abstract_local")
        or abstract_en
    ).strip()
    keywords_en = _keywords_text(config)
    keywords_id = _keywords_text(config, local=True)

    heading = _add_proto_paragraph(doc, prototypes["abstract_heading_en"])
    heading.add_run("Abstract")

    paragraph = _add_proto_paragraph(doc, prototypes["abstract_body_en"])
    _append_rich_text(paragraph, abstract_en)

    _keyword_paragraph(doc, prototypes, f"Keywords: {keywords_en}")

    heading = _add_proto_paragraph(doc, prototypes["abstract_heading_id"])
    run = heading.add_run("Abstrak")
    run.italic = True

    paragraph = _add_proto_paragraph(doc, prototypes["abstract_body_id"])
    _append_rich_text(paragraph, abstract_id)

    _keyword_paragraph(doc, prototypes, f"Kata Kunci: {keywords_id}", local=True)


def _add_section_heading(doc: Document, prototypes: dict[str, object], text: str):
    paragraph = _add_proto_paragraph(doc, prototypes["heading1"])
    _append_rich_text(paragraph, text)
    return paragraph


def _add_subsection_heading(doc: Document, prototypes: dict[str, object], text: str):
    paragraph = _add_proto_paragraph(doc, prototypes["heading2"])
    _append_rich_text(paragraph, text)
    return paragraph


def _add_subsubsection_heading(doc: Document, prototypes: dict[str, object], text: str):
    paragraph = _add_proto_paragraph(doc, prototypes["heading3"])
    _append_rich_text(paragraph, text)
    return paragraph


def _add_figure(doc: Document, prototypes: dict[str, object], item: dict, json_path: Path):

    paragraph = _add_proto_paragraph(doc, prototypes["figure"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    path_text = str(item.get("Path", "")).strip()
    image_path = _resolve_path(path_text, json_path) if path_text else None

    try:
        width_cm = float(item.get("WidthCm", MAX_FIGURE_WIDTH_CM))
    except (TypeError, ValueError):
        width_cm = MAX_FIGURE_WIDTH_CM
    width_cm = max(1.0, min(width_cm, MAX_FIGURE_WIDTH_CM))

    if image_path is not None and image_path.is_file():
        paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
    else:
        fallback = str(item.get("Prompt") or path_text or "Gambar tidak ditemukan").strip()
        _append_rich_text(paragraph, fallback)

    number = str(item.get("ImageNumber") or "").strip()
    title = str(item.get("Title") or "").strip()
    caption_text = f"Figure {number}. {title}" if number else title
    if caption_text:
        caption = _add_proto_paragraph(doc, prototypes["figure_caption"])
        _append_rich_text(caption, caption_text)


def _set_cell_text(cell, text: str, style_id: str, alignment: int):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_para_style(paragraph, style_id)
    paragraph.alignment = alignment
    _append_rich_text(paragraph, text)


def _add_table(doc: Document, prototypes: dict[str, object], item: dict):
    headers = list(item.get("Headers") or item.get("headers") or [])
    rows = list(item.get("Rows") or item.get("rows") or [])
    if not headers:
        return

    number = str(item.get("TableNumber") or "").strip()
    title = str(item.get("Title") or item.get("title") or "").strip()
    caption_text = f"Table {number}. {title}" if number else title
    if caption_text:
        caption = _add_proto_paragraph(doc, prototypes["table_caption"])
        _append_rich_text(caption, caption_text)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    _set_table_full_borders(table)
    _configure_data_table(table)

    for column_index, header in enumerate(headers):
        _set_cell_text(
            table.rows[0].cells[column_index], str(header), "tablecenter", WD_ALIGN_PARAGRAPH.CENTER
        )

    for row_index, row_values in enumerate(rows, start=1):
        for column_index in range(len(headers)):
            value = ""
            if column_index < len(row_values):
                value = str(row_values[column_index])
            style_id, alignment = _table_style_for_cell(value, is_header=False)
            _set_cell_text(table.rows[row_index].cells[column_index], value, style_id, alignment)


def _add_equation(doc: Document, prototypes: dict[str, object], item: dict):
    formula = str(item.get("latex") or item.get("text") or "").strip()
    if not formula:
        return

    table = _append_cloned_table(doc, prototypes["equation_table"])
    cells = table.rows[0].cells
    left_paragraph = _clear_cell(cells[0])
    middle_paragraph = _clear_cell(cells[1])
    right_paragraph = _clear_cell(cells[2])

    left_paragraph.add_run("")
    if not _append_inline_math(middle_paragraph, formula):
        run = middle_paragraph.add_run(formula)
        run.italic = True

    number = str(item.get("FormulaNumber") or item.get("NumberiOrLetter") or "").strip()
    if number:
        right_paragraph.add_run(f"({number})")


def _render_content_item(doc: Document, prototypes: dict[str, object], item: dict, json_path: Path):
    item_id = str(item.get("id", "")).strip().lower()
    if item_id == "text":
        text = str(item.get("text", "")).strip()
        if text:
            _body_paragraphs(doc, prototypes, text)
        return
    if item_id in {"gambar", "figure", "image"}:
        _add_figure(doc, prototypes, item, json_path)
        return
    if item_id in {"tabel", "table"}:
        _add_table(doc, prototypes, item)
        return
    if item_id in {"rumus", "formula", "equation"}:
        _add_equation(doc, prototypes, item)


def _render_subsubsection(
    doc: Document, prototypes: dict[str, object], section: dict, json_path: Path
):
    title = str(section.get("title", "")).strip()
    if title:
        _add_subsubsection_heading(doc, prototypes, title)
    content = section.get("content", [])
    if isinstance(content, str):
        if content.strip():
            _body_paragraphs(doc, prototypes, content.strip())
        return
    for item in content:
        if isinstance(item, dict):
            _render_content_item(doc, prototypes, item, json_path)
        elif isinstance(item, str) and item.strip():
            _body_paragraphs(doc, prototypes, item.strip())


def _render_subsection(
    doc: Document, prototypes: dict[str, object], section: dict, json_path: Path, sub_key: str
):
    title = str(section.get("title", "")).strip()
    if title:
        _add_subsection_heading(doc, prototypes, title)

    content = section.get("content", [])
    if isinstance(content, str):
        if content.strip():
            _body_paragraphs(doc, prototypes, content.strip())
    else:
        for item in content:
            if isinstance(item, dict):
                _render_content_item(doc, prototypes, item, json_path)
            elif isinstance(item, str) and item.strip():
                _body_paragraphs(doc, prototypes, item.strip())

    nested_keys = sorted(
        [key for key in section.keys() if re.fullmatch(rf"{re.escape(sub_key)}[a-z]+", key)]
    )
    for nested_key in nested_keys:
        nested_section = section.get(nested_key)
        if isinstance(nested_section, dict):
            _render_subsubsection(doc, prototypes, nested_section, json_path)


def _render_sections(doc: Document, prototypes: dict[str, object], config: dict, json_path: Path):
    section_keys = sorted(
        [key for key in config.keys() if re.fullmatch(r"section\d+", key)],
        key=lambda value: int(value.replace("section", "")),
    )

    for section_key in section_keys:
        section = config[section_key]
        if not isinstance(section, dict):
            continue
        title = str(section.get("title", "")).strip()
        if title:
            _add_section_heading(doc, prototypes, title)

        content = section.get("content", [])
        if isinstance(content, str):
            if content.strip():
                _body_paragraphs(doc, prototypes, content.strip())
        else:
            for item in content:
                if isinstance(item, dict):
                    _render_content_item(doc, prototypes, item, json_path)
                elif isinstance(item, str) and item.strip():
                    _body_paragraphs(doc, prototypes, item.strip())

        section_number = section_key.replace("section", "")
        subsection_keys = sorted(
            [key for key in section.keys() if re.fullmatch(rf"section{section_number}[a-z]", key)]
        )
        for subsection_key in subsection_keys:
            subsection = section.get(subsection_key)
            if isinstance(subsection, dict):
                _render_subsection(doc, prototypes, subsection, json_path, subsection_key)


def _add_references(doc: Document, prototypes: dict[str, object], config: dict):
    refs = config.get("references")
    if isinstance(refs, list):
        title = "REFERENCES"
        content = refs
    elif isinstance(refs, dict):
        title = str(refs.get("title", "REFERENCES")).strip() or "REFERENCES"
        content = refs.get("content") or refs.get("items") or []
    else:
        title = "REFERENCES"
        content = config.get("References", [])

    if not content:
        return

    _add_section_heading(doc, prototypes, title)

    for item in content:
        text = str(item.get("text", "")).strip() if isinstance(item, dict) else str(item).strip()
        text = re.sub(r"^\s*\[\d+\]\s*", "", text)
        if not text:
            continue
        paragraph = _add_proto_paragraph(doc, prototypes["reference"])
        _append_rich_text(paragraph, text)


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,

) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    prototypes = _collect_prototypes(doc)
    _clear_document_body(doc)

    _add_title_block(doc, prototypes, config)
    _add_authors(doc, prototypes, config)
    _add_abstracts(doc, prototypes, config)
    _render_sections(doc, prototypes, config, Path(json_path))
    _add_references(doc, prototypes, config)

    doc.save(str(final_output))
    print(f"Generated: {final_output}")
    return final_output


def build_pdf(json_path: Path, pdf_path: Path, template_path=None) -> Path:
    """Build a PDF for this journal template from a paper JSON.

    Calls build_document() to produce a .docx, then converts to .pdf
    via LibreOffice headless.  Final PDF is written to ``pdf_path``.
    """
    from ._render_pdf import build_pdf_from_builder
    return build_pdf_from_builder(build_document, json_path, pdf_path, template_path)

def main():
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        if not json_arg.exists():
            print(f"File tidak ditemukan: {json_arg}")
            sys.exit(1)
        result = build_document(json_arg, output_arg, template_arg)
        print(f"Selesai: {result}")
        return

    result = build_document()
    print(f"Selesai: {result}")


def _set_table_full_borders(table) -> None:
    """Pastikan tabel punya border tegas/visible (val=single, sz=4 = 0.5pt).

    Dipanggil setelah doc.add_table() supaya tabel data keliatan di Word.
    Auto-injected oleh _fix_table_borders.py untuk lulus audit border check.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


if __name__ == "__main__":
    main()