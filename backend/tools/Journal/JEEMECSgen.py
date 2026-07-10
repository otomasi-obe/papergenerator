"""
JEEMECSgen.py - generate a JEEMECS paper from JSON using the original JEEMECS.docx.

The generator keeps the original journal template, preserves header/footer and
package parts, rewrites the front matter in-place, then rebuilds the manuscript
body from JSON.
"""

from __future__ import annotations

import json
import re
import shutil
import string
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "JEEMECS.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

TITLE_FONT_PT = 17.0
AUTHOR_FONT_PT = 11.0
AFFILIATION_FONT_PT = 9.0
MAX_FIGURE_WIDTH_CM = 14.5
EQUATION_TAB_RIGHT_PT = 425.25
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

STYLE_XML_ID = {
    "TitleIJAIN": "TitleIJAIN",
    "Author": "Author",
    "AuthorAffiliation": "AuthorAffiliation",
    "AbstractHead": "AbstractHead",
    "AbstractText": "AbstractText",
    "Keyword": "Keyword",
    "KeywordHead": "KeywordHead",
    "BodyText": "BodyText",
    "Heading1": "Heading1",
    "Heading2": "Heading2",
    "Heading3": "Heading3",
    "Heading5": "Heading5",
    "equation": "equation",
    "figurecaption": "figurecaption",
    "tablehead": "tablehead",
    "tablecolhead": "tablecolhead",
    "tablecolsubhead": "tablecolsubhead",
    "tablecopy": "tablecopy",
    "tablefootnote": "tablefootnote",
    "references": "references",
}

NUM_HEADING = 1
NUM_FIGURE = 8
NUM_REFERENCE = 9
NUM_TABLE = 10

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]

_XSLT = None
_DRAWING_ID_NEXT = 1
_DOC_PR_PATTERN = re.compile(r'<wp:docPr\b[^>]*\bid="(\d+)"')


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _iter_body_blocks(parent: _DocumentType | _Cell):
    parent_elm = parent._element.body if isinstance(parent, _DocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _clear_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)


def _clear_cell(cell: _Cell) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(OxmlElement("w:p"))


def _actual_cell(table: Table, row_index: int, cell_index: int) -> _Cell:
    row = table._tbl.findall(qn("w:tr"))[row_index]
    tc = row.findall(qn("w:tc"))[cell_index]
    return _Cell(tc, table)


def _set_para_style(paragraph: Paragraph, style_name: str) -> None:
    style_id = STYLE_XML_ID.get(style_name, style_name)
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style_id)


def _set_num_pr(paragraph: Paragraph, *, num_id: int, ilvl: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)

    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))

    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def _set_run_format(
    run,
    *,
    size_pt: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    font_name: str | None = None,
    color_hex: str | None = None,
    superscript: bool = False,
) -> None:
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if font_name is not None:
        run.font.name = font_name
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), font_name)
    if color_hex is not None:
        run.font.color.rgb = RGBColor.from_string(color_hex)
    if superscript:
        run.font.superscript = True


def _append_run(
    paragraph: Paragraph,
    text: str,
    *,
    size_pt: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    font_name: str | None = None,
    color_hex: str | None = None,
    superscript: bool = False,
):
    run = paragraph.add_run(text)
    _set_run_format(
        run,
        size_pt=size_pt,
        bold=bold,
        italic=italic,
        underline=underline,
        font_name=font_name,
        color_hex=color_hex,
        superscript=superscript,
    )
    return run


def _normalize_text_commands(text: str) -> str:
    # Repair LLM streaming artifacts (collapsed integrals, bare math, etc.)
    try:
        from _math_omml import sanitize_llm_text_artifacts
        text = sanitize_llm_text_artifacts(text)
    except Exception:
        pass
    text = re.sub(r'\\\\n(?![a-z])', '\n', text)
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _split_text_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    parts = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or [normalized.strip()]


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue
        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\b")
                    index += 2
                    continue
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\i")
                    index += 2
                    continue
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                next_char = normalized[index + 2] if index + 2 < len(normalized) else ""
                if next_char.islower():
                    buffer.append("\\u")
                    index += 2
                    continue
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue
        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue
        buffer.append(char)
        index += 1
    yield from flush_buffer()


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _append_inline_math(paragraph: Paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _append_rich_text(paragraph: Paragraph, text: str) -> None:
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            paragraph.add_run().add_break()
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                _append_run(paragraph, token["value"], italic=True)
            continue
        run = paragraph.add_run(token["value"])
        _set_run_format(
            run,
            bold=token["bold"] or None,
            italic=token["italic"] or None,
            underline=token["underline"] or None,
        )


def _resolve_path(path_text: str, json_path: Path) -> Path:
    path = Path(path_text)
    if path.is_absolute():
        return path
    json_relative = json_path.parent / path
    if json_relative.exists():
        return json_relative
    return BASE_DIR / path


def _initialize_drawing_ids(doc: Document) -> None:
    global _DRAWING_ID_NEXT
    max_id = 0
    for part in doc.part.package.parts:
        part_name = str(getattr(part, "partname", ""))
        if not part_name.endswith(".xml"):
            continue
        try:
            xml_text = part.blob.decode("utf-8", errors="ignore")
        except Exception:
            continue
        for match in _DOC_PR_PATTERN.finditer(xml_text):
            max_id = max(max_id, int(match.group(1)))
    _DRAWING_ID_NEXT = max_id + 1


def _next_drawing_id() -> int:
    global _DRAWING_ID_NEXT
    drawing_id = _DRAWING_ID_NEXT
    _DRAWING_ID_NEXT += 1
    return drawing_id


def _assign_inline_drawing_id(inline) -> None:
    drawing_id = _next_drawing_id()
    inline._inline.docPr.set("id", str(drawing_id))
    for node in inline._inline.xpath(".//*[local-name()='cNvPr']"):
        node.set("id", str(drawing_id))


def _write_keywords_cell(cell: _Cell, keywords: list[str]) -> None:
    _clear_cell(cell)
    heading = cell.paragraphs[0]
    _set_para_style(heading, "Keyword")
    _append_run(heading, "Keywords", bold=True, font_name="Times New Roman")

    for _ in range(5):
        paragraph = cell.add_paragraph()
        _set_para_style(paragraph, "Keyword")
        _append_run(paragraph, "", font_name="Times New Roman")

    for _ in range(2):
        paragraph = cell.add_paragraph()
        _set_para_style(paragraph, "Articlehistory")

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _write_keyword_body_cell(cell: _Cell, keywords: list[str]) -> None:
    _clear_cell(cell)

    heading = cell.paragraphs[0]
    _set_para_style(heading, "KeywordHead")

    paragraph = cell.add_paragraph()
    _set_para_style(paragraph, "Keyword")
    if keywords:
        _append_run(paragraph, "; ".join(keywords), font_name="Times New Roman")


def _write_abstract_cell(cell: _Cell, abstract_text: str) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _set_para_style(paragraph, "AbstractText")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _append_rich_text(paragraph, abstract_text)

    blank_one = cell.add_paragraph()
    _set_para_style(blank_one, "Copyright")

    blank_two = cell.add_paragraph()
    _set_para_style(blank_two, "Copyright")
    _append_run(blank_two, "   ")

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _write_simple_cell(
    cell: _Cell,
    style_name: str,
    text: str,
    *,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    bold: bool | None = None,
) -> None:
    _clear_cell(cell)
    paragraph = cell.paragraphs[0]
    _set_para_style(paragraph, style_name)
    paragraph.alignment = align
    _append_run(paragraph, text, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _update_article_info_table(table: Table, config: dict) -> None:
    abstract = str(config.get("abstract", "")).strip()
    keywords = [str(item).strip() for item in config.get("keywords", []) if str(item).strip()]

    row0 = [_actual_cell(table, 0, index) for index in range(4)]
    row1 = [_actual_cell(table, 1, index) for index in range(4)]
    row2 = [_actual_cell(table, 2, index) for index in range(4)]

    _write_simple_cell(row0[0], "ArticleinfoHead", "")
    _write_simple_cell(row0[1], "ArticleinfoHead", "")
    _write_simple_cell(row0[2], "AbstractHead", "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER)
    _write_simple_cell(row0[3], "AbstractHead", "", align=WD_ALIGN_PARAGRAPH.CENTER)

    _write_keywords_cell(row1[0], keywords)
    _write_simple_cell(row1[1], "Keyword", "")
    _write_abstract_cell(row1[2], abstract)
    _write_simple_cell(row1[3], "AbstractText", "")

    _write_keyword_body_cell(row2[0], keywords)
    _write_simple_cell(row2[1], "Keyword", "")
    _write_simple_cell(row2[2], "Copyright", "")
    _write_simple_cell(row2[3], "Copyright", "")


def _select_corresponding_author(authors: list[dict]) -> dict | None:
    for author in authors:
        if author.get("corresponding"):
            return author
    return authors[0] if authors else None


def _format_title_block(doc: Document, config: dict) -> Table:
    blocks = list(_iter_body_blocks(doc))
    top_paragraphs: list[Paragraph] = []
    first_table: Table | None = None
    for block in blocks:
        if isinstance(block, Table):
            first_table = block
            break
        top_paragraphs.append(block)

    if first_table is None or len(top_paragraphs) < 6:
        raise RuntimeError("JEEMECS template structure is not the expected title-block layout")

    authors = [author for author in config.get("authors", []) if isinstance(author, dict)]
    corresponding = _select_corresponding_author(authors)

    title_paragraph = top_paragraphs[0]
    _clear_paragraph(title_paragraph)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _append_run(title_paragraph, str(config.get("title", "")).strip(), size_pt=TITLE_FONT_PT)

    author_paragraph = top_paragraphs[1]
    _clear_paragraph(author_paragraph)
    author_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    group_map: dict[tuple[str, str], str] = {}
    for index, author in enumerate(authors):
        key = (
            str(author.get("affiliation", "")).strip(),
            str(author.get("location", "")).strip(),
        )
        if key not in group_map:
            letter_index = len(group_map)
            label = (
                string.ascii_lowercase[letter_index]
                if letter_index < len(string.ascii_lowercase)
                else f"a{letter_index}"
            )
            group_map[key] = label

        if index > 0:
            _append_run(author_paragraph, ", ", size_pt=AUTHOR_FONT_PT)

        _append_run(author_paragraph, str(author.get("name", "")).strip(), size_pt=AUTHOR_FONT_PT)
        _append_run(author_paragraph, group_map[key], size_pt=AUTHOR_FONT_PT, superscript=True)
        _append_run(author_paragraph, ",", size_pt=AUTHOR_FONT_PT, superscript=True)
        _append_run(author_paragraph, str(index + 1), size_pt=AUTHOR_FONT_PT, superscript=True)
        if corresponding is author:
            _append_run(author_paragraph, ",", size_pt=AUTHOR_FONT_PT, superscript=True)
            _append_run(author_paragraph, "*", size_pt=AUTHOR_FONT_PT, superscript=True)

    affiliation_paragraphs = top_paragraphs[2:4]
    group_items = list(group_map.items())
    for paragraph in affiliation_paragraphs:
        _clear_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for paragraph, ((affiliation, location), label) in zip(affiliation_paragraphs, group_items):
        _append_run(paragraph, f"{label} ", size_pt=AFFILIATION_FONT_PT, superscript=True)
        details = ", ".join(part for part in [affiliation, location] if part)
        _append_run(paragraph, details, size_pt=AFFILIATION_FONT_PT)

    email_paragraph = top_paragraphs[4]
    _clear_paragraph(email_paragraph)
    email_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, author in enumerate(authors, start=1):
        if index > 1:
            _append_run(email_paragraph, "; ", size_pt=AFFILIATION_FONT_PT)
        _append_run(email_paragraph, str(index), size_pt=AFFILIATION_FONT_PT, superscript=True)
        email = str(author.get("email", "")).strip()
        if email:
            _append_run(email_paragraph, f" {email}", size_pt=AFFILIATION_FONT_PT)

    corr_paragraph = top_paragraphs[5]
    _clear_paragraph(corr_paragraph)
    corr_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if corresponding is not None:
        _append_run(corr_paragraph, "* corresponding author", size_pt=AFFILIATION_FONT_PT)

    if len(top_paragraphs) >= 7:
        _clear_paragraph(top_paragraphs[6])

    _update_article_info_table(first_table, config)
    return first_table


def _trim_template_body(doc: Document, first_table: Table) -> None:
    body = doc._element.body
    remove = False
    for child in list(body):
        if child is first_table._tbl:
            remove = True
            continue
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_cell_border(
    cell: _Cell, edge: str, value: str, size: str = "4", color: str = "auto"
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    tag = tc_borders.find(qn(f"w:{edge}"))
    if tag is None:
        tag = OxmlElement(f"w:{edge}")
        tc_borders.append(tag)
    tag.set(qn("w:val"), value)
    tag.set(qn("w:sz"), size)
    tag.set(qn("w:space"), "0")
    tag.set(qn("w:color"), color)


def _clear_cell_borders(cell: _Cell) -> None:
    for edge in ("top", "left", "bottom", "right"):
        _set_cell_border(cell, edge, "none", size="0")


def _new_styled_paragraph(
    doc: Document,
    style_name: str,
    *,
    num_id: int | None = None,
    ilvl: int | None = None,
    align=None,
) -> Paragraph:
    paragraph = doc.add_paragraph()
    _set_para_style(paragraph, style_name)
    if num_id is not None and ilvl is not None:
        _set_num_pr(paragraph, num_id=num_id, ilvl=ilvl)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def _body_paragraph(doc: Document, text: str) -> None:
    for block in _split_text_blocks(text):
        paragraph = _new_styled_paragraph(doc, "BodyText")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _append_rich_text(paragraph, block)


def _add_heading(doc: Document, title: str, *, style_name: str, ilvl: int) -> None:
    paragraph = _new_styled_paragraph(doc, style_name, num_id=NUM_HEADING, ilvl=ilvl)
    if style_name == "Heading2":
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        _append_run(
            paragraph,
            title,
            bold=True,
            size_pt=11.0,
            font_name="Times New Roman",
            color_hex="000000",
        )
        return
    if style_name == "Heading1":
        _append_run(paragraph, title, bold=True, font_name="Times New Roman")
        return
    _append_run(paragraph, title, font_name="Times New Roman")


def _set_table_paragraph(
    paragraph: Paragraph, style_name: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.CENTER
) -> None:
    _clear_paragraph(paragraph)
    _set_para_style(paragraph, style_name)
    paragraph.alignment = align
    for run in paragraph.runs:
        run.bold = bold


def _apply_table_row_border(row) -> None:
    for cell in row.cells:
        _set_cell_border(cell, "bottom", "single")


def _add_figure(doc: Document, item: dict, json_path: Path) -> None:

    image_path = _resolve_path(str(item.get("Path", "")).strip(), json_path)
    image_paragraph = doc.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path.is_file():
        inline = image_paragraph.add_run().add_picture(
            str(image_path), width=Cm(MAX_FIGURE_WIDTH_CM)
        )
        _assign_inline_drawing_id(inline)
    else:
        _append_run(
            image_paragraph, str(item.get("Path", "")).strip() or "[missing figure]", italic=True
        )

    caption = _new_styled_paragraph(
        doc,
        "figurecaption",
        num_id=NUM_FIGURE,
        ilvl=0,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _append_rich_text(caption, str(item.get("Title", "")).strip())


def _add_table(doc: Document, item: dict) -> None:
    headers = [str(value) for value in item.get("Headers", [])]
    rows = [[str(value) for value in row] for row in item.get("Rows", [])]
    if not headers:
        return

    caption = _new_styled_paragraph(
        doc,
        "tablehead",
        num_id=NUM_TABLE,
        ilvl=0,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _append_rich_text(caption, str(item.get("Title", "")).strip())

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders_match_template(table)

    for row in table.rows:
        for cell in row.cells:
            _clear_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for column_index, header in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        _clear_cell(cell)
        paragraph = cell.paragraphs[0]
        _set_para_style(paragraph, "tablecolhead")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _append_rich_text(paragraph, header)
        for run in paragraph.runs:
            run.bold = True
        _set_cell_border(cell, "top", "single")
        _set_cell_border(cell, "bottom", "single")

    for row_index, row_values in enumerate(rows, start=1):
        table_row = table.rows[row_index]
        for column_index in range(len(headers)):
            value = row_values[column_index] if column_index < len(row_values) else ""
            cell = table_row.cells[column_index]
            _clear_cell(cell)
            paragraph = cell.paragraphs[0]
            _set_para_style(paragraph, "tablecopy")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _append_rich_text(paragraph, value)
        _apply_table_row_border(table_row)


def _add_equation(doc: Document, item: dict) -> None:
    formula = str(item.get("latex") or item.get("text") or "").strip()
    number = str(item.get("FormulaNumber", "")).strip()
    if not formula:
        return

    paragraph = _new_styled_paragraph(doc, "equation")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Pt(EQUATION_TAB_RIGHT_PT), WD_TAB_ALIGNMENT.RIGHT
    )
    if not _append_inline_math(paragraph, formula):
        _append_run(paragraph, formula, italic=True)
    if number:
        paragraph.add_run().add_tab()
        _append_run(paragraph, f"({number})")


def _render_content_item(doc: Document, item: dict, json_path: Path) -> None:
    item_id = str(item.get("id", "")).strip().lower()
    if item_id == "text":
        text = str(item.get("text", "")).strip()
        if text:
            _body_paragraph(doc, text)
    elif item_id in {"gambar", "image", "figure"}:
        _add_figure(doc, item, json_path)
    elif item_id in {"tabel", "table"}:
        _add_table(doc, item)
    elif item_id in {"rumus", "formula", "equation"}:
        _add_equation(doc, item)


def _render_content(doc: Document, content, json_path: Path) -> None:
    if isinstance(content, str):
        if content.strip():
            _body_paragraph(doc, content.strip())
        return
    if not isinstance(content, list):
        return
    for item in content:
        if isinstance(item, str):
            if item.strip():
                _body_paragraph(doc, item.strip())
        elif isinstance(item, dict):
            _render_content_item(doc, item, json_path)


def _subsection_keys(container: dict, prefix: str) -> list[str]:
    pattern = re.compile(rf"^{re.escape(prefix)}[a-z]+$")
    return sorted([key for key in container if pattern.match(key)])


def _render_subsubsections(
    doc: Document, subsection: dict, json_path: Path, parent_key: str
) -> None:
    child_keys = _subsection_keys(subsection, parent_key)
    for child_key in child_keys:
        child = subsection[child_key]
        if not isinstance(child, dict):
            continue
        title = str(child.get("title", "")).strip()
        if title:
            _add_heading(doc, title, style_name="Heading3", ilvl=2)
        _render_content(doc, child.get("content", []), json_path)


def _render_sections(doc: Document, config: dict, json_path: Path) -> None:
    section_keys = sorted(
        [key for key in config.keys() if re.fullmatch(r"section\d+", key)],
        key=lambda value: int(value.replace("section", "")),
    )
    for section_key in section_keys:
        section = config[section_key]
        if not isinstance(section, dict):
            continue
        title = str(section.get("title", "")).strip()
        if title:
            _add_heading(doc, title, style_name="Heading1", ilvl=0)
        _render_content(doc, section.get("content", []), json_path)

        direct_children = [
            key
            for key in _subsection_keys(section, section_key)
            if len(key) == len(section_key) + 1
        ]
        for subsection_key in direct_children:
            subsection = section[subsection_key]
            if not isinstance(subsection, dict):
                continue
            subsection_title = str(subsection.get("title", "")).strip()
            if subsection_title:
                _add_heading(doc, subsection_title, style_name="Heading2", ilvl=1)
            _render_content(doc, subsection.get("content", []), json_path)
            _render_subsubsections(doc, subsection, json_path, subsection_key)


def _add_references(doc: Document, config: dict) -> None:
    references = config.get("references", {})
    items = (references.get("content") or references.get("items") or []) if isinstance(references, dict) else references
    if not items:
        return

    heading = _new_styled_paragraph(doc, "Heading5", align=WD_ALIGN_PARAGRAPH.CENTER)
    _append_run(
        heading, str(references.get("title", "References")).strip() if isinstance(references, dict) else "References", bold=True
    )

    for item in items:
        text = str(item.get("text", "")).strip() if isinstance(item, dict) else str(item).strip()
        if not text:
            continue
        paragraph = _new_styled_paragraph(doc, "references", num_id=NUM_REFERENCE, ilvl=0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _append_rich_text(paragraph, text)


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,

) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path) if output_path else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _initialize_drawing_ids(doc)
    article_table = _format_title_block(doc, config)
    _trim_template_body(doc, article_table)
    _render_sections(doc, config, Path(json_path))
    _add_references(doc, config)
    doc.save(str(final_output))
    print(f"Generated: {final_output}")
    return final_output


def build_pdf(json_path: Path, pdf_path: Path, template_path=None) -> Path:
    """Build a PDF for this journal template from a paper JSON.

    Calls build_document() to produce a .docx, then converts to .pdf
    via LibreOffice headless.  Final PDF is written to ``pdf_path``.
    """
    from ._render_pdf import build_pdf_from_builder
    return build_pdf_from_builder(build_document, json_path, pdf_path, template_path)

def main() -> None:
    if len(sys.argv) >= 2:
        json_arg = Path(sys.argv[1])
        output_arg = Path(sys.argv[2]) if len(sys.argv) >= 3 else None
        template_arg = Path(sys.argv[3]) if len(sys.argv) >= 4 else TEMPLATE_PATH
        if not json_arg.exists():
            print(f"File not found: {json_arg}")
            sys.exit(1)
        build_document(json_arg, output_arg, template_arg)
        return

    json_files = sorted(
        path
        for path in BASE_DIR.glob("*.json")
        if path.name.lower() not in {"package.json", "tsconfig.json", "settings.json"}
    )
    if not json_files:
        print("No JSON files found.")
        return

    for json_file in json_files:
        try:
            build_document(json_file)
        except Exception as error:
            print(f"Error generating {json_file.name}: {error}")


def _set_table_borders_match_template(table) -> None:
    """Set border tabel sesuai pattern template original: HORIZONTAL_ONLY (academic).

    Pattern booktabs / 3-line table: top + bottom visible, left/right/insideV nil,
    insideH visible (untuk garis di bawah header).
    Auto-injected oleh _fix_table_borders_v2.py.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    visible_sides = {"top", "bottom", "insideH"}
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        if edge in visible_sides:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
            el.set(qn("w:sz"), "0")


if __name__ == "__main__":
    main()