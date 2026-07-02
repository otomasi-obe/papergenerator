"""
ULTIMACOMPgen.py — Generate DOCX paper matching ULTIMACOMP.docx template.
ULTIMACOMP: ULTIMATICS – Jurnal Ilmu Komputer dan Informasi.
Template: A4, 2-section layout: cover (1-col) + body (2-col, IEEE style).

Page : top=85.05, bottom=56.70, left=85.05, right=56.70, header=36.00, footer=36.00
Section structure:
  1. Title  (1-col, no type)       : papertitle + papersubtitle
  2. Authors (1-col, continuous)   : Author + Affiliation lines
  3. Dates  (1-col, continuous)    : received / accepted / approved
  4. Body   (2-col, space=18, continuous) : Abstract, Keywords, sections, references
  5. Final  (1-col)                : empty trailing section

Paragraph styles used:
  papertitle, papersubtitle, Author, Affiliation,
  Abstract, keywords, Heading1, Heading2, Heading5,
  BodyText, bulletlist, Equation0,
  figurecaption, tablehead, tablecolhead, tablecopy, references.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _docx_base import (
    append_rich_text,
    build_sectpr,
    embed_sectpr,
    finalize_doc,
    open_template,
    para,
    render_sections,
    run_generator,
    setup_main_sectpr,
)

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "ULTIMACOMP.docx"

# ── Page dimensions (pt) ──────────────────────────────────────────────────────
PAGE_W = 595.45
PAGE_H = 841.70
TOP = 85.05
BOTTOM = 56.70
LEFT = 85.05
RIGHT = 56.70
HEADER = 36.00
FOOTER = 36.00
COL_SPACE = 18.00

# Body column width for the 2-column section
BODY_COL_WIDTH = (PAGE_W - LEFT - RIGHT - COL_SPACE) / 2  # ≈ 217.85 pt

# ── Style / rendering config ──────────────────────────────────────────────────
CFG = {
    "heading1": "Heading1",
    "heading2": "Heading2",
    "heading3": "Heading3",
    "body": "BodyText",
    "figure_caption": "figurecaption",
    "equation": "Equation0",
    "table_head": "tablehead",
    "table_col_head": "tablecolhead",
    "table_copy": "tablecopy",
    "bullet": "bulletlist",
    "references": "references",
    "fig_prefix": "Fig.",
    "max_fig_width_cm": 8.6,
    "col_width_pt": BODY_COL_WIDTH,
    # Heading1 style has numId=5 (upperRoman auto-numbering) and
    # Heading2 has numId=5 ilvl=1 (upperLetter auto-numbering),
    # so we must NOT add any prefix in the paragraph text.
    "section_heading_format": "plain_upper",  # Word’s list adds “I.”
    "subsection_no_prefix": True,  # Word’s list adds “A.”
    # tablehead numId=9 → “TABLE %1. ”, figurecaption numId=2 → “Fig. %1.”
    # — style already provides the label, so don’t write it in text.
    "table_auto_label": True,
    "figure_auto_label": True,
    # Heading2 rPr has italic=True; using bare add_run preserves that.
    "subsection_bare_run": True,
    "full_borders": True,  # All-border thin box for tables
}


# ── Front-matter builders ─────────────────────────────────────────────────────


def _format_reference(item) -> str:
    """Format a reference dict into a citation string."""
    if isinstance(item, str):
        return item.strip()
    if not isinstance(item, dict):
        return str(item).strip()
    text = item.get("text") or item.get("Text") or item.get("value")
    if text:
        return str(text).strip()
    parts = []
    authors = item.get("authors", [])
    if authors:
        parts.append(", ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors))
    year = item.get("year")
    if year:
        parts.append(f"({year})")
    title = item.get("title", "")
    if title:
        parts.append(f'"{title},"')
    jname = item.get("journal") or item.get("conference") or ""
    if jname:
        parts.append(str(jname) + ",")
    vol = item.get("volume", "")
    if vol:
        parts.append(f"vol. {vol},")
    issue = item.get("issue", "")
    if issue:
        parts.append(f"no. {issue},")
    pages = item.get("pages", "")
    if pages:
        parts.append(f"pp. {pages},")
    doi = item.get("doi", "")
    if doi:
        parts.append(f"doi: {doi}.")
    url = item.get("url", "")
    if url:
        accessed = item.get("accessed", "")
        parts.append(f"[Online]. Available: {url}" + (f" [Accessed: {accessed}]." if accessed else "."))
    publisher = item.get("publisher", "")
    if publisher and not jname:
        location = item.get("location", "")
        parts.append(f"{location}: {publisher}." if location else f"{publisher}.")
    result = " ".join(str(p) for p in parts if p).strip()
    result = result.replace(" , ", ", ").replace(" .", ".")
    if result.endswith(","):
        result = result[:-1] + "."
    if not result.endswith("."):
        result = result + "."
    return result

def _add_title(doc: Document, config: dict) -> None:
    title = config.get("title", "Untitled")
    p = para(doc, style_id="papertitle")
    append_rich_text(p, title)

    subtitle = config.get("subtitle", "")
    if subtitle:
        ps = para(doc, style_id="papersubtitle")
        append_rich_text(ps, subtitle)


def _add_authors(doc: Document, config: dict) -> None:
    authors = config.get("authors", [])
    if not authors:
        return

    # All author names on one line (Author style)
    # Numbers are superscript: Name¹, Name², …
    p = para(doc, style_id="Author")
    for i, a in enumerate(authors, 1):
        name = a.get("name", "")
        if i > 1:
            p.add_run(", ")
        p.add_run(name)
        # superscript number
        num_run = p.add_run(str(i))
        rpr = num_run._r.get_or_add_rPr()
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        rpr.append(vert)

    # Affiliations grouped by unique affiliation+location pair
    seen: dict[str, list[str]] = {}
    for i, a in enumerate(authors, 1):
        aff = a.get("affiliation", "")
        loc = a.get("location", "")
        full = f"{aff}, {loc}" if loc else aff
        seen.setdefault(full, []).append(str(i))

    for full, nums in seen.items():
        p = para(doc, style_id="Affiliation")
        nums_str = ",".join(nums)
        if nums_str:
            num_run = p.add_run(nums_str)
            rpr = num_run._r.get_or_add_rPr()
            vert = OxmlElement("w:vertAlign")
            vert.set(qn("w:val"), "superscript")
            rpr.append(vert)
            p.add_run(" ")
        append_rich_text(p, full)

    # Email line per author (Affiliation style)
    for a in authors:
        email = a.get("email", "")
        if email:
            p = para(doc, style_id="Affiliation")
            append_rich_text(p, email)


def _add_dates(doc: Document, config: dict) -> None:
    """Add received / accepted / approved date lines (Affiliation style)."""
    received = config.get("received", "")
    accepted = config.get("accepted", "")
    approved = config.get("approved", "")

    # Blank spacer (mirrors P[08] in the template)
    para(doc, style_id="Affiliation")

    if received:
        p = para(doc, style_id="Affiliation")
        append_rich_text(p, f"Received on {received}")
    if accepted:
        p = para(doc, style_id="Affiliation")
        append_rich_text(p, f"Accepted on {accepted}")
    if approved:
        p = para(doc, style_id="Affiliation")
        append_rich_text(p, f"Approved on {approved}")


def _bold_run(p, text: str, italic: bool = False) -> None:
    """Add a run with bold (and optionally italic) explicitly set."""
    run = p.add_run(text)
    run.bold = True
    run.italic = italic


def _add_abstract(doc: Document, config: dict) -> None:
    abstract = config.get("abstract", "")
    keywords = config.get("keywords", [])

    if abstract:
        # Abstract style has bold=True in rPr, but append_rich_text sets
        # bold=False explicitly on each run, overriding the style.  We
        # build all runs manually so every character is bold.
        p = para(doc, style_id="Abstract")
        _bold_run(p, "Abstract\u2014")  # em-dash, bold
        _bold_run(p, abstract)  # body text, also bold

    if keywords:
        # Normalize and sort keywords, join with semicolons.
        kws = [str(k).strip() for k in keywords if str(k).strip()]
        if kws:
            kws_sorted = sorted(kws, key=lambda s: s.lower())
            kw = "; ".join(kws_sorted)
            if not kw.endswith("."):
                kw += "."
            # Label is bold+italic; keywords are bold but not italic (after the em-dash).
            kp = para(doc, style_id="keywords")
            _bold_run(kp, "Index Terms\u2014", italic=True)
            _bold_run(kp, kw, italic=False)


# ── References ───────────────────────────────────────────────────────────────


def _add_references(doc: Document, config: dict) -> None:
    """Add references section.

    ULTIMACOMP's 'references' style has numId=8 (abstractNum with text '[%1]'),
    so Word auto-inserts '[1]', '[2]', … as list bullets.  We must NOT add the
    '[n]' prefix text ourselves, or it doubles: '[1] [1] Smith…'.
    Just write the bare reference text; the style numbering handles the label.
    """
    import re as _re

    refs_cfg = config.get("references", {})
    if isinstance(refs_cfg, dict):
        content = refs_cfg.get("content") or refs_cfg.get("items") or [])
    elif isinstance(refs_cfg, list):
        content = refs_cfg
    else:
        content = config.get("References", [])
    if not content:
        return

    # "References" heading (Heading5 — no auto-numbering)
    h = para(doc, style_id="Heading5")
    h.add_run("References")

    for ref in content:
        if isinstance(ref, dict):
            text = ref.get("text", "")
        else:
            text = str(ref)
        # Strip leading "[n]" or "[n] " if caller already embedded it
        text = _re.sub(r"^\s*\[\d+\]\s*", "", text).strip()
        if not text:
            continue
        p = para(doc, style_id="references")
        append_rich_text(p, text)


# ── Main document builder ─────────────────────────────────────────────────────


def build_document(json_path: Path, output_path: Path | None = None) -> Path:
    json_path = Path(json_path)
    config = json.loads(json_path.read_text(encoding="utf-8"))
    out = Path(output_path) if output_path else json_path.with_suffix(".docx")
    if out == TEMPLATE_PATH:
        out = json_path.parent / f"{json_path.stem}_ULTIMACOMP.docx"

    doc = open_template(TEMPLATE_PATH)

    # ── Section 1: Title (1-col, no section-type) ─────────────────────────────
    _add_title(doc, config)
    embed_sectpr(
        doc,
        build_sectpr(
            1,
            36.0,
            TOP,
            BOTTOM,
            LEFT,
            RIGHT,
            section_type=None,
            w_pt=PAGE_W,
            h_pt=PAGE_H,
            header_pt=HEADER,
            footer_pt=FOOTER,
        ),
        style_id="Author",
    )

    # ── Section 2: Authors & affiliations (1-col, continuous) ─────────────────
    _add_authors(doc, config)
    embed_sectpr(
        doc,
        build_sectpr(
            1,
            36.0,
            TOP,
            BOTTOM,
            LEFT,
            RIGHT,
            section_type="continuous",
            w_pt=PAGE_W,
            h_pt=PAGE_H,
            header_pt=HEADER,
            footer_pt=FOOTER,
        ),
        style_id="Affiliation",
    )

    # ── Section 3: Dates (1-col, continuous) ──────────────────────────────────
    _add_dates(doc, config)
    embed_sectpr(
        doc,
        build_sectpr(
            1,
            36.0,
            TOP,
            BOTTOM,
            LEFT,
            RIGHT,
            section_type="continuous",
            w_pt=PAGE_W,
            h_pt=PAGE_H,
            header_pt=HEADER,
            footer_pt=FOOTER,
        ),
    )

    # ── Section 4: Body + References (2-col, continuous) ──────────────────────
    _add_abstract(doc, config)
    render_sections(doc, config, json_path, BASE_DIR, CFG)
    _add_references(doc, config)

    # Close the 2-col section
    embed_sectpr(
        doc,
        build_sectpr(
            2,
            COL_SPACE,
            TOP,
            BOTTOM,
            LEFT,
            RIGHT,
            section_type="continuous",
            w_pt=PAGE_W,
            h_pt=PAGE_H,
            header_pt=HEADER,
            footer_pt=FOOTER,
        ),
    )

    # ── Final section: 1-col (document-level sectPr) ──────────────────────────
    setup_main_sectpr(
        doc,
        w_pt=PAGE_W,
        h_pt=PAGE_H,
        top_pt=TOP,
        bottom_pt=BOTTOM,
        left_pt=LEFT,
        right_pt=RIGHT,
        header_pt=HEADER,
        footer_pt=FOOTER,
        col_space_pt=36.0,
        num_cols=1,
        section_type="continuous",
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    finalize_doc(doc)
    doc.save(str(out))
    return out


if __name__ == "__main__":
    run_generator(BASE_DIR, TEMPLATE_PATH, build_document)
