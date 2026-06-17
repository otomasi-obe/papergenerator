"""
ICOSEGgen.py — Generator dokumen ICOSEG dari JSON.

Prinsip:
- Selalu mulai dari ICOSEG.docx asli (copy template) agar header/footer,
  styles, numbering, rels, theme, dsb tetap mengikuti dokumen sumber.
- Menulis ulang body menggunakan sample pPr/rPr yang di-clone dari template
  (setara dengan paste "Keep Formatting").
- Semua konten diambil dari _PLC-MediapipeID.json. Jika field tidak ada,
  tetap dibuat sesuai format lalu diisi placeholder.

Target: memenuhi kebutuhan format ICOSEG (font, align, spacing, numbering)
berdasarkan template.
"""

from __future__ import annotations

import json
import re
import shutil
import sys
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
JSON_PATH = BASE_DIR / "_PLC-MediapipeID.json"
TEMPLATE_PATH = BASE_DIR / "ICOSEG.docx"
JOURNAL_NAME = TEMPLATE_PATH.stem

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# Usable width in the ICOSEG template is ~470.3pt (see _analyse.py output).
RIGHT_TAB_PT = 470.0
CENTER_TAB_PT = 235.0

# Figure width: keep within text area (≈16.58 cm). Use a small safety margin.
MAX_FIGURE_WIDTH_CM = 16.0

NS_MAP_STRICT = {
    b"http://purl.oclc.org/ooxml/wordprocessingml/main": b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
    b"http://purl.oclc.org/ooxml/drawingml/wordprocessingDrawing": b"http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    b"http://purl.oclc.org/ooxml/officeDocument/math": b"http://schemas.openxmlformats.org/officeDocument/2006/math",
}

XSL_CANDIDATES = [
    BASE_DIR / "MML2OMML.XSL",
    ROOT_DIR / "MML2OMML.XSL",
    Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
]

_XSLT = None


@dataclass
class RenderState:
    figure_number: int = 0
    table_number: int = 0
    equation_number: int = 0


def _wq(tag: str) -> str:
    return f"{{{NS_W}}}{tag}"


def _strict_to_trans(data: bytes) -> bytes:
    for old, new in NS_MAP_STRICT.items():
        data = data.replace(old, new)
    return data


def _clone(element: etree._Element | None) -> etree._Element | None:
    return deepcopy(element) if element is not None else None


def _text_of_run(run_el: etree._Element) -> str:
    parts: list[str] = []
    for t in run_el.findall(f".//{_wq('t')}"):
        if t.text:
            parts.append(t.text)
    return "".join(parts)


def _text_of_paragraph(p_el: etree._Element) -> str:
    parts: list[str] = []
    for t in p_el.findall(f".//{_wq('t')}"):
        if t.text:
            parts.append(t.text)
    return "".join(parts)


def _find_paragraph_index(paragraphs: list[etree._Element], needle: str) -> int | None:
    n = needle.strip().lower()
    if not n:
        return None
    for idx, p in enumerate(paragraphs):
        if n in _text_of_paragraph(p).lower():
            return idx
    return None


def _first_text_run_rpr(p_el: etree._Element) -> etree._Element | None:
    for run in p_el.findall(_wq("r")):
        if _text_of_run(run).strip():
            return _clone(run.find(_wq("rPr")))
    runs = p_el.findall(_wq("r"))
    if runs:
        return _clone(runs[0].find(_wq("rPr")))
    return None


def _run_rpr_containing(p_el: etree._Element, needle: str) -> etree._Element | None:
    n = needle.strip().lower()
    if not n:
        return None
    for run in p_el.findall(_wq("r")):
        if n in _text_of_run(run).lower():
            return _clone(run.find(_wq("rPr")))
    return None


def _find_superscript_run_rpr(p_el: etree._Element) -> etree._Element | None:
    for run in p_el.findall(_wq("r")):
        rpr = run.find(_wq("rPr"))
        if rpr is None:
            continue
        va = rpr.find(_wq("vertAlign"))
        if va is not None and (va.get(_wq("val")) or "").lower() == "superscript":
            return _clone(rpr)
    return None


def _apply_sample_ppr(paragraph, sample_ppr: etree._Element | None) -> None:
    current = paragraph._p.find(qn("w:pPr"))
    if current is not None:
        paragraph._p.remove(current)
    if sample_ppr is not None:
        paragraph._p.insert(0, deepcopy(sample_ppr))


def _apply_sample_rpr(run, sample_rpr: etree._Element | None) -> None:
    current = run._r.find(qn("w:rPr"))
    if current is not None:
        run._r.remove(current)
    if sample_rpr is not None:
        run._r.insert(0, deepcopy(sample_rpr))


def _new_paragraph(doc: Document, sample_ppr: etree._Element | None = None):
    paragraph = doc.add_paragraph()
    if sample_ppr is not None:
        _apply_sample_ppr(paragraph, sample_ppr)
    return paragraph


def _add_sample_run(
    paragraph,
    text: str,
    sample_rpr: etree._Element | None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    underline: bool | None = None,
    superscript: bool = False,
):
    run = paragraph.add_run(text)
    _apply_sample_rpr(run, sample_rpr)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if superscript:
        run.font.superscript = True
    return run


def _clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_document_final_sectpr(doc: Document, sectpr: etree._Element) -> None:
    body = doc._element.body
    current = body.find(qn("w:sectPr"))
    if current is not None:
        body.remove(current)
    body.append(deepcopy(sectpr))


def _append_body_element_before_sectpr(doc: Document, element: etree._Element) -> None:
    body = doc._element.body
    sectpr = body.find(qn("w:sectPr"))
    if sectpr is not None:
        sectpr.addprevious(deepcopy(element))
    else:
        body.append(deepcopy(element))


def _set_spacing_min(
    paragraph, *, before_pt: float | None = None, after_pt: float | None = None
) -> None:
    pf = paragraph.paragraph_format
    if before_pt is not None:
        cur = pf.space_before.pt if pf.space_before is not None else 0.0
        if cur < before_pt:
            pf.space_before = Pt(before_pt)
    if after_pt is not None:
        cur = pf.space_after.pt if pf.space_after is not None else 0.0
        if cur < after_pt:
            pf.space_after = Pt(after_pt)


def _normalize_text_commands(text: str) -> str:
    text = text.replace("\\n", "\n")
    # Convert Markdown bold/italic to \b..\b / \i..\i toggle format
    text = re.sub(r"\*\*(.+?)\*\*", r"\\b\1\\b", text, flags=re.DOTALL)
    text = re.sub(r"\*([^*\n]+?)\*", r"\\i\1\\i", text)
    return text


def _iter_rich_tokens(text: str):
    normalized = _normalize_text_commands(text)
    buffer: list[str] = []
    bold = False
    italic = False
    underline = False
    index = 0

    def flush_buffer():
        nonlocal buffer
        content = "".join(buffer)
        buffer = []
        if content:
            yield {
                "kind": "text",
                "value": content,
                "bold": bold,
                "italic": italic,
                "underline": underline,
            }

    while index < len(normalized):
        char = normalized[index]
        if char == "\n":
            yield from flush_buffer()
            yield {"kind": "linebreak"}
            index += 1
            continue

        if char == "\\" and index + 1 < len(normalized):
            command = normalized[index + 1]
            if command == "\\":
                buffer.append("\\")
                index += 2
                continue
            if command == "b":
                yield from flush_buffer()
                bold = not bold
                index += 2
                continue
            if command == "i":
                yield from flush_buffer()
                italic = not italic
                index += 2
                continue
            if command == "u":
                yield from flush_buffer()
                underline = not underline
                index += 2
                continue

        if char == "$":
            closing = normalized.find("$", index + 1)
            if closing != -1:
                yield from flush_buffer()
                formula = normalized[index + 1 : closing]
                if formula:
                    yield {"kind": "math", "value": formula}
                index = closing + 1
                continue

        buffer.append(char)
        index += 1

    yield from flush_buffer()


def _get_xslt():
    global _XSLT
    if _XSLT is not None:
        return _XSLT
    for candidate in XSL_CANDIDATES:
        try:
            if candidate.exists():
                _XSLT = etree.XSLT(etree.parse(str(candidate)))
                return _XSLT
        except Exception:
            continue
    _XSLT = False
    return _XSLT


def _latex_to_omml(latex: str):
    try:
        import latex2mathml.converter
    except Exception:
        return None
    xslt = _get_xslt()
    if not xslt:
        return None
    try:
        mathml = latex2mathml.converter.convert(latex)
        root = etree.fromstring(mathml.encode("utf-8"))
        return xslt(root).getroot()
    except Exception:
        return None


def _set_math_rpr_defaults(rpr: etree._Element, half_points: int = 20) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Cambria Math")

    for tag in ("sz", "szCs"):
        size_el = rpr.find(qn(f"w:{tag}"))
        if size_el is None:
            size_el = OxmlElement(f"w:{tag}")
            rpr.append(size_el)
        size_el.set(qn("w:val"), str(half_points))


def _normalize_omml_math(omml: etree._Element, half_points: int = 20) -> etree._Element:
    for math_run in omml.findall(f".//{{{MATH_NS}}}r"):
        rpr = math_run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            math_run.insert(0, rpr)
        _set_math_rpr_defaults(rpr, half_points=half_points)

    for ctrl_pr in omml.findall(f".//{{{MATH_NS}}}ctrlPr"):
        rpr = ctrl_pr.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            ctrl_pr.insert(0, rpr)
        _set_math_rpr_defaults(rpr, half_points=half_points)

    return omml


def _append_inline_math(paragraph, latex: str) -> bool:
    omml = _latex_to_omml(latex)
    if omml is None:
        return False
    omml = _normalize_omml_math(omml, half_points=20)
    tag = omml.tag.split("}")[-1] if "}" in omml.tag else omml.tag
    if tag in ("oMath", "oMathPara"):
        paragraph._p.append(omml)
    else:
        wrapper = etree.fromstring(f'<m:oMath xmlns:m="{MATH_NS}"/>')
        wrapper.append(omml)
        paragraph._p.append(wrapper)
    return True


def _append_rich_text(
    paragraph,
    text: str,
    sample_rpr: etree._Element | None,
    *,
    base_bold: bool = False,
    base_italic: bool = False,
    base_underline: bool = False,
) -> None:
    for token in _iter_rich_tokens(text):
        if token["kind"] == "linebreak":
            paragraph.add_run().add_break()
            continue
        if token["kind"] == "math":
            if not _append_inline_math(paragraph, token["value"]):
                _add_sample_run(paragraph, token["value"], sample_rpr, italic=True)
            continue
        _add_sample_run(
            paragraph,
            token["value"],
            sample_rpr,
            bold=base_bold or token["bold"],
            italic=base_italic or token["italic"],
            underline=base_underline or token["underline"],
        )


def _split_body_blocks(text: str) -> list[str]:
    normalized = _normalize_text_commands(text).replace("\r\n", "\n").replace("\r", "\n")
    blocks = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return blocks or [normalized.strip()]


def _resolve_path(path_text: str, json_path: Path) -> Path:
    candidate = Path(path_text)
    if not candidate.is_absolute():
        candidate = (json_path.parent / candidate).resolve()
    return candidate


def _load_template_samples(template_path: Path) -> dict[str, etree._Element | None]:
    with zipfile.ZipFile(template_path) as archive:
        raw = archive.read("word/document.xml")
        if b"purl.oclc.org/ooxml" in raw:
            raw = _strict_to_trans(raw)

    root = etree.fromstring(raw)
    body = root.find(_wq("body"))
    if body is None:
        raise RuntimeError("Template ICOSEG tidak memiliki w:body")

    paragraphs = body.findall(_wq("p"))
    tables = body.findall(_wq("tbl"))

    def para(index: int) -> etree._Element:
        return paragraphs[index]

    def clone_ppr(index: int):
        return _clone(para(index).find(_wq("pPr")))

    def clone_first_text_rpr(index: int):
        return _first_text_run_rpr(para(index))

    def clone_sup_rpr(index: int):
        return _find_superscript_run_rpr(para(index))

    # Anchor-based lookup (fallback to known indices from _analyse_ICOSEG.txt)
    idx_title = _find_paragraph_index(paragraphs, "Title of Manuscript")
    idx_author = _find_paragraph_index(paragraphs, "Author")
    idx_aff1 = _find_paragraph_index(paragraphs, "Affiliation, Country")
    idx_corr = _find_paragraph_index(paragraphs, "Corresponding author")
    idx_email = _find_paragraph_index(paragraphs, "Email:")
    idx_abs = _find_paragraph_index(paragraphs, "Abstract.")
    idx_kw = _find_paragraph_index(paragraphs, "Keywords")
    idx_sec = _find_paragraph_index(paragraphs, "Introduction")
    idx_sub = _find_paragraph_index(paragraphs, "Sub Title")
    idx_body = _find_paragraph_index(paragraphs, "In Introduction")
    idx_fig_cap = _find_paragraph_index(paragraphs, "Figure 1")
    idx_fig_src = _find_paragraph_index(paragraphs, "Source: Source")
    idx_eq = _find_paragraph_index(paragraphs, "(Eq.1)")
    idx_where = _find_paragraph_index(paragraphs, "where:")
    idx_tbl_cap = _find_paragraph_index(paragraphs, "Table 1")
    idx_refs = _find_paragraph_index(paragraphs, "References")
    idx_ref_item = _find_paragraph_index(paragraphs, "Batty,")

    # Known stable indices (template ICOSEG.docx as analysed 2026-04-10)
    idx_title = 1 if idx_title is None else idx_title
    idx_author = 2 if idx_author is None else idx_author
    idx_aff1 = 3 if idx_aff1 is None else idx_aff1
    idx_corr = 6 if idx_corr is None else idx_corr
    idx_email = 7 if idx_email is None else idx_email
    idx_abs = 8 if idx_abs is None else idx_abs
    idx_kw = 9 if idx_kw is None else idx_kw
    idx_sec = 10 if idx_sec is None else idx_sec
    idx_body = 11 if idx_body is None else idx_body
    idx_sub = 13 if idx_sub is None else idx_sub
    idx_fig_src = 16 if idx_fig_src is None else idx_fig_src
    idx_fig_cap = 17 if idx_fig_cap is None else idx_fig_cap
    idx_eq = 20 if idx_eq is None else idx_eq
    idx_where = 21 if idx_where is None else idx_where
    idx_tbl_cap = 23 if idx_tbl_cap is None else idx_tbl_cap
    idx_refs = 38 if idx_refs is None else idx_refs
    idx_ref_item = 43 if idx_ref_item is None else idx_ref_item

    # Figure placeholder paragraph: the paragraph before the figure source line.
    idx_fig_placeholder = max(0, idx_fig_src - 1)

    # Table source line: the "Source:" after the table.
    idx_tbl_source = None
    if idx_tbl_cap is not None:
        for i in range(idx_tbl_cap + 1, len(paragraphs)):
            if "source:" in _text_of_paragraph(paragraphs[i]).lower():
                idx_tbl_source = i
                break
    if idx_tbl_source is None:
        idx_tbl_source = 26  # best-effort fallback (as analysed)

    # Table cell pPr samples (header/body) from the first table.
    header_cell_ppr = None
    body_cell_ppr = None
    if tables:
        rows = tables[0].findall(f".//{_wq('tr')}")
        if rows:
            cells0 = rows[0].findall(_wq("tc"))
            if cells0:
                p0 = cells0[0].find(f".//{_wq('p')}")
                if p0 is not None:
                    header_cell_ppr = _clone(p0.find(_wq("pPr")))
            if len(rows) > 1:
                cells1 = rows[1].findall(_wq("tc"))
                if cells1:
                    p1 = cells1[0].find(f".//{_wq('p')}")
                    if p1 is not None:
                        body_cell_ppr = _clone(p1.find(_wq("pPr")))

    final_sectpr = _clone(body.find(_wq("sectPr")))

    # Logo paragraph (ICOSEG logo is embedded as the first paragraph in template)
    logo_paragraph = _clone(para(0)) if paragraphs else None

    # Run-level samples
    title_rpr = clone_first_text_rpr(idx_title)

    _tmp = _run_rpr_containing(para(idx_author), "Author")
    author_name_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_author)
    _tmp = clone_sup_rpr(idx_author)
    author_sup_rpr = _tmp if _tmp is not None else author_name_rpr

    _tmp = clone_sup_rpr(idx_aff1)
    aff_sup_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_aff1)
    _tmp = _run_rpr_containing(para(idx_aff1), "Affiliation")
    aff_text_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_aff1)

    corr_rpr = clone_first_text_rpr(idx_corr)
    _tmp = _run_rpr_containing(para(idx_email), "Email")
    email_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_email)
    email_text_rpr = clone_first_text_rpr(idx_email)

    _tmp = _run_rpr_containing(para(idx_abs), "Abstract")
    abs_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_abs)
    abs_text_rpr = clone_first_text_rpr(idx_abs)

    _tmp = _run_rpr_containing(para(idx_kw), "Keywords")
    kw_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_kw)
    kw_text_rpr = clone_first_text_rpr(idx_kw)

    sec_rpr = clone_first_text_rpr(idx_sec)
    _tmp = clone_first_text_rpr(idx_sub)
    sub_rpr = _tmp if _tmp is not None else sec_rpr

    body_rpr = clone_first_text_rpr(idx_body)

    _tmp = _run_rpr_containing(para(idx_fig_cap), "Figure")
    fig_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_fig_cap)
    _tmp = _run_rpr_containing(para(idx_fig_cap), "Lorem")
    fig_text_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_fig_cap)

    fig_source_rpr = clone_first_text_rpr(idx_fig_src)

    _tmp = _run_rpr_containing(para(idx_eq), "(Eq.")
    eq_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_eq)
    eq_where_rpr = clone_first_text_rpr(idx_where)

    _tmp = _run_rpr_containing(para(idx_tbl_cap), "Table")
    tbl_label_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_tbl_cap)
    _tmp = _run_rpr_containing(para(idx_tbl_cap), "Lorem")
    tbl_text_rpr = _tmp if _tmp is not None else clone_first_text_rpr(idx_tbl_cap)

    tbl_source_rpr = clone_first_text_rpr(idx_tbl_source)

    ref_heading_rpr = clone_first_text_rpr(idx_refs)
    ref_item_rpr = clone_first_text_rpr(idx_ref_item)

    return {
        "logo_paragraph": logo_paragraph,
        "top_gap_ppr": clone_ppr(0),
        "title_ppr": clone_ppr(idx_title),
        "title_rpr": title_rpr,
        "author_ppr": clone_ppr(idx_author),
        "author_name_rpr": author_name_rpr,
        "author_sup_rpr": author_sup_rpr,
        "affiliation_ppr": clone_ppr(idx_aff1),
        "affiliation_sup_rpr": aff_sup_rpr,
        "affiliation_text_rpr": aff_text_rpr,
        "corresponding_ppr": clone_ppr(idx_corr),
        "corresponding_rpr": corr_rpr,
        "email_ppr": clone_ppr(idx_email),
        "email_label_rpr": email_label_rpr,
        "email_text_rpr": email_text_rpr,
        "abstract_ppr": clone_ppr(idx_abs),
        "abstract_label_rpr": abs_label_rpr,
        "abstract_text_rpr": abs_text_rpr,
        "keywords_ppr": clone_ppr(idx_kw),
        "keywords_label_rpr": kw_label_rpr,
        "keywords_text_rpr": kw_text_rpr,
        "section_heading_ppr": clone_ppr(idx_sec),
        "section_heading_rpr": sec_rpr,
        "subsection_heading_ppr": clone_ppr(idx_sub),
        "subsection_heading_rpr": sub_rpr,
        "body_ppr": clone_ppr(idx_body),
        "body_rpr": body_rpr,
        "figure_ppr": clone_ppr(idx_fig_placeholder),
        "figure_source_ppr": clone_ppr(idx_fig_src),
        "figure_source_rpr": fig_source_rpr,
        "figure_caption_ppr": clone_ppr(idx_fig_cap),
        "figure_label_rpr": fig_label_rpr,
        "figure_text_rpr": fig_text_rpr,
        "equation_ppr": clone_ppr(idx_eq),
        "equation_label_rpr": eq_label_rpr,
        "equation_where_ppr": clone_ppr(idx_where),
        "equation_where_rpr": eq_where_rpr,
        "table_caption_ppr": clone_ppr(idx_tbl_cap),
        "table_label_rpr": tbl_label_rpr,
        "table_text_rpr": tbl_text_rpr,
        "table_source_ppr": clone_ppr(idx_tbl_source),
        "table_source_rpr": tbl_source_rpr,
        "reference_heading_ppr": clone_ppr(idx_refs),
        "reference_heading_rpr": ref_heading_rpr,
        "reference_item_ppr": clone_ppr(idx_ref_item),
        "reference_item_rpr": ref_item_rpr,
        "table_header_cell_ppr": header_cell_ppr,
        "table_body_cell_ppr": body_cell_ppr,
        "final_sectpr": final_sectpr,
    }


def _title_text(config: dict) -> str:
    title = str(config.get("title") or "").strip()
    return title or "Judul Naskah (isi dari JSON tidak tersedia)"


def _parse_authors(config: dict) -> list[dict[str, str]]:
    raw = config.get("authors")
    authors: list[dict[str, str]] = []
    if isinstance(raw, list):
        for entry in raw:
            if not isinstance(entry, dict):
                continue
            authors.append(
                {
                    "name": str(entry.get("name") or "").strip(),
                    "affiliation": str(entry.get("affiliation") or "").strip(),
                    "location": str(entry.get("location") or "").strip(),
                    "email": str(entry.get("email") or "").strip(),
                }
            )
    # Ensure at least one placeholder author
    if not authors:
        authors.append(
            {
                "name": "Nama Penulis",
                "affiliation": "Afiliasi, Negara",
                "location": "",
                "email": "author@icoseg.ac.id",
            }
        )
    # Fill blanks
    for a in authors:
        if not a["name"]:
            a["name"] = "Nama Penulis"
        if not a["affiliation"] and not a["location"]:
            a["affiliation"] = "Afiliasi, Negara"
    return authors


def _abstract_text(config: dict) -> str:
    value = str(config.get("abstract") or "").strip()
    if value:
        return value
    return (
        "Tuliskan abstrak singkat yang memuat latar belakang, tujuan, metode, "
        "hasil utama, dan simpulan. (Field abstract tidak tersedia di JSON.)"
    )


def _keywords_list(config: dict) -> list[str]:
    raw = config.get("keywords")
    if isinstance(raw, list):
        items = [str(x).strip() for x in raw if str(x).strip()]
        if items:
            return items
    return ["Keyword 1", "Keyword 2", "Keyword 3"]


def _acknowledgment_text(config: dict) -> str:
    for key in ("acknowledgment", "acknowledgement", "ucapan_terima_kasih"):
        value = str(config.get(key) or "").strip()
        if value:
            return value
    return (
        "Penulis menyampaikan terima kasih kepada institusi/mitra yang telah "
        "memberikan dukungan fasilitas, perangkat, dan pendanaan untuk penelitian ini."
    )


def _add_front_matter(
    doc: Document, config: dict, samples: dict[str, etree._Element | None]
) -> None:
    # Preserve ICOSEG logo paragraph from template if available.
    logo_p = samples.get("logo_paragraph")
    if logo_p is not None:
        _append_body_element_before_sectpr(doc, logo_p)
    else:
        # Fallback: keep the original top-gap paragraph style.
        _new_paragraph(doc, samples["top_gap_ppr"])

    # Title
    title_para = _new_paragraph(doc, samples["title_ppr"])
    _append_rich_text(title_para, _title_text(config), samples["title_rpr"], base_bold=True)

    # Authors & affiliations
    authors = _parse_authors(config)

    # Group affiliations
    aff_map: dict[tuple[str, str], int] = {}
    aff_groups: list[dict[str, object]] = []
    author_aff_nums: list[int] = []

    for author in authors:
        key = (author.get("affiliation", ""), author.get("location", ""))
        num = aff_map.get(key)
        if num is None:
            num = len(aff_groups) + 1
            aff_map[key] = num
            aff_groups.append(
                {
                    "num": num,
                    "affiliation": author.get("affiliation", ""),
                    "location": author.get("location", ""),
                    "emails": [],
                }
            )
        author_aff_nums.append(num)
        email = author.get("email", "")
        if email:
            group = aff_groups[num - 1]
            emails = group["emails"]
            if email not in emails:
                emails.append(email)

    author_para = _new_paragraph(doc, samples["author_ppr"])
    for i, author in enumerate(authors):
        if i:
            _add_sample_run(author_para, ", ", samples["author_name_rpr"])
        _append_rich_text(author_para, author["name"], samples["author_name_rpr"], base_bold=True)
        _add_sample_run(
            author_para, str(author_aff_nums[i]), samples["author_sup_rpr"], superscript=True
        )
        if i == 0:
            _add_sample_run(author_para, "*", samples["author_sup_rpr"], superscript=True)

    for group in aff_groups:
        aff_para = _new_paragraph(doc, samples["affiliation_ppr"])
        _add_sample_run(
            aff_para, str(group["num"]), samples["affiliation_sup_rpr"], superscript=True
        )
        details = [
            str(group.get("affiliation") or "").strip(),
            str(group.get("location") or "").strip(),
        ]
        details = [d for d in details if d]
        aff_text = ", ".join(details) if details else "Afiliasi, Negara"
        _append_rich_text(aff_para, aff_text, samples["affiliation_text_rpr"])

    # Corresponding line
    corr_para = _new_paragraph(doc, samples["corresponding_ppr"])
    corr_text = "*Corresponding author(s) "
    _append_rich_text(corr_para, corr_text, samples["corresponding_rpr"], base_italic=False)

    # Email line
    email_para = _new_paragraph(doc, samples["email_ppr"])
    emails_all = [a.get("email", "") for a in authors if a.get("email", "")]
    email_value = "; ".join(emails_all) if emails_all else "author@icoseg.ac.id"
    _append_rich_text(email_para, "Email", samples["email_label_rpr"], base_bold=False)
    _append_rich_text(email_para, f": {email_value} ", samples["email_text_rpr"], base_bold=False)

    # Abstract
    abs_para = _new_paragraph(doc, samples["abstract_ppr"])
    _append_rich_text(abs_para, "Abstract. ", samples["abstract_label_rpr"], base_bold=True)
    _append_rich_text(
        abs_para, _abstract_text(config), samples["abstract_text_rpr"], base_bold=False
    )

    # Keywords
    kw_para = _new_paragraph(doc, samples["keywords_ppr"])
    _append_rich_text(kw_para, "Keywords", samples["keywords_label_rpr"], base_bold=True)
    _append_rich_text(kw_para, ": ", samples["keywords_text_rpr"], base_bold=False)
    kw_text = ", ".join(_keywords_list(config)).strip()
    _append_rich_text(kw_para, kw_text, samples["keywords_text_rpr"], base_bold=False)


def _add_section_heading(
    doc: Document, title: str, samples: dict[str, etree._Element | None]
) -> None:
    paragraph = _new_paragraph(doc, samples["section_heading_ppr"])
    _set_spacing_min(paragraph, before_pt=3.0, after_pt=3.0)
    _append_rich_text(paragraph, title, samples["section_heading_rpr"], base_bold=True)


def _add_subsection_heading(
    doc: Document, title: str, samples: dict[str, etree._Element | None]
) -> None:
    paragraph = _new_paragraph(doc, samples["subsection_heading_ppr"])
    _set_spacing_min(paragraph, before_pt=3.0, after_pt=3.0)
    _append_rich_text(
        paragraph, title, samples["subsection_heading_rpr"], base_bold=True, base_italic=True
    )


def _body_paragraph(doc: Document, text: str, samples: dict[str, etree._Element | None]) -> None:
    for block in _split_body_blocks(text):
        if not block:
            continue
        paragraph = _new_paragraph(doc, samples["body_ppr"])
        _append_rich_text(paragraph, block, samples["body_rpr"], base_bold=False)


def _add_missing_figure_notice(
    doc: Document, text: str, samples: dict[str, etree._Element | None]
) -> None:
    paragraph = _new_paragraph(doc, samples["body_ppr"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(paragraph, before_pt=3.0, after_pt=3.0)
    _append_rich_text(paragraph, text, samples["body_rpr"], base_italic=True)


def _add_figure(
    doc: Document,
    item: dict,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
) -> None:

    state.figure_number += 1
    number = str(
        item.get("ImageNumber") or item.get("number") or state.figure_number
    ).strip() or str(state.figure_number)
    title = str(item.get("Title") or item.get("title") or f"Figure title {number}").strip()
    path_text = str(item.get("Path") or item.get("path") or "").strip()
    prompt = str(item.get("Prompt") or "").strip()

    image_path = _resolve_path(path_text, json_path) if path_text else None

    fig_para = _new_paragraph(doc, samples["figure_ppr"])
    fig_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_para.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(fig_para, before_pt=3.0, after_pt=3.0)

    if image_path is not None and image_path.is_file():
        fig_para.add_run().add_picture(str(image_path), width=Cm(MAX_FIGURE_WIDTH_CM))
    else:
        fallback = prompt or title or f"Figure {number}"
        _add_missing_figure_notice(doc, f"[Figure {number} belum tersedia: {fallback}]", samples)

    # Source line (template has it between figure and caption)
    src_para = _new_paragraph(doc, samples["figure_source_ppr"])
    year = datetime.now().year
    src_para.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(src_para, before_pt=0.0, after_pt=3.0)
    _append_rich_text(
        src_para, f"Source: Authors, {year}", samples["figure_source_rpr"], base_italic=True
    )

    # Caption
    cap_para = _new_paragraph(doc, samples["figure_caption_ppr"])
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_para.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(cap_para, before_pt=3.0, after_pt=6.0)
    _add_sample_run(cap_para, f"Figure {number}", samples["figure_label_rpr"], bold=True)
    _add_sample_run(cap_para, ". ", samples["figure_text_rpr"])
    _append_rich_text(cap_para, title, samples["figure_text_rpr"], base_bold=False)


def _clear_cell(cell) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag != qn("w:tcPr"):
            tc.remove(child)
    tc.append(OxmlElement("w:p"))


def _fill_cell_text(
    cell,
    text: str,
    sample_ppr: etree._Element | None,
    sample_rpr: etree._Element | None,
    *,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    base_bold: bool = False,
    font_size_pt: float = 10.0,
) -> None:
    _clear_cell(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    _apply_sample_ppr(paragraph, sample_ppr)
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    # When cell rPr samples are unavailable (template table is empty), reuse body rPr.
    rpr = sample_rpr
    _append_rich_text(paragraph, text, rpr, base_bold=base_bold)
    for run in paragraph.runs:
        run.font.size = Pt(font_size_pt)


def _set_repeat_table_header(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tbl_header = trPr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        trPr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def _add_table(
    doc: Document,
    item: dict,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
) -> None:
    headers = [str(value) for value in (item.get("Headers") or item.get("headers") or [])]
    rows = [list(map(str, row)) for row in (item.get("Rows") or item.get("rows") or [])]
    if not headers:
        return

    state.table_number += 1
    number = str(
        item.get("TableNumber") or item.get("number") or state.table_number
    ).strip() or str(state.table_number)
    title = str(item.get("Title") or item.get("title") or f"Table title {number}").strip()

    # Caption (Table title above)
    cap_para = _new_paragraph(doc, samples["table_caption_ppr"])
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_para.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(cap_para, before_pt=6.0, after_pt=3.0)
    _add_sample_run(cap_para, f"Table {number}", samples["table_label_rpr"], bold=True)
    _add_sample_run(cap_para, ". ", samples["table_text_rpr"])
    _append_rich_text(cap_para, title, samples["table_text_rpr"], base_bold=False)

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))

    _set_table_full_borders(table)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Repeat header row on each page.
    _set_repeat_table_header(table.rows[0])

    header_cell_ppr = samples.get("table_header_cell_ppr")
    body_cell_ppr = samples.get("table_body_cell_ppr")

    # Reuse paragraph rPr from body as safe default.
    default_rpr = samples.get("body_rpr")

    for col_index, header in enumerate(headers):
        cell = table.rows[0].cells[col_index]
        _fill_cell_text(
            cell,
            header,
            header_cell_ppr,
            default_rpr,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            base_bold=True,
            font_size_pt=10.0,
        )

    for row_index, row_values in enumerate(rows, start=1):
        for col_index in range(len(headers)):
            value = row_values[col_index] if col_index < len(row_values) else ""
            cell = table.rows[row_index].cells[col_index]
            align = WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
            _fill_cell_text(
                cell,
                value,
                body_cell_ppr,
                default_rpr,
                align=align,
                base_bold=False,
                font_size_pt=10.0,
            )

    # Table source line (template has it under the table)
    src_para = _new_paragraph(doc, samples["table_source_ppr"])
    src_para.paragraph_format.first_line_indent = Pt(0)
    year = datetime.now().year
    _set_spacing_min(src_para, before_pt=0.0, after_pt=3.0)
    _append_rich_text(
        src_para, f"Source: Authors, {year}", samples["table_source_rpr"], base_italic=True
    )


def _add_equation(
    doc: Document,
    item: dict,
    samples: dict[str, etree._Element | None],
    state: RenderState,
) -> None:
    formula = str(item.get("latex") or item.get("text") or "").strip()
    if not formula:
        return

    state.equation_number += 1
    number = str(item.get("FormulaNumber") or state.equation_number).strip() or str(
        state.equation_number
    )

    paragraph = _new_paragraph(doc, samples["equation_ppr"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(paragraph, before_pt=3.0, after_pt=3.0)

    # Tab stops to center equation and right-align equation label.
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(CENTER_TAB_PT), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(RIGHT_TAB_PT), WD_TAB_ALIGNMENT.RIGHT)

    paragraph.add_run().add_tab()
    if not _append_inline_math(paragraph, formula):
        _append_rich_text(paragraph, formula, samples["body_rpr"], base_italic=True)
    paragraph.add_run().add_tab()
    _add_sample_run(paragraph, f"(Eq.{number})", samples["equation_label_rpr"], italic=True)


def _add_equation_where(
    doc: Document, text: str, samples: dict[str, etree._Element | None]
) -> None:
    paragraph = _new_paragraph(doc, samples["equation_where_ppr"])
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _set_spacing_min(paragraph, before_pt=0.0, after_pt=3.0)
    _append_rich_text(paragraph, text, samples["equation_where_rpr"], base_bold=False)


def _render_content_item(
    doc: Document,
    item,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
) -> None:
    if isinstance(item, str):
        text = item.strip()
        if text:
            _body_paragraph(doc, text, samples)
        return

    if not isinstance(item, dict):
        return

    item_id = str(item.get("id") or "").strip().lower()
    if item_id == "text":
        text = str(item.get("text") or "").strip()
        if text:
            _body_paragraph(doc, text, samples)
        return

    if item_id in {"gambar", "image", "figure"}:
        _add_figure(doc, item, json_path, samples, state)
        return

    if item_id in {"tabel", "table"}:
        _add_table(doc, item, json_path, samples, state)
        return

    if item_id in {"rumus", "formula", "equation"}:
        _add_equation(doc, item, samples, state)
        return


def _render_content_sequence(
    doc: Document,
    content,
    json_path: Path,
    samples: dict[str, etree._Element | None],
    state: RenderState,
) -> None:
    if content is None:
        return
    if isinstance(content, str):
        _body_paragraph(doc, content, samples)
        return
    if isinstance(content, list):
        for item in content:
            _render_content_item(doc, item, json_path, samples, state)
        return


_SECTION_RE = re.compile(r"^section(\d+)$", re.IGNORECASE)
_SUBSECTION_RE = re.compile(r"^section(\d+)([a-z]+)$", re.IGNORECASE)


def _iter_section_keys(config: dict) -> list[str]:
    keys = []
    for key in config.keys():
        m = _SECTION_RE.match(str(key))
        if m:
            keys.append(key)
    keys.sort(key=lambda k: int(_SECTION_RE.match(k).group(1)))
    return keys


def _iter_subsection_keys(section_obj: dict, section_number: int) -> list[str]:
    keys = []
    prefix = f"section{section_number}".lower()
    for key in section_obj.keys():
        k = str(key).lower()
        if k.startswith(prefix) and k != prefix:
            # section2a, section2b, ...
            if _SUBSECTION_RE.match(k):
                keys.append(key)
    keys.sort(key=lambda k: str(k).lower())
    return keys


def _render_sections(
    doc: Document,
    config: dict,
    json_path: Path,
    samples: dict[str, etree._Element | None],
) -> None:
    state = RenderState()

    for skey in _iter_section_keys(config):
        section = config.get(skey)
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or section.get("Title") or "").strip()
        if not title:
            title = f"Section {skey}"
        _add_section_heading(doc, title, samples)

        _render_content_sequence(doc, section.get("content"), json_path, samples, state)

        sec_num = int(_SECTION_RE.match(skey).group(1))
        for subkey in _iter_subsection_keys(section, sec_num):
            subsection = section.get(subkey)
            if not isinstance(subsection, dict):
                continue
            stitle = str(subsection.get("title") or subsection.get("Title") or "").strip()
            if not stitle:
                stitle = f"Subsection {subkey}"
            _add_subsection_heading(doc, stitle, samples)
            _render_content_sequence(doc, subsection.get("content"), json_path, samples, state)

    # Acknowledgments (always present per template)
    _add_section_heading(doc, "Acknowledgments", samples)
    _body_paragraph(doc, _acknowledgment_text(config), samples)

    # References (always present)
    refs = config.get("references")
    ref_items: list[str] = []
    if isinstance(refs, dict):
        content = refs.get("content")
        if isinstance(content, list):
            for entry in content:
                if isinstance(entry, str) and entry.strip():
                    ref_items.append(entry.strip())
                elif isinstance(entry, dict):
                    text = str(entry.get("text") or entry.get("Text") or "").strip()
                    if text:
                        ref_items.append(text)

    _add_section_heading(doc, "References", samples)
    if not ref_items:
        p = _new_paragraph(doc, samples["reference_item_ppr"])
        _append_rich_text(
            p,
            "Tambahkan referensi sesuai gaya ICOSEG.",
            samples["reference_item_rpr"],
            base_bold=False,
        )
        return

    for item in ref_items:
        p = _new_paragraph(doc, samples["reference_item_ppr"])
        _append_rich_text(p, item, samples["reference_item_rpr"], base_bold=False)


def _collect_json_stats(config: dict) -> dict[str, int]:
    stats = {"fig": 0, "tbl": 0, "eq": 0, "text": 0}

    def walk(node):
        if isinstance(node, dict):
            # content item
            if "id" in node:
                item_id = str(node.get("id") or "").strip().lower()
                if item_id == "text":
                    stats["text"] += 1
                elif item_id in {"gambar", "image", "figure"}:
                    stats["fig"] += 1
                elif item_id in {"tabel", "table"}:
                    stats["tbl"] += 1
                elif item_id in {"rumus", "formula", "equation"}:
                    stats["eq"] += 1
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)
        elif isinstance(node, str):
            if node.strip():
                stats["text"] += 1

    walk(config)
    return stats


def verify_document(template_path: Path, out_path: Path, config: dict) -> tuple[bool, list[str]]:
    issues: list[str] = []

    # 1) Header/footer should remain semantically the same.
    #    Note: python-docx may reserialize these parts on save, so byte equality
    #    is too strict. We compare extracted text + paragraph count instead.
    def _hf_signature(z: zipfile.ZipFile, inner: str) -> tuple[str, int] | None:
        if inner not in z.namelist():
            return None
        raw = z.read(inner)
        if b"purl.oclc.org/ooxml" in raw:
            raw2 = _strict_to_trans(raw)
        else:
            raw2 = raw
        root = etree.fromstring(raw2)
        text = "".join(t.text for t in root.iter(_wq("t")) if t.text).strip()
        npara = len(root.findall(f".//{_wq('p')}"))
        return text, npara

    with zipfile.ZipFile(template_path) as zt, zipfile.ZipFile(out_path) as zo:
        for part in (
            "word/header1.xml",
            "word/footer1.xml",
            "word/footer2.xml",
            "word/footer3.xml",
        ):
            sig_t = _hf_signature(zt, part)
            sig_o = _hf_signature(zo, part)
            if sig_t is None or sig_o is None:
                continue
            if sig_t != sig_o:
                issues.append(
                    f"Header/footer semantic mismatch: {part} (template={sig_t}, output={sig_o})"
                )

    # 2) Basic content presence
    with zipfile.ZipFile(out_path) as z:
        raw = z.read("word/document.xml")
        if b"purl.oclc.org/ooxml" in raw:
            raw = _strict_to_trans(raw)
        root = etree.fromstring(raw)
        text = "".join(t.text for t in root.iter(_wq("t")) if t.text)

        # 2b) Logo presence: ensure the original ICOSEG logo (image1.png) is still referenced.
        try:
            rels_raw = z.read("word/_rels/document.xml.rels")
            rels_root = etree.fromstring(rels_raw)
            pkg_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            logo_rid = None
            for rel in rels_root.findall(f".//{{{pkg_ns}}}Relationship"):
                target = (rel.get("Target") or "").replace("\\\\", "/")
                if target.endswith("media/image1.png"):
                    logo_rid = rel.get("Id")
                    break
            if logo_rid is not None:
                if logo_rid.encode("utf-8") not in raw:
                    issues.append(
                        "ICOSEG logo missing: image1.png relationship not referenced in document.xml"
                    )
        except Exception:
            # Don't hard-fail verification if relationship parsing isn't possible.
            pass

    title = _title_text(config)
    if title and title[:40] not in text:
        issues.append("Title from JSON not found in output document text")

    for author in _parse_authors(config):
        name = author.get("name", "")
        if name and name not in text:
            issues.append(f"Author name missing: {name}")

    # 3) Count figures/tables/equations rendered (heuristic-based)
    want = _collect_json_stats(config)

    with zipfile.ZipFile(out_path) as z:
        raw = z.read("word/document.xml")
        if b"purl.oclc.org/ooxml" in raw:
            raw = _strict_to_trans(raw)
        root = etree.fromstring(raw)
        body = root.find(_wq("body"))
        paragraphs = body.findall(_wq("p")) if body is not None else []
        tables = body.findall(_wq("tbl")) if body is not None else []

    got_fig = sum(
        1 for p in paragraphs if _text_of_paragraph(p).strip().lower().startswith("figure ")
    )
    got_tbl = sum(
        1 for p in paragraphs if _text_of_paragraph(p).strip().lower().startswith("table ")
    )
    got_eq = sum(1 for p in paragraphs if "(eq." in _text_of_paragraph(p).lower())

    if got_fig < want["fig"]:
        issues.append(f"Figure captions found {got_fig} < expected {want['fig']}")
    if got_tbl < want["tbl"]:
        issues.append(f"Table captions found {got_tbl} < expected {want['tbl']}")
    if got_eq < want["eq"]:
        issues.append(f"Equations found {got_eq} < expected {want['eq']}")

    # 3b) Ensure every table repeats its header row.
    for idx, tbl in enumerate(tables, start=1):
        first_tr = tbl.find(_wq("tr"))
        if first_tr is None:
            continue
        trpr = first_tr.find(_wq("trPr"))
        tbl_header = trpr.find(_wq("tblHeader")) if trpr is not None else None
        if tbl_header is None:
            issues.append(f"Table #{idx} does not have repeat header enabled (missing w:tblHeader)")

    # 4) Spacing minima checks (twips):
    #    - Table caption: before>=6pt (120tw), after>=3pt (60tw)
    #    - Figure caption: before>=3pt (60tw), after>=6pt (120tw)
    #    - Section/subsection: before>=3pt (60tw), after>=3pt (60tw)
    #    - Equation: before>=3pt (60tw), after>=3pt (60tw)
    def twips(v: str | None) -> int:
        try:
            return int(v) if v is not None else 0
        except Exception:
            return 0

    def spacing_of(p_el: etree._Element) -> tuple[int, int]:
        ppr = p_el.find(_wq("pPr"))
        if ppr is None:
            return 0, 0
        sp = ppr.find(_wq("spacing"))
        if sp is None:
            return 0, 0
        return twips(sp.get(_wq("before"))), twips(sp.get(_wq("after")))

    def ilvl_of(p_el: etree._Element) -> int | None:
        ppr = p_el.find(_wq("pPr"))
        if ppr is None:
            return None
        numpr = ppr.find(_wq("numPr"))
        if numpr is None:
            return None
        ilvl = numpr.find(_wq("ilvl"))
        if ilvl is None:
            return None
        try:
            return int(ilvl.get(_wq("val")))
        except Exception:
            return None

    for p in paragraphs:
        t = _text_of_paragraph(p).strip()
        lower = t.lower()
        before, after = spacing_of(p)
        if lower.startswith("table "):
            if before < 120 or after < 60:
                issues.append(
                    f"Table caption spacing too small: before={before}, after={after}, text='{t[:60]}'"
                )
        if lower.startswith("figure "):
            if before < 60 or after < 120:
                issues.append(
                    f"Figure caption spacing too small: before={before}, after={after}, text='{t[:60]}'"
                )
        if "(eq." in lower:
            if before < 60 or after < 60:
                issues.append(
                    f"Equation spacing too small: before={before}, after={after}, text='{t[:60]}'"
                )
        ilvl = ilvl_of(p)
        if ilvl in (0, 1):
            if before < 60 or after < 60:
                issues.append(
                    f"Heading spacing too small (ilvl={ilvl}): before={before}, after={after}, text='{t[:60]}'"
                )

    ok = not issues
    return ok, issues


def build_document(
    json_path: Path = JSON_PATH,
    output_path: Path | None = None,
    template_path: Path = TEMPLATE_PATH,
    *,
    verify: bool = False,
) -> Path:
    config = json.loads(Path(json_path).read_text(encoding="utf-8"))
    final_output = (
        Path(output_path)
        if output_path is not None
        else Path(json_path).parent / f"{JOURNAL_NAME}_output.docx"
    )
    final_output.parent.mkdir(parents=True, exist_ok=True)

    samples = _load_template_samples(template_path)
    if samples["final_sectpr"] is None:
        raise RuntimeError("Template ICOSEG tidak memiliki final sectPr.")

    shutil.copy(str(template_path), str(final_output))
    doc = Document(str(final_output))
    _clear_document_body(doc)
    _set_document_final_sectpr(doc, samples["final_sectpr"])

    _add_front_matter(doc, config, samples)
    _render_sections(doc, config, Path(json_path), samples)

    doc.save(str(final_output))

    if verify:
        ok, issues = verify_document(template_path, final_output, config)
        if not ok:
            raise RuntimeError("Verification failed:\n- " + "\n- ".join(issues))

    print(f"Generated: {final_output}")
    return final_output


def main() -> None:
    verify = "--verify" in sys.argv
    args = [a for a in sys.argv[1:] if a != "--verify"]

    json_arg = Path(args[0]) if len(args) >= 1 else JSON_PATH
    out_arg = Path(args[1]) if len(args) >= 2 else None
    tpl_arg = Path(args[2]) if len(args) >= 3 else TEMPLATE_PATH

    if not json_arg.exists():
        raise SystemExit(f"JSON tidak ditemukan: {json_arg}")
    if not tpl_arg.exists():
        raise SystemExit(f"Template tidak ditemukan: {tpl_arg}")

    build_document(json_arg, out_arg, tpl_arg, verify=verify)


def _set_table_full_borders(table) -> None:
    """Pastikan tabel punya border tegas/visible (val=single, sz=4 = 0.5pt).

    Dipanggil setelah doc.add_table() supaya tabel data keliatan di Word.
    Auto-injected oleh _fix_table_borders.py untuk lulus audit border check.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_borders = tbl_pr.find(qn("w:tblBorders"))
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = tbl_borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tbl_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


if __name__ == "__main__":
    main()
